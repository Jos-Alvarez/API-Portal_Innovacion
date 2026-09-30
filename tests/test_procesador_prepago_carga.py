"""Pruebas del procesador `Prepago_Carga` (transferencias entre cuentas).

El extracto real no está en este repositorio y no puede estarlo (ADR 0015),
así que estas pruebas arman libros sintéticos con `openpyxl`, guardados sin
extensión (`entrada_0`) como los temporales reales de `app/recepcion.py`.
Prueban las **reglas**; la paridad es `tests/paridad/test_prepago_carga.py`.
"""

from __future__ import annotations

import csv
from datetime import date, datetime
from pathlib import Path
from typing import Any

import pytest
from openpyxl import Workbook, load_workbook

from app.core.errores import ErrorContenido, TipoError
from app.core.tipos import ArchivoEntrada
from app.procesadores.prepago_carga.modulo import (
    COLUMNAS,
    CUENTA_COMISION,
    FIN_DE_LINEA,
    HOJA_DE_SALIDA,
    MOTIVO_BANCO_ENTRADA,
    MOTIVO_FECHA_CONTABLE,
    MOTIVO_FECHA_DOCUMENTO,
    NOMBRE_DESCARTES,
    NOMBRE_TXT,
    NOMBRE_XLSX,
    PrepagoCarga,
    _a_importe,
    _celda_de_txt,
)

_SALIDA = 1101017075
_ENTRADA = 1101017016
_TEXTO = "TRANSFERENCIA ENTRE CUENTAS"

_ENCABEZADO = [
    COLUMNAS.banco_salida,
    COLUMNAS.banco_entrada,
    COLUMNAS.fecha_documento,
    COLUMNAS.fecha_contable,
    COLUMNAS.importe_salida,
    "TIPO",
    COLUMNAS.comision,
    COLUMNAS.importe_entrada,
    "TIPO BANCO SALIDA",
    "TIPO BANCO ENTRADA",
]


def _fila(
    *,
    salida: object = _SALIDA,
    entrada: object = _ENTRADA,
    documento: object = "13.02.2026",
    contable: object = "16.03.2026",
    importe_salida: object = 120161.9,
    comision: object = 4.3,
    importe_entrada: object = 120157.59999999999,
) -> list[object]:
    return [
        salida,
        entrada,
        documento,
        contable,
        importe_salida,
        "SALIDA",
        comision,
        importe_entrada,
        "BCP",
        "IBK",
    ]


def _libro(tmp_path: Path, filas: list[list[object]], encabezado: list[str] | None = None) -> Path:
    """Guarda el libro SIN extensión, como `app/recepcion.py` guarda los temporales."""
    ruta = tmp_path / "entrada_0"
    libro = Workbook()
    hoja = libro.active
    assert hoja is not None
    hoja.append(encabezado if encabezado is not None else _ENCABEZADO)
    for fila in filas:
        hoja.append(fila)
    with ruta.open("wb") as destino:
        libro.save(destino)
    return ruta


def _entrada(ruta: Path, nombre: str = "transferencias.xlsx") -> ArchivoEntrada:
    return ArchivoEntrada(
        nombre_original=nombre,
        ruta_temporal=ruta,
        tamano_comprimido=ruta.stat().st_size,
        formato="xlsx",
    )


def _salidas_por_nombre(procesador: PrepagoCarga, entrada: ArchivoEntrada) -> dict[str, Path]:
    return {s.nombre_propuesto: s.ruta_temporal for s in procesador.procesar([entrada])}


def _filas_xlsx(ruta: Path) -> list[list[Any]]:
    libro = load_workbook(ruta)
    try:
        return [list(f) for f in libro.worksheets[0].iter_rows(values_only=True)]
    finally:
        libro.close()


def _procesar(
    procesador: PrepagoCarga, tmp_path: Path, filas: list[list[object]]
) -> list[list[Any]]:
    salidas = _salidas_por_nombre(procesador, _entrada(_libro(tmp_path, filas)))
    return _filas_xlsx(salidas[NOMBRE_XLSX])


@pytest.fixture
def procesador() -> PrepagoCarga:
    return PrepagoCarga()


# --- El camino feliz y el layout de las cuatro líneas ------------------------


def test_con_comision_un_registro_produce_cinco_filas(
    procesador: PrepagoCarga, tmp_path: Path
) -> None:
    """Cabecera, salida, entrada, comisión y separadora, de 26 columnas cada una."""
    salidas = _salidas_por_nombre(procesador, _entrada(_libro(tmp_path, [_fila()])))
    libro = load_workbook(salidas[NOMBRE_XLSX])
    assert libro.sheetnames == [HOJA_DE_SALIDA]
    filas = [list(f) for f in libro.worksheets[0].iter_rows(values_only=True)]
    assert len(filas) == 5
    assert all(len(f) == 26 for f in filas)


def test_el_layout_de_las_cuatro_lineas_es_el_del_script(
    procesador: PrepagoCarga, tmp_path: Path
) -> None:
    """Las posiciones exactas. Es la regla de negocio, no un detalle de formato."""
    cabecera, salida, entrada, comision, separadora = _procesar(procesador, tmp_path, [_fila()])

    assert cabecera[0] == 1
    assert cabecera[2] == "VPR1"
    assert cabecera[4] == 2026
    assert cabecera[5] == "SA"
    assert cabecera[6] == "20260213"  # fecha documento
    assert cabecera[7] == "20260316"  # fecha contable
    assert cabecera[8] == "03"  # el mes sale de la contable
    assert cabecera[10] == "TRANS ENTRE CUENTAS"
    assert cabecera[11] == "PEN"
    ocupadas = {0, 2, 4, 5, 6, 7, 8, 10, 11}
    assert all(cabecera[i] is None for i in range(26) if i not in ocupadas)

    assert [salida[0], salida[12], salida[13], salida[17], salida[25]] == [
        2,
        50,
        _SALIDA,
        120161.9,
        _TEXTO,
    ]
    assert [entrada[0], entrada[12], entrada[13], entrada[17], entrada[25]] == [
        2,
        40,
        _ENTRADA,
        120157.6,  # round(abs(120157.59999999999), 2)
        _TEXTO,
    ]
    assert [comision[0], comision[12], comision[13], comision[17]] == [
        2,
        40,
        CUENTA_COMISION,
        4.3,
    ]
    assert comision[25] is None  # la línea de comisión NO lleva texto
    assert all(celda is None for celda in separadora)


def test_el_anio_sale_del_documento_y_el_mes_de_la_contable(
    procesador: PrepagoCarga, tmp_path: Path
) -> None:
    """Rareza del script conservada: en un cruce de año se nota."""
    fila = _fila(documento="31.12.2025", contable="02.01.2026")
    cabecera = _procesar(procesador, tmp_path, [fila])[0]
    assert cabecera[4] == 2025
    assert cabecera[8] == "01"
    assert cabecera[6] == "20251231"
    assert cabecera[7] == "20260102"


@pytest.mark.parametrize("comision", [0, 0.0, None, "", "no es un número"])
def test_sin_comision_no_hay_cuarta_linea(
    procesador: PrepagoCarga, tmp_path: Path, comision: object
) -> None:
    """Cero, vacío o no parseable: la comisión vale 0 y la línea no se emite."""
    filas = _procesar(procesador, tmp_path, [_fila(comision=comision)])
    assert len(filas) == 4
    assert CUENTA_COMISION not in [f[13] for f in filas]


def test_una_comision_negativa_se_emite_en_valor_absoluto(
    procesador: PrepagoCarga, tmp_path: Path
) -> None:
    """El script compara `!= 0.0`, no `> 0`: una comisión negativa también sale."""
    filas = _procesar(procesador, tmp_path, [_fila(comision=-7.126)])
    assert len(filas) == 5
    assert filas[3][13] == CUENTA_COMISION
    assert filas[3][17] == 7.13


def test_los_importes_salen_positivos_y_redondeados(
    procesador: PrepagoCarga, tmp_path: Path
) -> None:
    fila = _fila(importe_salida=-1000.456, importe_entrada="1,000.451")
    filas = _procesar(procesador, tmp_path, [fila])
    assert filas[1][17] == 1000.46
    assert filas[2][17] == 1000.45  # la coma de miles se quita antes de parsear


def test_un_banco_float_se_vuelca_como_entero(procesador: PrepagoCarga, tmp_path: Path) -> None:
    """`int(banco) if isinstance(banco, float)`: trunca, y el TXT no lleva decimales.

    openpyxl devuelve `int` para un número entero guardado en el libro, así
    que la rama sólo se ejercita con decimales: `int()` los TRUNCA, como en
    el script.
    """
    fila = _fila(salida=1101017075.7, entrada=1101017016.2)
    salidas = _salidas_por_nombre(procesador, _entrada(_libro(tmp_path, [fila])))
    lineas = salidas[NOMBRE_TXT].read_bytes().decode("utf-8").split(FIN_DE_LINEA)
    assert lineas[1].split("\t")[13] == "1101017075"
    assert lineas[2].split("\t")[13] == "1101017016"


def test_un_banco_de_texto_se_vuelca_tal_cual(procesador: PrepagoCarga, tmp_path: Path) -> None:
    filas = _procesar(procesador, tmp_path, [_fila(entrada="CTA-001")])
    assert filas[2][13] == "CTA-001"


def test_varios_registros_se_encadenan_en_orden(procesador: PrepagoCarga, tmp_path: Path) -> None:
    filas = _procesar(
        procesador,
        tmp_path,
        [_fila(), _fila(salida=1101017085, comision=None), _fila(salida=1101017005)],
    )
    assert len(filas) == 5 + 4 + 5
    assert [filas[i][13] for i in (1, 6, 10)] == [_SALIDA, 1101017085, 1101017005]


# --- Columnas por nombre -----------------------------------------------------


@pytest.mark.parametrize(
    "faltante",
    [
        COLUMNAS.banco_salida,
        COLUMNAS.banco_entrada,
        COLUMNAS.fecha_documento,
        COLUMNAS.fecha_contable,
        COLUMNAS.importe_salida,
        COLUMNAS.comision,
        COLUMNAS.importe_entrada,
    ],
)
def test_falta_una_columna_obligatoria(
    procesador: PrepagoCarga, tmp_path: Path, faltante: str
) -> None:
    encabezado = [c for c in _ENCABEZADO if c != faltante]
    error = procesador.validar([_entrada(_libro(tmp_path, [], encabezado=encabezado))])
    assert isinstance(error, ErrorContenido)
    assert error.tipo is TipoError.CONTENIDO
    assert error.contexto == {
        "archivo": "transferencias.xlsx",
        "motivo": "columna_faltante",
        "columna": faltante,
    }


def test_la_comision_corregida_no_se_reconoce(procesador: PrepagoCarga, tmp_path: Path) -> None:
    """El encabezado real dice `COMISON`. "Corregirlo" en el Excel rompe la lectura."""
    encabezado = ["COMISION" if c == COLUMNAS.comision else c for c in _ENCABEZADO]
    error = procesador.validar([_entrada(_libro(tmp_path, [], encabezado=encabezado))])
    assert isinstance(error, ErrorContenido)
    assert error.contexto == {
        "archivo": "transferencias.xlsx",
        "motivo": "columna_faltante",
        "columna": "COMISON",
    }


def test_una_columna_insertada_no_corre_la_lectura(
    procesador: PrepagoCarga, tmp_path: Path
) -> None:
    encabezado = ["Columna nueva", *_ENCABEZADO]
    ruta = _libro(tmp_path, [["basura", *_fila()]], encabezado=encabezado)
    assert procesador.validar([_entrada(ruta)]) is None
    filas = _filas_xlsx(_salidas_por_nombre(procesador, _entrada(ruta))[NOMBRE_XLSX])
    assert filas[1][13] == _SALIDA


def test_los_encabezados_se_comparan_tras_strip(procesador: PrepagoCarga, tmp_path: Path) -> None:
    encabezado = [f"  {c} " for c in _ENCABEZADO]
    ruta = _libro(tmp_path, [_fila()], encabezado=encabezado)
    assert procesador.validar([_entrada(ruta)]) is None


# --- TXT: formato y fin de línea ---------------------------------------------


def test_el_txt_usa_crlf_explicito(procesador: PrepagoCarga, tmp_path: Path) -> None:
    salidas = _salidas_por_nombre(procesador, _entrada(_libro(tmp_path, [_fila()])))
    crudo = salidas[NOMBRE_TXT].read_bytes()
    assert FIN_DE_LINEA.encode() == b"\r\n"
    assert crudo.endswith(b"\r\n")
    assert crudo.count(b"\r\n") == 5
    assert b"\n" not in crudo.replace(b"\r\n", b"")


def test_el_txt_es_tabulado_y_con_el_formato_del_script(
    procesador: PrepagoCarga, tmp_path: Path
) -> None:
    fila = _fila(comision=20, importe_salida=1000, importe_entrada=980)
    salidas = _salidas_por_nombre(procesador, _entrada(_libro(tmp_path, [fila])))
    with salidas[NOMBRE_TXT].open(encoding="utf-8", newline="") as archivo:
        filas = list(csv.reader(archivo, delimiter="\t", lineterminator="\r\n"))
    assert all(len(f) == 26 for f in filas)
    assert filas[0][:12] == [
        "1",
        "",
        "VPR1",
        "",
        "2026",
        "SA",
        "20260213",
        "20260316",
        "03",
        "",
        "TRANS ENTRE CUENTAS",
        "PEN",
    ]
    # Los importes son float: un entero se vuelca con `.0`, como en el script.
    assert (filas[1][13], filas[1][17]) == ("1101017075", "1000.0")
    assert (filas[2][13], filas[2][17]) == ("1101017016", "980.0")
    assert (filas[3][13], filas[3][17], filas[3][25]) == ("7101013001", "20.0", "")
    assert filas[4] == [""] * 26


@pytest.mark.parametrize(
    ("celda", "esperado"),
    [
        ("", ""),
        (1234.0, "1234.0"),
        (2000000000.0, "2000000000"),
        (7101013001, "7101013001"),
        (4.3, "4.3"),
        ("texto", "texto"),
    ],
)
def test_formateo_de_celda_del_txt(celda: object, esperado: str) -> None:
    assert _celda_de_txt(celda) == esperado


@pytest.mark.parametrize(
    ("crudo", "esperado"),
    [
        ("1,234.56", 1234.56),
        ("  -42.0 ", -42.0),
        ("", 0.0),
        ("-", 0.0),
        ("NA", 0.0),
        ("nan", 0.0),
        ("no es un número", 0.0),
        (None, 0.0),
        (20, 20.0),
        (1250.5, 1250.5),
    ],
)
def test_limpieza_de_importes(crudo: object, esperado: float) -> None:
    """Equivalente de `limpiar_monto`: nulo o no parseable vale cero, siempre `float`."""
    resultado = _a_importe(crudo)
    assert resultado == esperado
    assert isinstance(resultado, float)


# --- Fechas ------------------------------------------------------------------


@pytest.mark.parametrize(
    ("crudo", "esperado"),
    [
        ("13.02.2026", "20260213"),
        ("13/02/2026", "20260213"),
        ("13-02-2026", "20260213"),
        ("2026-02-13", "20260213"),
        ("01.02.2026", "20260201"),  # día primero: 1 de febrero
        (datetime(2026, 2, 13, 10, 30), "20260213"),
        (date(2026, 2, 13), "20260213"),
    ],
)
def test_fechas_se_leen_con_el_dia_primero(
    procesador: PrepagoCarga, tmp_path: Path, crudo: object, esperado: str
) -> None:
    cabecera = _procesar(procesador, tmp_path, [_fila(documento=crudo, contable=crudo)])[0]
    assert cabecera[6] == esperado
    assert cabecera[7] == esperado


# --- Descartes ---------------------------------------------------------------


def test_sin_descartes_no_hay_tercer_archivo(procesador: PrepagoCarga, tmp_path: Path) -> None:
    salidas = _salidas_por_nombre(procesador, _entrada(_libro(tmp_path, [_fila()])))
    assert sorted(salidas) == sorted((NOMBRE_XLSX, NOMBRE_TXT))


@pytest.mark.parametrize(
    ("fila", "motivo", "valor"),
    [
        (_fila(entrada=None), MOTIVO_BANCO_ENTRADA, ""),
        (_fila(entrada="NA"), MOTIVO_BANCO_ENTRADA, "NA"),
        (_fila(documento="no es fecha"), MOTIVO_FECHA_DOCUMENTO, "no es fecha"),
        (_fila(documento=None), MOTIVO_FECHA_DOCUMENTO, ""),
        (_fila(documento=45000), MOTIVO_FECHA_DOCUMENTO, "45000"),
        (_fila(contable="32.01.2026"), MOTIVO_FECHA_CONTABLE, "32.01.2026"),
    ],
)
def test_cada_descarte_se_reporta_con_motivo_fila_y_valor(
    procesador: PrepagoCarga, tmp_path: Path, fila: list[object], motivo: str, valor: str
) -> None:
    salidas = _salidas_por_nombre(procesador, _entrada(_libro(tmp_path, [_fila(), fila])))
    reporte = salidas[NOMBRE_DESCARTES].read_bytes().decode("utf-8")
    lineas = reporte.split("\r\n")
    assert lineas[0] == f"Filas descartadas: 1 ({motivo}=1)"
    assert lineas[2] == "fila\tmotivo\tvalor"
    assert lineas[3] == f"3\t{motivo}\t{valor}"  # la fila 3 del Excel
    assert reporte.endswith("\r\n")
    # La fila válida sigue saliendo.
    assert len(_filas_xlsx(salidas[NOMBRE_XLSX])) == 5


def test_el_banco_entrada_se_revisa_antes_que_las_fechas(
    procesador: PrepagoCarga, tmp_path: Path
) -> None:
    """Mismo orden que el script: un solo motivo por fila, el primero que falla."""
    fila = _fila(entrada=None, documento="mal", contable="mal")
    salidas = _salidas_por_nombre(procesador, _entrada(_libro(tmp_path, [_fila(), fila])))
    reporte = salidas[NOMBRE_DESCARTES].read_text(encoding="utf-8")
    assert f"{MOTIVO_BANCO_ENTRADA}=1" in reporte
    assert "Filas descartadas: 1" in reporte


def test_el_reporte_cuenta_por_motivo(procesador: PrepagoCarga, tmp_path: Path) -> None:
    filas = [
        _fila(),
        _fila(entrada=None),
        _fila(documento="x"),
        _fila(contable="y"),
        _fila(contable="z"),
    ]
    salidas = _salidas_por_nombre(procesador, _entrada(_libro(tmp_path, filas)))
    reporte = salidas[NOMBRE_DESCARTES].read_text(encoding="utf-8")
    assert reporte.splitlines()[0] == (
        f"Filas descartadas: 4 ({MOTIVO_BANCO_ENTRADA}=1, "
        f"{MOTIVO_FECHA_CONTABLE}=2, {MOTIVO_FECHA_DOCUMENTO}=1)"
    )


def test_el_descarte_se_registra_en_el_log(
    procesador: PrepagoCarga, tmp_path: Path, caplog: pytest.LogCaptureFixture
) -> None:
    filas = [_fila(), _fila(comision=None), _fila(documento="mal")]
    with caplog.at_level("INFO", logger="app"):
        procesador.procesar([_entrada(_libro(tmp_path, filas))])

    (registro,) = [r for r in caplog.records if getattr(r, "evento", None) == "descartes"]
    assert registro.clave == "prepago-carga"  # type: ignore[attr-defined]
    assert registro.filas_descartadas == 1  # type: ignore[attr-defined]
    assert registro.filas_procesadas == 2  # type: ignore[attr-defined]
    assert registro.por_motivo == {MOTIVO_FECHA_DOCUMENTO: 1}  # type: ignore[attr-defined]


def test_sin_descartes_no_se_registra_nada(
    procesador: PrepagoCarga, tmp_path: Path, caplog: pytest.LogCaptureFixture
) -> None:
    with caplog.at_level("INFO", logger="app"):
        procesador.procesar([_entrada(_libro(tmp_path, [_fila()]))])
    assert not [r for r in caplog.records if getattr(r, "evento", None) == "descartes"]


@pytest.mark.parametrize("vacio", [None, "", "NA"])
def test_las_filas_sin_banco_salida_se_filtran_sin_contar(
    procesador: PrepagoCarga, tmp_path: Path, vacio: object
) -> None:
    """El filtro previo del script (columna A vacía) no es un descarte: es ruido."""
    filas = [_fila(), _fila(salida=vacio, documento="basura")]
    salidas = _salidas_por_nombre(procesador, _entrada(_libro(tmp_path, filas)))
    assert NOMBRE_DESCARTES not in salidas
    assert len(_filas_xlsx(salidas[NOMBRE_XLSX])) == 5


# --- Cero filas y archivos ilegibles -----------------------------------------


def test_cero_filas_procesadas_es_error_de_contenido(
    procesador: PrepagoCarga, tmp_path: Path
) -> None:
    ruta = _libro(tmp_path, [_fila(documento="mal"), _fila(salida=None)])
    with pytest.raises(ErrorContenido) as capturado:
        procesador.procesar([_entrada(ruta)])
    assert capturado.value.contexto == {
        "archivo": "transferencias.xlsx",
        "motivo": "cero_filas",
        "columna": None,
    }
    assert sorted(p.name for p in tmp_path.iterdir()) == ["entrada_0"]  # nada escrito


def test_solo_encabezado_es_cero_filas(procesador: PrepagoCarga, tmp_path: Path) -> None:
    with pytest.raises(ErrorContenido) as capturado:
        procesador.procesar([_entrada(_libro(tmp_path, []))])
    assert capturado.value.contexto["motivo"] == "cero_filas"


def test_un_excel_ilegible_es_error_de_contenido(procesador: PrepagoCarga, tmp_path: Path) -> None:
    ruta = tmp_path / "entrada_0"
    ruta.write_bytes(b"esto no es un xlsx")
    error = procesador.validar([_entrada(ruta)])
    assert isinstance(error, ErrorContenido)
    assert error.contexto["motivo"] == "cero_filas"
    with pytest.raises(ErrorContenido):
        procesador.procesar([_entrada(ruta)])


def test_una_fila_que_rompe_no_se_traga(
    procesador: PrepagoCarga, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """El script tenía `except Exception: continue`. Acá se propaga."""
    import app.procesadores.prepago_carga.modulo as modulo

    def _explota(*_args: Any, **_kwargs: Any) -> float:
        raise RuntimeError("importe imposible")

    monkeypatch.setattr(modulo, "_a_importe", _explota)
    with pytest.raises(RuntimeError, match="importe imposible"):
        procesador.procesar([_entrada(_libro(tmp_path, [_fila()]))])


# --- Coherencia con el registro ----------------------------------------------


def test_la_clave_coincide_con_el_registry() -> None:
    from app.core.contrato import obtener_contrato
    from app.registry import REGISTRY

    procesador = REGISTRY["prepago-carga"]
    assert procesador.clave == "prepago-carga"
    assert isinstance(procesador, PrepagoCarga)

    contrato = obtener_contrato("prepago-carga")
    assert contrato is not None
    assert contrato.entradas_min == 1
    assert contrato.entradas_max == 1
    assert "xlsx" in contrato.formatos_aceptados


# --- De punta a punta, por la ruta real y el proceso hijo --------------------


class TestPuntaAPunta:
    """La tubería completa: HTTP → temporal sin extensión → `spawn` → ZIP.

    Obligatoria para todo procesador nuevo (ver `test_procesador_contado_carga`):
    el temporal sin extensión es una trampa estructural que sólo se ve acá.
    """

    @staticmethod
    def _subir(cliente: Any, contenido: bytes) -> Any:
        return cliente.post(
            "/interno/procesadores/prepago-carga",
            files={"archivos": ("transferencias.xlsx", contenido, "application/octet-stream")},
            headers={"Authorization": "Bearer sentinela-token-de-pruebas-3f9c2a"},
        )

    @staticmethod
    def _bytes_de_libro(filas: list[list[object]]) -> bytes:
        import io

        libro = Workbook()
        hoja = libro.active
        assert hoja is not None
        hoja.append(_ENCABEZADO)
        for fila in filas:
            hoja.append(fila)
        memoria = io.BytesIO()
        libro.save(memoria)
        return memoria.getvalue()

    def test_un_excel_legitimo_vuelve_como_zip(self, token_sentinela: str) -> None:
        import io
        import zipfile

        from fastapi.testclient import TestClient

        from app.main import crear_app

        with TestClient(crear_app()) as cliente:
            respuesta = self._subir(cliente, self._bytes_de_libro([_fila()]))

        assert respuesta.status_code == 200, respuesta.text
        with zipfile.ZipFile(io.BytesIO(respuesta.content)) as paquete:
            assert sorted(paquete.namelist()) == sorted((NOMBRE_XLSX, NOMBRE_TXT))
            assert paquete.read(NOMBRE_TXT).count(b"\r\n") == 5

    def test_el_descarte_viaja_en_el_zip(self, token_sentinela: str) -> None:
        import io
        import zipfile

        from fastapi.testclient import TestClient

        from app.main import crear_app

        filas = [_fila(), _fila(contable="fecha rota")]
        with TestClient(crear_app()) as cliente:
            respuesta = self._subir(cliente, self._bytes_de_libro(filas))

        assert respuesta.status_code == 200, respuesta.text
        with zipfile.ZipFile(io.BytesIO(respuesta.content)) as paquete:
            assert NOMBRE_DESCARTES in paquete.namelist()
            assert b"fecha rota" in paquete.read(NOMBRE_DESCARTES)

    def test_un_archivo_que_no_es_excel_es_422(self, token_sentinela: str) -> None:
        from fastapi.testclient import TestClient

        from app.main import crear_app

        with TestClient(crear_app()) as cliente:
            respuesta = self._subir(cliente, b"esto no es un xlsx")

        assert respuesta.status_code == 422
        assert respuesta.json()["tipo"] == "contenido"
