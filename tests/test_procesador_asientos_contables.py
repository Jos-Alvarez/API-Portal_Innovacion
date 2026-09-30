"""Pruebas del procesador `AsientosContables_Carga`.

Los siete pares reales no están en este repositorio y no pueden estarlo (ADR
0015), así que estas pruebas arman libros sintéticos con `openpyxl`, guardados
con nombres SIN extensión (`entrada_0`) como los que escribe `app/recepcion.py`.
Prueban las **reglas**; la paridad contra las salidas históricas es
`tests/paridad/test_asientos_contables.py`.

Cada prueba de layout fija las posiciones exactas de las líneas porque ésa es
la regla de negocio, no un detalle de formato. Las rarezas del script que se
migraron verbatim tienen su propia prueba: son justo las que alguien
"arreglaría" sin querer.
"""

from __future__ import annotations

import io
import zipfile
from collections.abc import Mapping, Sequence
from datetime import date, datetime
from pathlib import Path
from typing import Any

import pytest
from openpyxl import Workbook, load_workbook

from app.core.errores import ErrorContenido, TipoError
from app.core.tipos import ArchivoEntrada
from app.procesadores.asientos_contables.modulo import (
    CLAVE,
    FIN_DE_LINEA,
    MOTIVO_CATEGORIA,
    MOTIVO_CUENTA,
    MOTIVO_FECHA,
    MOTIVO_FECHA_CIERRE,
    NOMBRE_OBSERVACIONES,
    AsientosContables,
    _a_fecha,
    _a_importe,
    _suma_de_grupo,
    _suma_de_serie,
    nombre_de_salida,
)

HOY = date(2025, 8, 26)
_CIERRE_HOY = "20250826"

_ENCABEZADO_COMISION: list[object] = ["Fecha1", "Fecha2", "Monto", "Categoria", "Fecha Cierre"]
_ENCABEZADO_COMPENSACION: list[object] = [
    "Nº documento",
    "Clase de documento",
    "Fecha de documento",
    "Importe en moneda local",
    "Texto",
    "OBS",
    "Fecha Cierre",
]
_ENCABEZADO_COMPENSACION_IZIPAY: list[object] = [
    "Nº documento",
    "Clase de documento",
    "Fecha de documento",
    "Importe en moneda local",
    "Texto",
    "Referencia",
    "fecha dia",
    "Fecha cierre",
]

Hojas = dict[str, list[list[object]]]
HojasDeEntrada = Mapping[str, Sequence[Sequence[object]]]


def _d(mes: int, dia: int, anio: int = 2025) -> datetime:
    return datetime(anio, mes, dia)


def _comision(
    fecha: object, monto: object, categoria: object, cierre: object = None
) -> list[object]:
    return [fecha, None, monto, categoria, cierre]


def _compensacion(
    clase: object, fecha: object, importe: object, obs: object = None, cierre: object = None
) -> list[object]:
    return ["100000001", clase, fecha, importe, "Tarjeta", obs, cierre]


def _compensacion_izipay(
    clase: object,
    fecha: object,
    importe: object,
    texto: object,
    referencia: object = None,
    dia: object = None,
    cierre: object = None,
) -> list[object]:
    return ["100000001", clase, fecha, importe, texto, referencia, dia, cierre]


def _hojas_estandar(
    tipo: str,
    comision: list[list[object]] | None = None,
    compensacion: list[list[object]] | None = None,
) -> Hojas:
    return {
        tipo: [_ENCABEZADO_COMISION, *(comision or [])],
        "COMPENSACION": [_ENCABEZADO_COMPENSACION, *(compensacion or [])],
        "Hoja1": [["no se lee"]],
    }


def _hojas_izipay(
    comision: list[list[object]] | None = None, compensacion: list[list[object]] | None = None
) -> Hojas:
    return {
        "IZIPAY": [_ENCABEZADO_COMISION, *(comision or [])],
        "COMPENSACION": [_ENCABEZADO_COMPENSACION_IZIPAY, *(compensacion or [])],
    }


def _hojas_efectivo(compensacion: list[list[object]]) -> Hojas:
    return {
        "Abonos BCP": [["Fecha", "Monto"]],
        "Efectivo": [["no se lee"]],
        "COMPENSACION": [_ENCABEZADO_COMPENSACION, *compensacion],
    }


def _bytes_de_libro(hojas: HojasDeEntrada) -> bytes:
    libro = Workbook()
    libro.remove(libro["Sheet"])
    for nombre, filas in hojas.items():
        hoja = libro.create_sheet(nombre)
        for fila in filas:
            hoja.append(fila)
    memoria = io.BytesIO()
    libro.save(memoria)
    return memoria.getvalue()


class _Fabrica:
    """Escribe libros como `entrada_N`, sin extensión, igual que la recepción real."""

    def __init__(self, directorio: Path) -> None:
        self.directorio = directorio
        self.contador = 0

    def entrada(self, hojas: HojasDeEntrada, nombre: str = "archivo.xlsx") -> ArchivoEntrada:
        ruta = self.directorio / f"entrada_{self.contador}"
        self.contador += 1
        ruta.write_bytes(_bytes_de_libro(hojas))
        return ArchivoEntrada(
            nombre_original=nombre,
            ruta_temporal=ruta,
            tamano_comprimido=ruta.stat().st_size,
            formato="xlsx",
        )


@pytest.fixture
def fabrica(tmp_path: Path) -> _Fabrica:
    return _Fabrica(tmp_path)


@pytest.fixture
def procesador() -> AsientosContables:
    return AsientosContables(hoy=lambda: HOY)


def _procesar(procesador: AsientosContables, *entradas: ArchivoEntrada) -> dict[str, Path]:
    assert procesador.validar(list(entradas)) is None
    return {s.nombre_propuesto: s.ruta_temporal for s in procesador.procesar(list(entradas))}


def _leer(ruta: Path) -> dict[str, list[list[Any]]]:
    libro = load_workbook(ruta)
    try:
        return {h.title: [list(f) for f in h.iter_rows(values_only=True)] for h in libro.worksheets}
    finally:
        libro.close()


def _cabecera(fila: list[Any]) -> list[Any]:
    """Las columnas con dato de una línea tipo 1."""
    return [fila[i] for i in (0, 2, 4, 5, 6, 7, 8, 9, 10, 11)]


def _detalle(fila: list[Any]) -> tuple[Any, ...]:
    """Las columnas con dato de una línea tipo 2: tipo, clave, cuenta, monto, Z001, texto."""
    return (fila[0], fila[12], fila[13], fila[17], fila[20], fila[25])


def _es_separadora(fila: list[Any]) -> bool:
    return len(fila) == 26 and all(celda is None for celda in fila)


# --- Hojas estándar: COMISION ------------------------------------------------


def test_comision_estandar_layout_completo(
    procesador: AsientosContables, fabrica: _Fabrica
) -> None:
    """Grupos por fecha ascendente; comisiones, su suma, y los dos signos de extorno."""
    entrada = fabrica.entrada(
        _hojas_estandar(
            "VISA",
            comision=[
                _comision(_d(6, 2), -10.5, "Comisión"),
                _comision(_d(6, 1), -3, "Comisión"),
                _comision(_d(6, 1), 7, "Extorno"),
                _comision(_d(6, 1), -2, "Extorno"),
                _comision(_d(6, 1), 100, "Ventas según SAP"),
                _comision(_d(6, 1), 0, "Comisión"),
            ],
        )
    )
    filas = _leer(_procesar(procesador, entrada)[nombre_de_salida("VISA")])["COMISION"]

    assert all(len(f) == 26 for f in filas)
    assert _cabecera(filas[0]) == [
        1,
        "PEXP",
        2025,
        "AB",
        "20250601",
        _CIERRE_HOY,
        "08",
        "COMISION",
        "VISA",
        "PEN",
    ]
    assert [_detalle(f) for f in filas[1:7]] == [
        (2, 40, 1106012030, 3, None, "VISA"),
        (2, 50, 1101020065, 3, None, "VISA"),
        (2, 40, 1101020064, 7, None, "VISA"),  # extorno positivo
        (2, 50, 2107808010, 7, None, "VISA"),
        (2, 40, 2107808010, 2, None, "VISA"),  # extorno negativo
        (2, 50, 1101020065, 2, None, "VISA"),
    ]
    assert _es_separadora(filas[7])
    assert filas[8][6] == "20250602"
    assert [_detalle(f) for f in filas[9:11]] == [
        (2, 40, 1106012030, 10.5, None, "VISA"),
        (2, 50, 1101020065, 10.5, None, "VISA"),
    ]
    assert _es_separadora(filas[11])
    assert len(filas) == 12


@pytest.mark.parametrize(
    ("tipo", "cuenta"),
    [
        ("AMEX", 1106012033),
        ("DINNERS", 1106012034),
        ("MASTERCARD", 1106012032),
        ("SAFETYPAY", 1106012031),
        ("VISA", 1106012030),
    ],
)
def test_cada_tarjeta_usa_su_cuenta_de_comision(
    procesador: AsientosContables, fabrica: _Fabrica, tipo: str, cuenta: int
) -> None:
    entrada = fabrica.entrada(_hojas_estandar(tipo, comision=[_comision(_d(6, 1), -1, "Comisión")]))
    salidas = _procesar(procesador, entrada)
    filas = _leer(salidas[nombre_de_salida(tipo)])["COMISION"]
    assert filas[1][13] == cuenta
    assert filas[0][10] == tipo


def test_comision_en_minuscula_pasa_el_filtro_pero_no_genera_linea(
    procesador: AsientosContables, fabrica: _Fabrica
) -> None:
    """Rareza verbatim: filtro sin mayúsculas, bucle con `"Comisión"` exacto.

    La fila crea la cabecera del día y no produce detalle; se reporta.
    """
    entrada = fabrica.entrada(
        _hojas_estandar("AMEX", comision=[_comision(_d(6, 1), -5, "comisión cobrada")])
    )
    salidas = _procesar(procesador, entrada)
    filas = _leer(salidas[nombre_de_salida("AMEX")])["COMISION"]
    assert len(filas) == 2
    assert filas[0][0] == 1
    assert _es_separadora(filas[1])
    reporte = salidas[NOMBRE_OBSERVACIONES].read_text(encoding="utf-8")
    assert f"{MOTIVO_CATEGORIA}=1" in reporte
    assert "comisión cobrada" in reporte


def test_la_fecha_de_cierre_real_manda_sobre_hoy(
    procesador: AsientosContables, fabrica: _Fabrica
) -> None:
    """Año y mes de la cabecera salen del CIERRE: la primera no nula del grupo."""
    entrada = fabrica.entrada(
        _hojas_estandar(
            "AMEX",
            comision=[
                _comision(_d(6, 30), -1, "Comisión", cierre=None),
                _comision(_d(6, 30), -1, "Comisión", cierre=_d(7, 1)),
                _comision(_d(6, 30), -1, "Comisión", cierre=_d(9, 9)),
            ],
        )
    )
    cabecera = _leer(_procesar(procesador, entrada)[nombre_de_salida("AMEX")])["COMISION"][0]
    assert (cabecera[4], cabecera[6], cabecera[7], cabecera[8]) == (
        2025,
        "20250630",
        "20250701",
        "07",
    )


def test_sin_fecha_de_cierre_se_usa_el_reloj_inyectado(fabrica: _Fabrica) -> None:
    """La regla real: el cierre vacío es el día de la corrida (`datetime.now()`)."""
    hojas = _hojas_estandar("VISA", comision=[_comision(_d(6, 1), -1, "Comisión")])
    otro_dia = AsientosContables(hoy=lambda: date(2031, 1, 15))
    cabecera = _leer(_procesar(otro_dia, fabrica.entrada(hojas))[nombre_de_salida("VISA")])[
        "COMISION"
    ][0]
    assert (cabecera[4], cabecera[7], cabecera[8]) == (2031, "20310115", "01")


def test_el_reloj_por_defecto_es_el_de_hoy(fabrica: _Fabrica) -> None:
    hojas = _hojas_estandar("VISA", comision=[_comision(_d(6, 1), -1, "Comisión")])
    antes = date.today()
    cabecera = _leer(
        _procesar(AsientosContables(), fabrica.entrada(hojas))["resultado_visa_procesado.xlsx"]
    )["COMISION"][0]
    assert cabecera[7] in {antes.strftime("%Y%m%d"), date.today().strftime("%Y%m%d")}


# --- Hojas estándar: COBRANZA y RECARGA --------------------------------------


def test_cobranza_suma_las_dz_por_fecha(procesador: AsientosContables, fabrica: _Fabrica) -> None:
    """`contains('DZ')` sin mayúsculas: "dz" entra; RG no. Importes en absoluto."""
    entrada = fabrica.entrada(
        _hojas_estandar(
            "MASTERCARD",
            compensacion=[
                _compensacion("DZ", _d(6, 3), -100),
                _compensacion("dz", _d(6, 3), 50),
                _compensacion("RG", _d(6, 3), 999),
                _compensacion("DZ", _d(6, 1), 20),
            ],
        )
    )
    filas = _leer(_procesar(procesador, entrada)[nombre_de_salida("MASTERCARD")])["COBRANZA"]
    assert _cabecera(filas[0]) == [
        1,
        "PEXP",
        2025,
        "DZ",
        "20250601",
        _CIERRE_HOY,
        "08",
        "COBRANZA",
        "MASTERCARD",
        "PEN",
    ]
    assert _detalle(filas[1]) == (2, 40, 1101020064, 20, None, "MASTERCARD")
    assert _detalle(filas[2]) == (2, 11, 1001417, 20, "Z001", "MASTERCARD")
    assert _es_separadora(filas[3])
    assert filas[4][6] == "20250603"
    assert _detalle(filas[5])[3] == 150
    assert len(filas) == 8


def test_recarga_estandar_signos_y_filtros(
    procesador: AsientosContables, fabrica: _Fabrica
) -> None:
    """RG y AJUSTE como subcadenas; el cero cuenta como POSITIVO (`>= 0`)."""
    entrada = fabrica.entrada(
        _hojas_estandar(
            "DINNERS",
            compensacion=[
                _compensacion("RG", _d(6, 5), 30, "AJUSTE"),
                _compensacion("RG", _d(6, 6), 0, "ajuste manual"),
                _compensacion("RG", _d(6, 7), -40, "AJUSTE"),
                _compensacion("DZ", _d(6, 8), 5, "AJUSTE"),
                _compensacion("RG", _d(6, 9), 5, None),
            ],
        )
    )
    filas = _leer(_procesar(procesador, entrada)[nombre_de_salida("DINNERS")])["RECARGA"]
    assert len(filas) == 12
    positiva, cero, negativa = filas[0:4], filas[4:8], filas[8:12]

    assert (positiva[0][5], positiva[0][9]) == ("RG", "RECARGA POSITIVO")
    assert _detalle(positiva[1]) == (2, "01", 1001418, 30, None, "DINNERS")
    assert isinstance(positiva[1][12], str)  # "01" es TEXTO, las demás claves enteros
    assert _detalle(positiva[2]) == (2, 31, 3000013956, 30, None, "DINNERS")

    assert cero[0][9] == "RECARGA POSITIVO"

    assert negativa[0][9] == "RECARGA NEGATIVO"
    assert _detalle(negativa[1]) == (2, 11, 1001418, 40, "Z001", "DINNERS")
    assert _detalle(negativa[2]) == (2, 21, 3000013956, 40, None, "DINNERS")
    assert _es_separadora(negativa[3])


def test_las_hojas_vacias_no_se_escriben(procesador: AsientosContables, fabrica: _Fabrica) -> None:
    """Sólo las hojas con filas, en el orden COMISION, COBRANZA, RECARGA."""
    entrada = fabrica.entrada(
        _hojas_estandar(
            "SAFETYPAY",
            comision=[_comision(_d(6, 1), -1, "Comisión")],
            compensacion=[_compensacion("RG", _d(6, 5), 30, "AJUSTE")],
        )
    )
    assert list(_leer(_procesar(procesador, entrada)[nombre_de_salida("SAFETYPAY")])) == [
        "COMISION",
        "RECARGA",
    ]


# --- IZIPAY ------------------------------------------------------------------


def test_izipay_comision_agrupa_todas_las_filas_y_mapea_cuentas(
    procesador: AsientosContables, fabrica: _Fabrica
) -> None:
    """Cabecera por CADA día (aunque no tenga comisiones) y cuenta por subcadena."""
    entrada = fabrica.entrada(
        _hojas_izipay(
            comision=[
                _comision(_d(3, 12), 0, "Ventas según SAP"),
                _comision(_d(3, 12), -1357.99, "Comisión MC/VISA", cierre=_d(7, 1)),
                _comision(_d(3, 12), -21.31, "Comisión Diners"),
                _comision(_d(3, 12), -79.12, "Comisión Amex"),
                _comision(_d(3, 12), -5, "Comisión Otra"),
                _comision(_d(3, 13), 100, "Ventas según SAP"),
            ]
        )
    )
    salidas = _procesar(procesador, entrada)
    filas = _leer(salidas[nombre_de_salida("IZIPAY")])["COMISION"]

    # La PRIMERA fila del día decide el cierre, esté vacía o no.
    assert filas[0][7] == _CIERRE_HOY
    assert [_detalle(f) for f in filas[1:6]] == [
        (2, 40, 1106012032, 1357.99, None, "IZIPAY"),
        (2, 40, 1106012034, 21.31, None, "IZIPAY"),
        (2, 40, 1106012033, 79.12, None, "IZIPAY"),
        (2, 40, 0, 5, None, "IZIPAY"),  # sin cuenta: 0, como el script
        (2, 50, 1101020065, 1357.99 + 21.31 + 79.12 + 5, None, "IZIPAY"),
    ]
    assert _es_separadora(filas[6])
    assert filas[7][6] == "20250313"  # día sin comisiones: cabecera y separadora
    assert _es_separadora(filas[8])
    assert len(filas) == 9

    reporte = salidas[NOMBRE_OBSERVACIONES].read_text(encoding="utf-8")
    assert f"{MOTIVO_CUENTA}=1" in reporte
    assert "Comisión Otra" in reporte


def test_izipay_cobranza_categorias_y_signo_invertido(
    procesador: AsientosContables, fabrica: _Fabrica
) -> None:
    """DZ EXACTO; categorías por texto exacto; MASTERCARD+IZIPAY negativo → "01"."""
    dia1, dia2 = _d(6, 27), _d(6, 30)
    entrada = fabrica.entrada(
        _hojas_izipay(
            compensacion=[
                _compensacion_izipay("DZ", _d(6, 1), 1000, 12345, dia=dia1),
                _compensacion_izipay("dz", _d(6, 1), 500, 1, dia=dia1),
                _compensacion_izipay("RG", _d(6, 1), 600, "VISA", dia=dia1),
                _compensacion_izipay("RG", _d(6, 1), 50, "MASTERCARD", dia=dia1),
                _compensacion_izipay("RG", _d(6, 1), -80, "IZIPAY", dia=dia1),
                _compensacion_izipay("RG", _d(6, 1), 700, "VISA", dia=None),
                _compensacion_izipay("RG", _d(6, 1), 20, "MASTERCARD", dia=dia2),
                _compensacion_izipay("RG", _d(6, 1), 5, "AMEX", dia=dia2),
            ]
        )
    )
    salidas = _procesar(procesador, entrada)
    filas = _leer(salidas[nombre_de_salida("IZIPAY")])["COBRANZA"]

    assert _cabecera(filas[0])[3:5] == ["DZ", "20250627"]
    assert [_detalle(f) for f in filas[1:4]] == [
        (2, 40, 1101020064, 1000, None, "IZIPAY"),
        (2, 11, 1001416, 600, "Z001", "IZIPAY"),
        (2, "01", 1001417, 30, None, "IZIPAY"),
    ]
    assert _es_separadora(filas[4])
    assert filas[5][6] == "20250630"
    assert [_detalle(f) for f in filas[6:8]] == [
        (2, 11, 1001420, 5, "Z001", "IZIPAY"),
        (2, 11, 1001417, 20, "Z001", "IZIPAY"),
    ]
    assert len(filas) == 9
    # Una `fecha dia` vacía es cómo el operador deja una fila afuera: no se reporta.
    assert NOMBRE_OBSERVACIONES not in salidas


def test_izipay_recarga_ajuste_exacto_y_cero_negativo(
    procesador: AsientosContables, fabrica: _Fabrica
) -> None:
    """`== "AJUSTE"` sin mirar la clase; el cero es NEGATIVO (`> 0`); cuenta 1001417."""
    entrada = fabrica.entrada(
        _hojas_izipay(
            compensacion=[
                _compensacion_izipay("RG", _d(3, 24), 1974, "IZIPAY", "AJUSTE"),
                _compensacion_izipay("XX", _d(3, 25), 0, "IZIPAY", "AJUSTE"),
                _compensacion_izipay("RG", _d(3, 26), 10, "IZIPAY", "ajuste"),
            ]
        )
    )
    filas = _leer(_procesar(procesador, entrada)[nombre_de_salida("IZIPAY")])["RECARGA"]
    assert len(filas) == 8
    assert filas[0][9] == "RECARGA POSITIVO"
    assert _detalle(filas[1]) == (2, "01", 1001417, 1974, None, "IZIPAY")
    assert filas[4][9] == "RECARGA NEGATIVO"
    assert _detalle(filas[5]) == (2, 11, 1001417, 0, "Z001", "IZIPAY")


def test_izipay_lee_fechas_de_texto_con_el_mes_primero(
    procesador: AsientosContables, fabrica: _Fabrica
) -> None:
    """`pd.to_datetime` pelado: "02/03/2025" es 3 de febrero; "13/03/2025", 13 de marzo."""
    entrada = fabrica.entrada(
        _hojas_izipay(
            compensacion=[
                _compensacion_izipay("DZ", _d(1, 1), 1, 1, dia="02/03/2025"),
                _compensacion_izipay("DZ", _d(1, 1), 1, 1, dia="13/03/2025"),
            ]
        )
    )
    filas = _leer(_procesar(procesador, entrada)[nombre_de_salida("IZIPAY")])["COBRANZA"]
    assert [f[6] for f in filas if f[0] == 1] == ["20250203", "20250313"]


# --- EFECTIVO ------------------------------------------------------------------


def test_efectivo_sale_sin_comision(procesador: AsientosContables, fabrica: _Fabrica) -> None:
    entrada = fabrica.entrada(
        _hojas_efectivo(
            [
                _compensacion("DZ", _d(7, 1), 11215),
                _compensacion("RG", _d(7, 31), -34, "AJUSTE"),
            ]
        ),
        nombre="7. EFECTIVO - JULIO.xlsx",
    )
    salidas = _procesar(procesador, entrada)
    assert list(salidas) == ["resultado_efectivo_procesado.xlsx"]
    hojas = _leer(salidas["resultado_efectivo_procesado.xlsx"])
    assert list(hojas) == ["COBRANZA", "RECARGA"]
    assert _detalle(hojas["COBRANZA"][2]) == (2, 11, 1001415, 11215, "Z001", "EFECTIVO")
    assert _detalle(hojas["RECARGA"][1]) == (2, 11, 1001415, 34, "Z001", "EFECTIVO")


# --- Detección del tipo ------------------------------------------------------


def _error_de_validar(procesador: AsientosContables, *entradas: ArchivoEntrada) -> ErrorContenido:
    error = procesador.validar(list(entradas))
    assert isinstance(error, ErrorContenido)
    assert error.tipo is TipoError.CONTENIDO
    return error


def test_la_hoja_de_la_tarjeta_manda_sobre_el_nombre(
    procesador: AsientosContables, fabrica: _Fabrica
) -> None:
    hojas = _hojas_estandar("VISA", comision=[_comision(_d(6, 1), -1, "Comisión")])
    salidas = _procesar(procesador, fabrica.entrada(hojas, nombre="EFECTIVO de AMEX.xlsx"))
    assert list(salidas) == [nombre_de_salida("VISA")]


@pytest.mark.parametrize(
    ("hojas", "nombre"),
    [
        (_hojas_efectivo([_compensacion("DZ", _d(7, 1), 1)]), "sin pistas.xlsx"),
        (
            {"COMPENSACION": [_ENCABEZADO_COMPENSACION, _compensacion("DZ", _d(7, 1), 1)]},
            "Caja Efectivo julio.xlsx",
        ),
    ],
    ids=["por_hoja_efectivo", "por_nombre_con_compensacion"],
)
def test_efectivo_se_reconoce_por_hoja_o_por_nombre(
    procesador: AsientosContables, fabrica: _Fabrica, hojas: Hojas, nombre: str
) -> None:
    salidas = _procesar(procesador, fabrica.entrada(hojas, nombre=nombre))
    assert list(salidas) == ["resultado_efectivo_procesado.xlsx"]


def test_dos_hojas_de_tarjeta_es_tipo_no_reconocido(
    procesador: AsientosContables, fabrica: _Fabrica
) -> None:
    hojas = {**_hojas_estandar("VISA"), "AMEX": [_ENCABEZADO_COMISION]}
    error = _error_de_validar(procesador, fabrica.entrada(hojas, nombre="AMEX.xlsx"))
    assert error.contexto == {
        "archivo": "AMEX.xlsx",
        "motivo": "tipo_no_reconocido",
        "columna": None,
    }


def test_sin_hoja_ni_nombre_reconocible_es_tipo_no_reconocido(
    procesador: AsientosContables, fabrica: _Fabrica
) -> None:
    hojas = {"COMPENSACION": [_ENCABEZADO_COMPENSACION]}
    error = _error_de_validar(procesador, fabrica.entrada(hojas, nombre="reporte.xlsx"))
    assert error.contexto["motivo"] == "tipo_no_reconocido"


def test_un_nombre_con_dos_tipos_es_tipo_no_reconocido(
    procesador: AsientosContables, fabrica: _Fabrica
) -> None:
    hojas = {"Hoja1": [["x"]]}
    error = _error_de_validar(procesador, fabrica.entrada(hojas, nombre="VISA y AMEX.xlsx"))
    assert error.contexto["motivo"] == "tipo_no_reconocido"


def test_el_nombre_sin_su_hoja_es_hoja_faltante(
    procesador: AsientosContables, fabrica: _Fabrica
) -> None:
    """El respaldo por nombre existe para esto: decir QUÉ hoja falta."""
    hojas = {"COMPENSACION": [_ENCABEZADO_COMPENSACION]}
    error = _error_de_validar(procesador, fabrica.entrada(hojas, nombre="1. amex - agosto.xlsx"))
    assert error.contexto == {
        "archivo": "1. amex - agosto.xlsx",
        "motivo": "hoja_faltante",
        "columna": "AMEX",
    }


def test_falta_compensacion_es_hoja_faltante(
    procesador: AsientosContables, fabrica: _Fabrica
) -> None:
    hojas = {"IZIPAY": [_ENCABEZADO_COMISION]}
    error = _error_de_validar(procesador, fabrica.entrada(hojas, nombre="izipay.xlsx"))
    assert error.contexto == {
        "archivo": "izipay.xlsx",
        "motivo": "hoja_faltante",
        "columna": "COMPENSACION",
    }


def test_dos_archivos_del_mismo_tipo_es_tipo_duplicado(
    procesador: AsientosContables, fabrica: _Fabrica
) -> None:
    """Sus salidas se llamarían igual. El error nombra al SEGUNDO."""
    hojas = _hojas_estandar("VISA", comision=[_comision(_d(6, 1), -1, "Comisión")])
    primero = fabrica.entrada(hojas, nombre="visa junio.xlsx")
    segundo = fabrica.entrada(hojas, nombre="visa julio.xlsx")
    error = _error_de_validar(procesador, primero, segundo)
    assert error.contexto == {
        "archivo": "visa julio.xlsx",
        "motivo": "tipo_duplicado",
        "columna": None,
    }
    with pytest.raises(ErrorContenido) as capturado:
        procesador.procesar([primero, segundo])
    assert capturado.value.contexto["motivo"] == "tipo_duplicado"


# --- Columnas por nombre -----------------------------------------------------


def test_falta_una_columna_obligatoria(procesador: AsientosContables, fabrica: _Fabrica) -> None:
    hojas = _hojas_estandar("VISA")
    hojas["VISA"] = [[c for c in _ENCABEZADO_COMISION if c != "Monto"]]
    error = _error_de_validar(procesador, fabrica.entrada(hojas, nombre="visa.xlsx"))
    assert error.contexto == {
        "archivo": "visa.xlsx",
        "motivo": "columna_faltante",
        "columna": "Monto",
    }


def test_la_fecha_de_cierre_es_obligatoria(
    procesador: AsientosContables, fabrica: _Fabrica
) -> None:
    """Decisión de la migración: un encabezado renombrado no cae a hoy en silencio."""
    hojas = _hojas_estandar("VISA")
    hojas["COMPENSACION"] = [_ENCABEZADO_COMPENSACION[:6]]
    error = _error_de_validar(procesador, fabrica.entrada(hojas))
    assert error.contexto == {
        "archivo": "archivo.xlsx",
        "motivo": "columna_faltante",
        "columna": "Fecha Cierre",
    }


def test_los_alias_de_encabezado_se_aceptan(
    procesador: AsientosContables, fabrica: _Fabrica
) -> None:
    """VISA real: la sexta columna se llama `Ajuste`; mayúsculas y espacios no importan."""
    encabezado = [*_ENCABEZADO_COMPENSACION[:5], "Ajuste", "  fecha CIERRE "]
    hojas = _hojas_estandar("VISA")
    hojas["COMPENSACION"] = [encabezado, _compensacion("RG", _d(6, 5), 30, "AJUSTE")]
    filas = _leer(_procesar(procesador, fabrica.entrada(hojas))[nombre_de_salida("VISA")])
    assert list(filas) == ["RECARGA"]


def test_una_columna_insertada_no_corre_la_lectura(
    procesador: AsientosContables, fabrica: _Fabrica
) -> None:
    """Leyendo por posición —lo que hacía el script— esto salía mal sin fallar."""
    hojas = _hojas_estandar("VISA")
    hojas["VISA"] = [
        ["Nueva", *_ENCABEZADO_COMISION],
        ["basura", *_comision(_d(6, 1), -4, "Comisión")],
    ]
    filas = _leer(_procesar(procesador, fabrica.entrada(hojas))[nombre_de_salida("VISA")])
    assert _detalle(filas["COMISION"][1])[3] == 4


# --- Cero filas y archivos ilegibles -----------------------------------------


def test_cero_filas_es_error_y_no_deja_archivos(
    procesador: AsientosContables, fabrica: _Fabrica, tmp_path: Path
) -> None:
    hojas = _hojas_estandar(
        "VISA",
        comision=[_comision(_d(6, 1), 100, "Ventas según SAP")],
        compensacion=[_compensacion("RG", _d(6, 1), 10, None)],
    )
    entrada = fabrica.entrada(hojas, nombre="visa.xlsx")
    assert procesador.validar([entrada]) is None
    with pytest.raises(ErrorContenido) as capturado:
        procesador.procesar([entrada])
    assert capturado.value.contexto == {
        "archivo": "visa.xlsx",
        "motivo": "cero_filas",
        "columna": None,
    }
    assert sorted(p.name for p in tmp_path.iterdir()) == ["entrada_0"]


def test_un_segundo_archivo_sin_filas_hace_fallar_todo(
    procesador: AsientosContables, fabrica: _Fabrica, tmp_path: Path
) -> None:
    """No hay éxito parcial: ni siquiera el Excel del primer archivo queda escrito."""
    bueno = fabrica.entrada(_hojas_estandar("VISA", comision=[_comision(_d(6, 1), -1, "Comisión")]))
    vacio = fabrica.entrada(_hojas_estandar("AMEX"), nombre="amex.xlsx")
    with pytest.raises(ErrorContenido) as capturado:
        procesador.procesar([bueno, vacio])
    assert capturado.value.contexto == {
        "archivo": "amex.xlsx",
        "motivo": "cero_filas",
        "columna": None,
    }
    assert sorted(p.name for p in tmp_path.iterdir()) == ["entrada_0", "entrada_1"]


def test_un_excel_ilegible_es_error_de_contenido(
    procesador: AsientosContables, tmp_path: Path
) -> None:
    ruta = tmp_path / "entrada_0"
    ruta.write_bytes(b"esto no es un xlsx")
    entrada = ArchivoEntrada(
        nombre_original="corrupto.xlsx", ruta_temporal=ruta, tamano_comprimido=18, formato="xlsx"
    )
    error = _error_de_validar(procesador, entrada)
    assert error.contexto["motivo"] == "cero_filas"


# --- Observaciones -----------------------------------------------------------


def test_entrada_limpia_sin_observaciones(procesador: AsientosContables, fabrica: _Fabrica) -> None:
    hojas = _hojas_estandar("VISA", comision=[_comision(_d(6, 1), -1, "Comisión")])
    assert list(_procesar(procesador, fabrica.entrada(hojas))) == [nombre_de_salida("VISA")]


def test_las_perdidas_silenciosas_se_reportan(
    procesador: AsientosContables, fabrica: _Fabrica
) -> None:
    """Fecha ilegible en una fila filtrada, y un cierre ilegible que cae a hoy."""
    hojas = _hojas_estandar(
        "VISA",
        compensacion=[
            _compensacion("DZ", _d(6, 1), 10),
            _compensacion("DZ", "no es una fecha", 20),
            _compensacion("DZ", _d(6, 2), 30, cierre="tampoco"),
            _compensacion("RG", None, 40, "AJUSTE"),
        ],
    )
    salidas = _procesar(procesador, fabrica.entrada(hojas, nombre="visa.xlsx"))

    cobranza = _leer(salidas[nombre_de_salida("VISA")])["COBRANZA"]
    assert [f[6] for f in cobranza if f[0] == 1] == ["20250601", "20250602"]
    assert cobranza[4][7] == _CIERRE_HOY

    crudo = salidas[NOMBRE_OBSERVACIONES].read_bytes()
    assert crudo.endswith(b"\r\n")
    assert b"\n" not in crudo.replace(b"\r\n", b"")
    lineas = crudo.decode("utf-8").split(FIN_DE_LINEA)
    assert lineas[0] == (f"Observaciones: 3 ({MOTIVO_FECHA_CIERRE}=1, {MOTIVO_FECHA}=2)")
    assert lineas[2] == "archivo\thoja\tfila\tmotivo\tvalor"
    assert lineas[3:6] == [
        f"visa.xlsx\tCOMPENSACION\t3\t{MOTIVO_FECHA}\tno es una fecha",
        f"visa.xlsx\tCOMPENSACION\t4\t{MOTIVO_FECHA_CIERRE}\ttampoco",
        f"visa.xlsx\tCOMPENSACION\t5\t{MOTIVO_FECHA}\t",
    ]


def test_las_observaciones_se_registran_en_el_log(
    procesador: AsientosContables, fabrica: _Fabrica, caplog: pytest.LogCaptureFixture
) -> None:
    hojas = _hojas_estandar("VISA", compensacion=[_compensacion("DZ", "ilegible", 1)])
    hojas["VISA"].append(_comision(_d(6, 1), -1, "Comisión"))
    with caplog.at_level("INFO", logger="app"):
        _procesar(procesador, fabrica.entrada(hojas, nombre="visa.xlsx"))

    (registro,) = [r for r in caplog.records if getattr(r, "evento", None) == "descartes"]
    assert registro.clave == CLAVE  # type: ignore[attr-defined]
    assert registro.archivo == "visa.xlsx"  # type: ignore[attr-defined]
    assert registro.por_motivo == {MOTIVO_FECHA: 1}  # type: ignore[attr-defined]
    assert registro.filas_procesadas == 3  # type: ignore[attr-defined]


def test_sin_observaciones_no_se_registra_nada(
    procesador: AsientosContables, fabrica: _Fabrica, caplog: pytest.LogCaptureFixture
) -> None:
    hojas = _hojas_estandar("VISA", comision=[_comision(_d(6, 1), -1, "Comisión")])
    with caplog.at_level("INFO", logger="app"):
        _procesar(procesador, fabrica.entrada(hojas))
    assert not [r for r in caplog.records if getattr(r, "evento", None) == "descartes"]


# --- Varios archivos por ejecución ---------------------------------------------


def test_un_excel_por_archivo_con_el_nombre_del_script(
    procesador: AsientosContables, fabrica: _Fabrica
) -> None:
    entradas = [
        fabrica.entrada(_hojas_estandar(tipo, comision=[_comision(_d(6, 1), -1, "Comisión")]))
        for tipo in ("AMEX", "DINNERS")
    ]
    salidas = _procesar(procesador, *entradas)
    assert list(salidas) == ["resultado_amex_procesado.xlsx", "resultado_dinners_procesado.xlsx"]


# --- Ninguna excepción se traga ------------------------------------------------


def test_una_fila_que_rompe_no_se_traga(
    procesador: AsientosContables, fabrica: _Fabrica, monkeypatch: pytest.MonkeyPatch
) -> None:
    """El script tenía `try/except: print` por hoja. Acá se propaga."""
    import app.procesadores.asientos_contables.modulo as modulo

    def _explota(*_args: Any, **_kwargs: Any) -> float:
        raise RuntimeError("importe imposible")

    monkeypatch.setattr(modulo, "_a_importe", _explota)
    hojas = _hojas_estandar("VISA", comision=[_comision(_d(6, 1), -1, "Comisión")])
    with pytest.raises(RuntimeError, match="importe imposible"):
        procesador.procesar([fabrica.entrada(hojas)])


# --- Conversiones que emulan a pandas ------------------------------------------


@pytest.mark.parametrize(
    ("crudo", "esperado"),
    [
        ("1,234.56", 1234.56),
        ("  -42.5 ", -42.5),
        ("                       -  ", 0),
        ("", 0),
        ("nan", 0),
        ("NA", 0),
        ("no es un número", 0),
        ("1_000", 0),
        (None, 0),
        (True, 0),
        (1250.5, 1250.5),
        (200.0, 200),
        (-13, -13),
    ],
)
def test_limpieza_de_importes(crudo: object, esperado: float) -> None:
    """Equivalente de `limpiar_montos`: lo que no parsea vale cero."""
    assert _a_importe(crudo) == esperado


@pytest.mark.parametrize(
    ("crudo", "dia_primero", "esperado"),
    [
        ("02/03/2025", True, datetime(2025, 3, 2)),
        ("02/03/2025", False, datetime(2025, 2, 3)),
        ("13/03/2025", False, datetime(2025, 3, 13)),
        ("2025-03-02", True, datetime(2025, 3, 2)),
        (date(2025, 3, 2), True, datetime(2025, 3, 2)),
        (datetime(2025, 3, 2, 10, 30), True, datetime(2025, 3, 2, 10, 30)),
        ("IZIPAY 01.04.25", True, None),
        ("NaT", True, None),
        (45000, True, None),
        (None, True, None),
    ],
)
def test_lectura_de_fechas(crudo: object, dia_primero: bool, esperado: datetime | None) -> None:
    assert _a_fecha(crudo, dia_primero=dia_primero) == esperado


@pytest.mark.parametrize(
    ("valores", "groupby", "serie"),
    [
        ([0.1] * 10, 1.0, 1.0),
        ([0.1] * 20, 2.0, 2.0000000000000004),
        ([0.1] * 300, 30.0, 29.999999999999996),
        ([0.7] * 130, 91.0, 91.00000000000001),
    ],
)
def test_las_sumas_reproducen_las_de_pandas(
    valores: list[float], groupby: float, serie: float
) -> None:
    """Valores esperados obtenidos de pandas 3.0.6, no calculados a mano.

    Un bucle ingenuo da 30.000000000000156 para `[0.1] * 300`: `groupby().sum()`
    compensa (Kahan) y `Series.sum()` suma por pares, y los tres difieren.
    """
    assert _suma_de_grupo(valores) == groupby
    assert _suma_de_serie(valores) == serie
    assert _suma_de_grupo([1, 2, 3]) == 6 and isinstance(_suma_de_grupo([1, 2, 3]), int)


def test_los_nuevos_motivos_conservan_la_forma_del_contexto() -> None:
    """`hoja_faltante` reutiliza `columna`; los otros dos la dejan en `None`."""
    assert ErrorContenido.hoja_faltante(archivo="a.xlsx", hoja="COMPENSACION").contexto == {
        "archivo": "a.xlsx",
        "motivo": "hoja_faltante",
        "columna": "COMPENSACION",
    }
    for error, motivo in (
        (ErrorContenido.tipo_no_reconocido(archivo="a.xlsx"), "tipo_no_reconocido"),
        (ErrorContenido.tipo_duplicado(archivo="a.xlsx"), "tipo_duplicado"),
    ):
        assert error.contexto == {"archivo": "a.xlsx", "motivo": motivo, "columna": None}


# --- Coherencia con el registro ----------------------------------------------


def test_la_clave_coincide_con_el_registry() -> None:
    from app.core.contrato import obtener_contrato
    from app.registry import REGISTRY

    procesador = REGISTRY[CLAVE]
    assert procesador.clave == CLAVE == "asientos-contables"
    assert isinstance(procesador, AsientosContables)

    contrato = obtener_contrato(CLAVE)
    assert contrato is not None
    assert (contrato.entradas_min, contrato.entradas_max) == (1, 7)
    assert contrato.formatos_aceptados == ("xlsx",)


# --- De punta a punta, por la ruta real y el proceso hijo -------------------


class TestPuntaAPunta:
    """La tubería completa: HTTP → temporales sin extensión → `spawn` → ZIP.

    La lección de `contado_carga`: todo procesador nuevo necesita una prueba
    que suba por HTTP y cruce el hijo de `spawn`, porque la trampa del
    temporal sin extensión sólo se ve ahí.
    """

    @staticmethod
    def _subir(cliente: Any, archivos: list[tuple[str, bytes]]) -> Any:
        return cliente.post(
            f"/interno/procesadores/{CLAVE}",
            files=[
                ("archivos", (nombre, contenido, "application/octet-stream"))
                for nombre, contenido in archivos
            ],
            headers={"Authorization": "Bearer sentinela-token-de-pruebas-3f9c2a"},
        )

    def test_dos_excel_vuelven_como_zip_con_dos_resultados(self, token_sentinela: str) -> None:
        from fastapi.testclient import TestClient

        from app.main import crear_app

        visa = _bytes_de_libro(
            _hojas_estandar("VISA", comision=[_comision(_d(6, 1), -1, "Comisión")])
        )
        efectivo = _bytes_de_libro(_hojas_efectivo([_compensacion("DZ", _d(7, 1), 11215)]))
        with TestClient(crear_app()) as cliente:
            respuesta = self._subir(
                cliente, [("5. VISA.xlsx", visa), ("7. EFECTIVO - JULIO.xlsx", efectivo)]
            )

        assert respuesta.status_code == 200, respuesta.text
        with zipfile.ZipFile(io.BytesIO(respuesta.content)) as paquete:
            assert sorted(paquete.namelist()) == [
                "resultado_efectivo_procesado.xlsx",
                "resultado_visa_procesado.xlsx",
            ]
            libro = load_workbook(io.BytesIO(paquete.read("resultado_visa_procesado.xlsx")))
            assert libro.sheetnames == ["COMISION"]

    def test_un_tipo_duplicado_es_422(self, token_sentinela: str) -> None:
        """El error tipificado cruza el hijo intacto y llega como 422 de contenido."""
        from fastapi.testclient import TestClient

        from app.main import crear_app

        visa = _bytes_de_libro(
            _hojas_estandar("VISA", comision=[_comision(_d(6, 1), -1, "Comisión")])
        )
        with TestClient(crear_app()) as cliente:
            respuesta = self._subir(cliente, [("a.xlsx", visa), ("b.xlsx", visa)])

        assert respuesta.status_code == 422
        cuerpo = respuesta.json()
        assert cuerpo["tipo"] == "contenido"
        assert cuerpo["contexto"] == {
            "archivo": "b.xlsx",
            "motivo": "tipo_duplicado",
            "columna": None,
        }
