"""Fábricas de entradas sintéticas para los procesadores de medios de pago.

Los reportes reales no están en este repositorio y no pueden estarlo (ADR
0015): estas fábricas arman CSV de Izipay con los títulos EXACTOS de las
muestras de mayo de 2025 (con la coma final que deja una columna `Unnamed`),
libros de SafetyPay con títulos en la fila 7, extractos BBVA con registros
`00`, `11`, `22` y `23`, y maestros de abonos. Todo se guarda con nombres SIN
extensión (`entrada_0`) como los que escribe `app/recepcion.py`.
"""

from __future__ import annotations

import io
from collections.abc import Mapping, Sequence
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook

from app.core.tipos import ArchivoEntrada

TITULOS_IZIPAY: dict[str, tuple[str, ...]] = {
    "mastercard": tuple(
        "Codigo,Producto,Tipo_Mov,Fecha_Proceso,Fecha_Lote,Lote_Manual,Lote_Pos,Terminal,Voucher,"
        "Autorizacion,Cuotas,Tarjeta,Origen,Transaccion,Fecha_Consumo,Importe,Status,Comision,"
        "Comision_Afecta,IGV,Neto_Parcial,Neto_Total,Fecha_Abono,Fecha_Abono_8Dig,Observaciones".split(
            ","
        )
    ),
    "amex": tuple(
        "Codigo,Producto,Tipo_Mov,Fecha_Proceso,Fecha_Lote,Lote_Manual,Lote_Pos,Terminal,Voucher,"
        "Autorizacion,Cuotas,Tarjeta,Origen,Transaccion,Fecha_Consumo,Importe,Comision,"
        "Comision_Merchant,IGV,Neto_Parcial,Neto_Total,Fecha_Abono,Observaciones".split(",")
    ),
    "dinner": tuple(
        "Codigo,Producto,Tipo_Mov,Fecha_Proceso,Fecha_Lote,Lote_Manual,Lote_Pos,Terminal,Voucher,"
        "Autorizacion,Cuotas,Tarjeta,Origen,Transaccion,Fecha_Consumo,Importe,Status,Comision,IGV,"
        "Neto_Parcial,Neto_Total,Fecha_Abono,Fecha_Abono_8Dig,Observaciones".split(",")
    ),
}

ANTIGUO = "004010761"
IZIPAY = "001042409"


def fila_izipay(
    tipo: str, abono: str, *, codigo: str = IZIPAY, neto: str = "13.87", **otros: str
) -> dict[str, str]:
    """Una fila con el aspecto de las reales; `abono` es `AAAAMMDD`."""
    fila = {
        "Codigo": codigo,
        "Producto": "VA",
        "Tipo_Mov": "A",
        "Fecha_Proceso": "02/05/2025",
        "Lote_Pos": "0378",
        "Terminal": "PS0005",
        "Voucher": "0004034",
        "Autorizacion": "072277",
        "Tarjeta": "477289******7763",
        "Importe": "14.00",
        "Comision": "0.12",
        "IGV": "0.01",
        "Neto_Parcial": "13.87",
        "Neto_Total": neto,
    }
    fila["Fecha_Abono" if tipo == "amex" else "Fecha_Abono_8Dig"] = abono
    fila.update(otros)
    return fila


def csv_izipay(
    tipo: str,
    filas: Sequence[Mapping[str, str]],
    *,
    titulos: Sequence[str] | None = None,
    coma_final: bool = True,
    extra: Sequence[str] = (),
) -> bytes:
    """El CSV como lo descarga Izipay: UTF-8, LF y, como en las muestras, una
    coma al final de cada línea. `extra` son líneas crudas agregadas al final."""
    columnas = list(titulos or TITULOS_IZIPAY[tipo])
    coma = "," if coma_final else ""
    lineas = [",".join(columnas) + coma]
    for fila in filas:
        lineas.append(",".join(_campo(fila.get(c, "")) for c in columnas) + coma)
    lineas.extend(extra)
    return ("\n".join(lineas) + "\n").encode("utf-8")


def _campo(valor: str) -> str:
    """Comillas sólo donde hacen falta, como las pone cualquier exportador."""
    if "," in valor or '"' in valor:
        return '"' + valor.replace('"', '""') + '"'
    return valor


TITULOS_SAFETYPAY: tuple[str, ...] = (
    "Purchase Complete Date (GMT Merchant)",
    "Operation ID",
    "Transaction ID",
    "Type Of Operation",
    "Merchant Sales ID",
    "Merchant Order Number",
    "Sale Currency",
    "Sale Amount",
    "SafetyPay Commission ",
    "Net To Merchant",
    "Reconciled",
    "Merchant Settlement Date",
    "Payment Batch Number",
    "Created By",
)


def fecha_safetypay(dia: date) -> str:
    """`"Monday, 26 May 2025 00:00:00"`, la forma de las muestras."""
    dias = ("Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday")
    meses = (
        "January",
        "February",
        "March",
        "April",
        "May",
        "June",
        "July",
        "August",
        "September",
        "October",
        "November",
        "December",
    )
    return f"{dias[dia.weekday()]}, {dia.day:02d} {meses[dia.month - 1]} {dia.year} 00:00:00"


def fila_safetypay(liquidacion: object, **otros: object) -> dict[str, object]:
    fila: dict[str, object] = {
        "Purchase Complete Date (GMT Merchant)": "05/22/2025 22:48:25",
        "Operation ID": "0125143389223723",
        "Transaction ID": "389223",
        "Type Of Operation": "Sales",
        "Merchant Sales ID": "00000000000000390223",
        "Merchant Order Number": "00000000000000390223",
        "Sale Currency": "PEN",
        "Sale Amount": "20.00",
        "SafetyPay Commission ": "0.30",
        "Net To Merchant": "19.65",
        "Reconciled": "Y  ",
        "Merchant Settlement Date": liquidacion,
        "Payment Batch Number": "177614",
    }
    fila.update(otros)
    return fila


def libro_safetypay(
    filas: Sequence[Mapping[str, object]],
    *,
    pie: bool = True,
    titulos: Sequence[str] = TITULOS_SAFETYPAY,
) -> bytes:
    """Un `MPFinancialReport`: seis filas de cabecera, títulos en la séptima y,
    como en las muestras, una fila final `Total records:` sin fecha."""
    libro = Workbook()
    hoja = libro.active
    assert hoja is not None
    hoja.title = "Financial Report"
    for _ in range(3):
        hoja.append([None])
    hoja.append(["Merchant name:", "PEX"])
    hoja.append(["Total records:", str(len(filas))])
    hoja.append([None])
    hoja.append(list(titulos))
    for fila in filas:
        hoja.append([fila.get(c) for c in titulos])
    if pie:
        hoja.append(["Total records:", str(len(filas))])
    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()


def linea_22(
    fecha: str,
    importe_centimos: int,
    *,
    codigo: int = 118,
    concepto: str = "R20432405525PROCESOS DE MEDI",
    doc: str = "0814",
    oficina: str = "0000026498",
) -> str:
    """Un registro de movimiento del extracto; `fecha` es `aammdd`."""
    return (
        f"22,0011,{doc},{fecha},{fecha},00,{codigo:03d},2,"
        f"{importe_centimos:014d},{oficina},{concepto}"
    )


def extracto_bbva(lineas: Sequence[str]) -> bytes:
    """El `bbva.csv_`: cabecera `00` y `11`, y un `23` después de cada `22`,
    en latin1 como lo entrega el banco."""
    todas = [
        "00,0081,250528,M- M- M-",
        "11,0011,0397,0100004984,250502,250502,2,0,211,3,PEX PERU SAC,",
    ]
    for linea in lineas:
        todas.append(linea)
        if linea.startswith("22"):
            todas.append("23,01,BCA. AUTOMATICA      MOV.0000026498,")
    return ("\n".join(todas) + "\n").encode("latin1")


TITULOS_MAESTRO: tuple[str, ...] = (
    "F. Operación",
    "F. Valor",
    "Código",
    "Importe",
    "Oficina",
    "Concepto",
    "Nº. Doc.",
    "Descripción",
    "Observacion",
    "FECHA (VAN Y VIENE DE ABONO)",
)


def libro_maestro(filas: Sequence[Sequence[object]]) -> bytes:
    """Un `Abonos_BBVA.xlsx`: un CATEGORIZADO de una corrida anterior."""
    libro = Workbook()
    hoja = libro.active
    assert hoja is not None
    hoja.append(list(TITULOS_MAESTRO))
    for fila in filas:
        hoja.append(list(fila))
    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()


@dataclass
class Fabrica:
    """Escribe entradas como `entrada_N`, sin extensión."""

    directorio: Path
    contador: int = 0
    creadas: list[ArchivoEntrada] = field(default_factory=list)

    def __call__(self, nombre: str, contenido: bytes) -> ArchivoEntrada:
        ruta = self.directorio / f"entrada_{self.contador}"
        self.contador += 1
        ruta.write_bytes(contenido)
        _nombre, punto, extension = nombre.rpartition(".")
        entrada = ArchivoEntrada(
            nombre_original=nombre,
            ruta_temporal=ruta,
            tamano_comprimido=len(contenido),
            formato=extension.lower() if punto else "",
        )
        self.creadas.append(entrada)
        return entrada


def leer_xlsx(ruta: Path | bytes) -> list[list[Any]]:
    """Las filas de la primera hoja, con valores (por descriptor si es ruta)."""
    fuente: Any = io.BytesIO(ruta) if isinstance(ruta, bytes) else ruta.open("rb")
    try:
        libro = load_workbook(fuente)
        hoja = libro.worksheets[0]
        return [list(f) for f in hoja.iter_rows(values_only=True)]
    finally:
        if not isinstance(ruta, bytes):
            fuente.close()


def celdas_xlsx(ruta: Path) -> list[list[Any]]:
    """Las celdas (no los valores) de la primera hoja: tipo y formato."""
    with ruta.open("rb") as descriptor:
        libro = load_workbook(descriptor)
        return [list(f) for f in libro.worksheets[0].iter_rows()]


def como_fecha(texto: str) -> datetime:
    return datetime.strptime(texto, "%d/%m/%Y")
