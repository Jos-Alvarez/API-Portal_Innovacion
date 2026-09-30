"""Paridad de `Prepago_Carga` contra el extracto real (ADR 0015).

**Estado: 1 par de los 3 que exige el PRD.** El par 01 es el extracto de
transferencias entre cuentas de febrero de 2026 con las salidas HISTÓRICAS
que el script produjo el 2026-03-20, no regeneradas. Con `FIXTURES_PARIDAD`
sin definir todo esto saltea con aviso, y falla en CI.

**El TXT se compara byte a byte y el Excel por contenido normalizado**, tal
como fija ADR 0015 y como hace `test_contado_carga.py`.
"""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import load_workbook

from app.core.tipos import ArchivoEntrada
from app.procesadores.prepago_carga.modulo import (
    NOMBRE_DESCARTES,
    NOMBRE_TXT,
    NOMBRE_XLSX,
    PrepagoCarga,
)
from tests.paridad.manifiesto import (
    VARIABLE_DE_UBICACION,
    ParDeParidad,
    exigir_pares,
    resolver_raiz,
)

PROCESADOR = "prepago_carga"


def _pares() -> list[ParDeParidad]:
    """Compuerta de ADR 0015. Corta la prueba —fallando o salteando— si falta algo."""
    return exigir_pares(PROCESADOR)


def _directorio() -> Path:
    raiz = resolver_raiz()
    assert raiz is not None, f"{VARIABLE_DE_UBICACION} sin definir pese a haber pasado la compuerta"
    return raiz / PROCESADOR


def _ejecutar(par: ParDeParidad, destino: Path) -> dict[str, Path]:
    """Corre el módulo sobre la entrada del par, en un directorio de trabajo propio.

    La copia se llama `entrada_0`, sin extensión, como el temporal real.
    """
    origen = _directorio() / par.archivos["entrada"]
    trabajo = destino / par.id
    trabajo.mkdir(parents=True)
    copia = trabajo / "entrada_0"
    copia.write_bytes(origen.read_bytes())

    entrada = ArchivoEntrada(
        nombre_original=origen.name,
        ruta_temporal=copia,
        tamano_comprimido=copia.stat().st_size,
        formato="xlsx",
    )
    return {s.nombre_propuesto: s.ruta_temporal for s in PrepagoCarga().procesar([entrada])}


def _celdas(ruta: Path) -> list[list[object]]:
    libro = load_workbook(ruta, data_only=True)
    try:
        return [list(fila) for fila in libro.worksheets[0].iter_rows(values_only=True)]
    finally:
        libro.close()


def test_el_txt_coincide_byte_a_byte(tmp_path: Path) -> None:
    """Codificación y fin de línea incluidos. Es el punto entero de ADR 0015."""
    for par in _pares():
        salidas = _ejecutar(par, tmp_path)
        esperado = (_directorio() / par.archivos["salida_txt"]).read_bytes()
        obtenido = salidas[NOMBRE_TXT].read_bytes()
        assert obtenido == esperado, (
            f"par {par.id}: el TXT no coincide byte a byte. Procedencia del esperado: "
            f"{par.procedencia.generado_por} el {par.procedencia.generado_el} en "
            f"{par.procedencia.generado_en}, fin de línea {par.procedencia.fin_de_linea}."
        )


def test_el_excel_coincide_por_contenido(tmp_path: Path) -> None:
    """Celda por celda y con tipos: `"02"` contra `2` es una diferencia."""
    for par in _pares():
        salidas = _ejecutar(par, tmp_path)
        esperado = _celdas(_directorio() / par.archivos["salida_xlsx"])
        obtenido = _celdas(salidas[NOMBRE_XLSX])
        assert obtenido == esperado, f"par {par.id}: el Excel no coincide"
        assert [[type(c) for c in f] for f in obtenido] == [
            [type(c) for c in f] for f in esperado
        ], f"par {par.id}: el Excel coincide en valor pero no en tipos"


def test_ningun_par_real_tiene_descartes(tmp_path: Path) -> None:
    """Un par con descartes codifica el comportamiento que la migración corrige."""
    for par in _pares():
        salidas = _ejecutar(par, tmp_path)
        assert NOMBRE_DESCARTES not in salidas, (
            f"par {par.id}: el módulo endurecido descarta filas que el script viejo perdía en "
            f"silencio. Revisar {salidas.get(NOMBRE_DESCARTES)} y decidir explícitamente si "
            "manda la paridad byte a byte o el endurecimiento obligatorio."
        )


@pytest.mark.parametrize("papel", ["entrada", "salida_xlsx", "salida_txt"])
def test_cada_par_declara_los_tres_archivos(papel: str) -> None:
    """El manifiesto es el contrato: un par a medias no es un par."""
    for par in _pares():
        assert papel in par.archivos, f"par {par.id}: falta '{papel}' en el manifiesto"
