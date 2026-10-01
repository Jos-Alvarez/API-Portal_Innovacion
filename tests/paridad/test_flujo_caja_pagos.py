"""Paridad de `flujo-caja-pagos` contra la salida del script (ADR 0015).

**Estado: 1 par con tres entradas y una salida.** Las entradas son tres libros
semanales reales de pagos programados LIMA EXPRESA (julio de 2025) y la
salida esperada es `Consolidado_Pagos.xlsx` tal como la produjo el script
—no regenerada—. El módulo la reproduce celda por celda con igualdad EXACTA,
en las dos hojas, con el orden de filas de la concatenación. Con
`FIXTURES_PARIDAD` sin definir, todo esto saltea con aviso, y falla en CI.

**Los nombres originales importan**: salen en la columna `Excel` y deciden el
orden de los libros (el de NTFS), así que se conservan en `_NOMBRES_ORIGINALES`
y los libros se entregan en orden inverso para que el orden de subida no
pueda disimular un error de orden.

**El Excel se compara por contenido, sin tolerancia**, con la misma única
excepción que las demás paridades: `66` y `66.0` son la misma celda. Las
columnas leídas con `dtype=str` se comparan como TEXTO: texto contra número
es una diferencia real.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from app.core.tipos import ArchivoEntrada
from app.procesadores.flujo_caja_pagos.modulo import (
    NOMBRE_CONSOLIDADO,
    NOMBRE_OBSERVACIONES,
    FlujoCajaPagos,
)
from tests.paridad.manifiesto import (
    VARIABLE_DE_UBICACION,
    ParDeParidad,
    exigir_pares,
    resolver_raiz,
)

PROCESADOR = "flujo_caja_pagos"

_NOMBRES_ORIGINALES: dict[str, str] = {
    "entrada_1": "01.-Pagos Programados LIMA EXPRESA del 01 al 06.07.xlsm",
    "entrada_2": "02.-Pagos Programados LIMA EXPRESA del 07 al 13.07.xlsm",
    "entrada_3": "03.-Pagos Programados LIMA EXPRESA del 14 al 20.07.xlsm",
}
_SALIDA: str = "salida_consolidado"


def _pares() -> list[ParDeParidad]:
    """Compuerta de ADR 0015. Corta la prueba —fallando o salteando— si falta algo."""
    return exigir_pares(PROCESADOR)


def _directorio() -> Path:
    raiz = resolver_raiz()
    assert raiz is not None, f"{VARIABLE_DE_UBICACION} sin definir pese a haber pasado la compuerta"
    return raiz / PROCESADOR


def _ejecutar(par: ParDeParidad, destino: Path) -> dict[str, Path]:
    """Corre el módulo sobre las entradas del par, como `entrada_N` sin extensión.

    Se llama a `procesar` directo: lo que la paridad verifica son las reglas de
    negocio, no la plomería del proceso hijo.
    """
    trabajo = destino / par.id
    trabajo.mkdir(parents=True)
    entradas: list[ArchivoEntrada] = []
    for indice, papel in enumerate(reversed(_NOMBRES_ORIGINALES)):
        copia = trabajo / f"entrada_{indice}"
        copia.write_bytes((_directorio() / par.archivos[papel]).read_bytes())
        entradas.append(
            ArchivoEntrada(
                nombre_original=_NOMBRES_ORIGINALES[papel],
                ruta_temporal=copia,
                tamano_comprimido=copia.stat().st_size,
                formato="xlsm",
            )
        )
    procesador = FlujoCajaPagos()
    assert procesador.validar(entradas) is None, f"par {par.id}: las entradas no validan"
    return {s.nombre_propuesto: s.ruta_temporal for s in procesador.procesar(entradas)}


def _celdas(ruta: Path) -> dict[str, list[list[Any]]]:
    libro = load_workbook(ruta, read_only=True, data_only=True)
    try:
        return {h.title: [list(f) for f in h.iter_rows(values_only=True)] for h in libro.worksheets}
    finally:
        libro.close()


def _normalizado(celda: Any) -> Any:
    """`66` y `66.0` son la misma celda (ver la docstring)."""
    if isinstance(celda, float) and celda.is_integer():
        return int(celda)
    return celda


def _tipado(hojas: dict[str, list[list[Any]]]) -> dict[str, list[list[tuple[type, Any]]]]:
    return {
        nombre: [[(type(_normalizado(c)), _normalizado(c)) for c in fila] for fila in filas]
        for nombre, filas in hojas.items()
    }


def test_el_consolidado_coincide_celda_por_celda(tmp_path: Path) -> None:
    for par in _pares():
        salidas = _ejecutar(par, tmp_path)
        assert NOMBRE_CONSOLIDADO in salidas, f"par {par.id}: el módulo no produjo el consolidado"
        obtenido = _celdas(salidas[NOMBRE_CONSOLIDADO])
        esperado = _celdas(_directorio() / par.archivos[_SALIDA])
        assert list(obtenido) == list(esperado), f"par {par.id}: hojas distintas"
        for hoja in esperado:
            assert len(obtenido[hoja]) == len(esperado[hoja]), (
                f"par {par.id}/{hoja}: {len(obtenido[hoja])} filas, "
                f"se esperaban {len(esperado[hoja])}"
            )
        assert _tipado(obtenido) == _tipado(esperado), (
            f"par {par.id}: el Excel no coincide. Procedencia del esperado: "
            f"{par.procedencia.generado_por} el {par.procedencia.generado_el} en "
            f"{par.procedencia.generado_en}."
        )


def test_una_semana_real_no_deja_observaciones(tmp_path: Path) -> None:
    """Las semanas reales traen hojas sin datos (el script decía "no contiene
    datos"), pero eso no se reporta: la salida es el Excel solo. Una
    observación acá significa que la salida esperada codifica una pérdida
    silenciosa que la migración vino a hacer visible, y hay que decidir cuál
    requisito manda."""
    for par in _pares():
        salidas = _ejecutar(par, tmp_path)
        assert NOMBRE_OBSERVACIONES not in salidas, (
            f"par {par.id}: observaciones inesperadas en una entrada real"
        )


def test_cada_par_declara_sus_tres_entradas_y_su_salida() -> None:
    """El manifiesto es el contrato: un par a medias no es un par."""
    for par in _pares():
        assert {*_NOMBRES_ORIGINALES, _SALIDA} <= set(par.archivos), f"par {par.id}: incompleto"
