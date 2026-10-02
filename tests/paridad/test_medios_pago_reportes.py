"""Paridad de `medios-pago-reportes` contra la salida del script (ADR 0015).

**Estado: 1 par con cuatro entradas** (los reportes crudos reales de mayo de
2025: Mastercard, AMEX, Diners y SafetyPay) **y 99 salidas**, guardadas como
un ZIP en el recurso compartido. Las salidas esperadas se REGENERARON
corriendo el script original sobre estas mismas entradas (ver la procedencia
en `manifiesto.toml`). Con `FIXTURES_PARIDAD` sin definir, todo esto saltea con
aviso, y falla en CI.

Los CSV se comparan **byte por byte** (incluido el CRLF de `to_csv` en
Windows). Los libros, celda por celda: valor con su tipo, formato de número y
estilo de los títulos, con la misma única tolerancia que las demás
paridades: `66` y `66.0` son la misma celda.
"""

from __future__ import annotations

import io
import zipfile
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from app.core.tipos import ArchivoEntrada
from app.procesadores.medios_pago.comun import NOMBRE_OBSERVACIONES
from app.procesadores.medios_pago.reportes import MediosPagoReportes
from tests.paridad.manifiesto import (
    VARIABLE_DE_UBICACION,
    ParDeParidad,
    exigir_pares,
    resolver_raiz,
)

PROCESADOR = "medios_pago_reportes"

NOMBRES_ORIGINALES: dict[str, str] = {
    "entrada_mc": "mc_052025009428105 (13).csv",
    "entrada_amex": "movi_amex052025009428105 (13).csv",
    "entrada_dinner": "servicios052025009428105 (13).csv",
    "entrada_safetypay": "MPFinancialReport (1).xlsx",
}
_SALIDA: str = "salida_zip"


def _pares() -> list[ParDeParidad]:
    """Compuerta de ADR 0015. Corta la prueba —fallando o salteando— si falta algo."""
    return exigir_pares(PROCESADOR)


def directorio(procesador: str = PROCESADOR) -> Path:
    raiz = resolver_raiz()
    assert raiz is not None, f"{VARIABLE_DE_UBICACION} sin definir pese a haber pasado la compuerta"
    return raiz / procesador


def preparar(
    par: ParDeParidad, nombres: dict[str, str], destino: Path, procesador: str = PROCESADOR
) -> list[ArchivoEntrada]:
    """Las entradas del par como `entrada_N` sin extensión, en orden INVERSO
    al alfabético para que el orden de subida no disimule un error de orden."""
    trabajo = destino / par.id
    trabajo.mkdir(parents=True)
    entradas: list[ArchivoEntrada] = []
    papeles = [p for p in reversed(sorted(nombres)) if p in par.archivos]
    for indice, papel in enumerate(papeles):
        copia = trabajo / f"entrada_{indice}"
        copia.write_bytes((directorio(procesador) / par.archivos[papel]).read_bytes())
        _base, _punto, formato = nombres[papel].rpartition(".")
        entradas.append(
            ArchivoEntrada(
                nombre_original=nombres[papel],
                ruta_temporal=copia,
                tamano_comprimido=copia.stat().st_size,
                formato=formato.lower(),
            )
        )
    return entradas


def esperadas(par: ParDeParidad, procesador: str = PROCESADOR) -> dict[str, bytes]:
    with zipfile.ZipFile(directorio(procesador) / par.archivos[_SALIDA]) as contenedor:
        return {nombre: contenedor.read(nombre) for nombre in contenedor.namelist()}


def _normalizado(celda: Any) -> Any:
    """`66` y `66.0` son la misma celda (ver la docstring)."""
    if isinstance(celda, float) and celda.is_integer():
        return int(celda)
    return celda


def celdas(datos: bytes) -> dict[str, list[list[tuple[Any, ...]]]]:
    """Valor con su tipo, formato de número y estilo de cada celda."""
    libro = load_workbook(io.BytesIO(datos))
    try:
        return {
            hoja.title: [
                [
                    (
                        type(_normalizado(c.value)),
                        _normalizado(c.value),
                        None if c.value is None else c.number_format,
                        None
                        if c.value is None
                        else (c.font.b, c.border.left.style, c.alignment.horizontal),
                    )
                    for c in fila
                ]
                for fila in hoja.iter_rows()
            ]
            for hoja in libro.worksheets
        }
    finally:
        libro.close()


def comparar(obtenidas: dict[str, Path], esperado: dict[str, bytes], par: ParDeParidad) -> None:
    """Mismos nombres; CSV byte por byte; libros celda por celda."""
    assert NOMBRE_OBSERVACIONES not in obtenidas, f"par {par.id}: observaciones inesperadas"
    assert sorted(obtenidas) == sorted(esperado), f"par {par.id}: otros archivos de salida"
    distintos = []
    for nombre, datos in esperado.items():
        propios = obtenidas[nombre].read_bytes()
        if nombre.endswith(".csv"):
            igual = propios == datos
        else:
            igual = celdas(propios) == celdas(datos)
        if not igual:
            distintos.append(nombre)
    assert not distintos, (
        f"par {par.id}: {len(distintos)} salidas no coinciden, p. ej. {distintos[:5]}. "
        f"Procedencia del esperado: {par.procedencia.generado_por} el "
        f"{par.procedencia.generado_el} en {par.procedencia.generado_en}."
    )


def test_las_99_particiones_coinciden(tmp_path: Path) -> None:
    for par in _pares():
        entradas = preparar(par, NOMBRES_ORIGINALES, tmp_path)
        procesador = MediosPagoReportes()
        assert procesador.validar(entradas) is None, f"par {par.id}: las entradas no validan"
        obtenidas = {s.nombre_propuesto: s.ruta_temporal for s in procesador.procesar(entradas)}
        comparar(obtenidas, esperadas(par), par)


def test_cada_par_declara_sus_entradas_y_su_salida() -> None:
    """El manifiesto es el contrato: un par a medias no es un par."""
    for par in _pares():
        assert {*NOMBRES_ORIGINALES, _SALIDA} <= set(par.archivos), f"par {par.id}: incompleto"
