"""Paridad de `AsientosContables_Carga` contra las 7 salidas históricas (ADR 0015).

**Estado: 7 pares, uno por tipo.** Las salidas esperadas son las que el
script produjo el 2025-08-26 —no regeneradas— y el módulo las reproduce
celda por celda con igualdad EXACTA en las 20 hojas. Con `FIXTURES_PARIDAD`
sin definir, todo esto saltea con aviso, y falla en CI.

**El reloj es parte del par.** Casi todas las cabeceras llevan como fecha de
cierre el día en que corrió el script (`datetime.now()`), así que la prueba
fija el reloj del módulo en `procedencia.generado_el`. Sin eso, la paridad
fallaría todos los días menos uno.

**El Excel se compara por contenido, sin tolerancia.** Los importes son
sumas de pandas que el módulo emula (Kahan en `groupby`, por pares en
`Series.sum`), y sobre estos siete pares la emulación da el mismo `float` bit
a bit. Si un par futuro difiriera en el último decimal, la decisión de
tolerarlo tiene que ser explícita, no un `approx` puesto de antemano. Sí se
tolera `200` contra `200.0`: pandas elige `int` o `float` por COLUMNA y el
módulo por valor, y para Excel son el mismo número. La clave `"01"` se
compara con su tipo: texto contra entero es una diferencia real.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from app.core.tipos import ArchivoEntrada
from app.procesadores.asientos_contables.modulo import NOMBRE_OBSERVACIONES, AsientosContables
from tests.paridad.manifiesto import (
    VARIABLE_DE_UBICACION,
    ParDeParidad,
    exigir_pares,
    resolver_raiz,
)

PROCESADOR = "asientos_contables"


def _pares() -> list[ParDeParidad]:
    """Compuerta de ADR 0015. Corta la prueba —fallando o salteando— si falta algo."""
    return exigir_pares(PROCESADOR)


def _directorio() -> Path:
    raiz = resolver_raiz()
    assert raiz is not None, f"{VARIABLE_DE_UBICACION} sin definir pese a haber pasado la compuerta"
    return raiz / PROCESADOR


def _ejecutar(par: ParDeParidad, destino: Path) -> dict[str, Path]:
    """Corre el módulo sobre la entrada del par, con el reloj en el día del par.

    La copia se llama `entrada_0`, sin extensión, como los temporales de
    `app/recepcion.py`. Se llama a `procesar` directo: lo que la paridad
    verifica son las reglas de negocio, no la plomería del proceso hijo.
    """
    origen = _directorio() / par.archivos["entrada"]
    trabajo = destino / par.id
    trabajo.mkdir(parents=True)
    copia = trabajo / "entrada_0"
    copia.write_bytes(origen.read_bytes())

    entrada = ArchivoEntrada(
        nombre_original=origen.name,
        ruta_temporal=copia,
        tamano_comprimido=copia.stat().st_size,
        formato="xlsx",
    )
    procesador = AsientosContables(hoy=lambda: par.procedencia.generado_el)
    assert procesador.validar([entrada]) is None, f"par {par.id}: la entrada no valida"
    return {s.nombre_propuesto: s.ruta_temporal for s in procesador.procesar([entrada])}


def _celdas(ruta: Path) -> dict[str, list[list[Any]]]:
    libro = load_workbook(ruta, data_only=True)
    try:
        return {h.title: [list(f) for f in h.iter_rows(values_only=True)] for h in libro.worksheets}
    finally:
        libro.close()


def _normalizado(celda: Any) -> Any:
    """`200` y `200.0` son la misma celda; `"01"` y `1`, no (ver la docstring)."""
    if isinstance(celda, float) and celda.is_integer():
        return int(celda)
    return celda


def _tipado(hojas: dict[str, list[list[Any]]]) -> dict[str, list[list[tuple[type, Any]]]]:
    return {
        nombre: [[(type(_normalizado(c)), _normalizado(c)) for c in fila] for fila in filas]
        for nombre, filas in hojas.items()
    }


def _salida_xlsx(salidas: dict[str, Path]) -> Path:
    (ruta,) = [ruta for nombre, ruta in salidas.items() if nombre.endswith(".xlsx")]
    return ruta


def test_el_excel_coincide_celda_por_celda(tmp_path: Path) -> None:
    for par in _pares():
        obtenido = _celdas(_salida_xlsx(_ejecutar(par, tmp_path)))
        esperado = _celdas(_directorio() / par.archivos["salida_xlsx"])
        assert list(obtenido) == list(esperado), f"par {par.id}: hojas distintas"
        assert _tipado(obtenido) == _tipado(esperado), (
            f"par {par.id}: el Excel no coincide. Procedencia del esperado: "
            f"{par.procedencia.generado_por} el {par.procedencia.generado_el} en "
            f"{par.procedencia.generado_en}."
        )


def test_ningun_par_real_tiene_observaciones(tmp_path: Path) -> None:
    """El riesgo abierto del TECH-DESIGN, convertido en comprobación.

    Un par con observaciones significa que la salida esperada codifica una
    pérdida silenciosa que la migración vino a hacer visible: hay que decidir
    cuál requisito manda, y esta prueba existe para que esa decisión no se
    tome por omisión.
    """
    for par in _pares():
        salidas = _ejecutar(par, tmp_path)
        assert NOMBRE_OBSERVACIONES not in salidas, (
            f"par {par.id}: el módulo reporta observaciones sobre una entrada real. Revisar "
            f"{salidas.get(NOMBRE_OBSERVACIONES)}."
        )


def test_cada_par_declara_entrada_y_salida() -> None:
    """El manifiesto es el contrato: un par a medias no es un par."""
    for par in _pares():
        assert {"entrada", "salida_xlsx"} <= set(par.archivos), f"par {par.id}: incompleto"
