"""Paridad de `flujo-caja-ingresos` contra las dos salidas históricas (ADR 0015).

**Estado: 1 par con dos entradas y dos salidas.** Las entradas son los libros
reales de cuentas Fideicomisas y Operativas de 2025 (hasta fines de julio) y
las salidas esperadas son las que los dos scripts produjeron en agosto de
2025 —no regeneradas—: `Consolidado_Ingresos.xlsx` y `Banco_Nacion.xlsx`. El
módulo las reproduce celda por celda con igualdad EXACTA, en las tres hojas y
con el orden de filas de pandas (empates incluidos). Con `FIXTURES_PARIDAD`
sin definir, todo esto saltea con aviso, y falla en CI.

**El Excel se compara por contenido, sin tolerancia**, con la misma única
excepción que `asientos_contables`: `66` y `66.0` son la misma celda (Excel
no distingue enteros de decimales y openpyxl escribe los dos igual). Las
fechas se comparan como `datetime` y `F. Operación` de Banco de la Nación
como TEXTO: texto contra fecha es una diferencia real.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from app.core.tipos import ArchivoEntrada
from app.procesadores.flujo_caja_ingresos.modulo import (
    NOMBRE_BANCO_NACION,
    NOMBRE_CONSOLIDADO,
    NOMBRE_OBSERVACIONES,
    FlujoCajaIngresos,
)
from tests.paridad.manifiesto import (
    VARIABLE_DE_UBICACION,
    ParDeParidad,
    exigir_pares,
    resolver_raiz,
)

PROCESADOR = "flujo_caja_ingresos"

_ENTRADAS: tuple[str, ...] = ("entrada_fideicomisas", "entrada_operativas")
_SALIDAS: dict[str, str] = {
    "salida_consolidado": NOMBRE_CONSOLIDADO,
    "salida_banco_nacion": NOMBRE_BANCO_NACION,
}
"""Papel del manifiesto → nombre con que el módulo propone la salida."""

_NOMBRES_ORIGINALES: dict[str, str] = {
    "entrada_fideicomisas": "LAMSAC - Mov Ctas Fideicomisas 2025 FL.xlsx",
    "entrada_operativas": "LAMSAC - Mov Ctas Operativas 2025 FL_ffff.xlsx",
}
"""Los nombres con que llegaron los libros. El tipo se detecta por las hojas,
así que el nombre no decide nada; se conserva para que las observaciones, si
las hubiera, digan lo mismo que diría el portal."""


def _pares() -> list[ParDeParidad]:
    """Compuerta de ADR 0015. Corta la prueba —fallando o salteando— si falta algo."""
    return exigir_pares(PROCESADOR)


def _directorio() -> Path:
    raiz = resolver_raiz()
    assert raiz is not None, f"{VARIABLE_DE_UBICACION} sin definir pese a haber pasado la compuerta"
    return raiz / PROCESADOR


def _ejecutar(par: ParDeParidad, destino: Path) -> dict[str, Path]:
    """Corre el módulo sobre las dos entradas del par, como `entrada_0` y `entrada_1`.

    Sin extensión, como los temporales de `app/recepcion.py`. Se llama a
    `procesar` directo: lo que la paridad verifica son las reglas de negocio,
    no la plomería del proceso hijo.
    """
    trabajo = destino / par.id
    trabajo.mkdir(parents=True)
    entradas: list[ArchivoEntrada] = []
    for indice, papel in enumerate(_ENTRADAS):
        copia = trabajo / f"entrada_{indice}"
        copia.write_bytes((_directorio() / par.archivos[papel]).read_bytes())
        entradas.append(
            ArchivoEntrada(
                nombre_original=_NOMBRES_ORIGINALES[papel],
                ruta_temporal=copia,
                tamano_comprimido=copia.stat().st_size,
                formato="xlsx",
            )
        )
    procesador = FlujoCajaIngresos()
    assert procesador.validar(entradas) is None, f"par {par.id}: las entradas no validan"
    return {s.nombre_propuesto: s.ruta_temporal for s in procesador.procesar(entradas)}


def _celdas(ruta: Path) -> dict[str, list[list[Any]]]:
    libro = load_workbook(ruta, read_only=True, data_only=True)
    try:
        return {h.title: [list(f) for f in h.iter_rows(values_only=True)] for h in libro.worksheets}
    finally:
        libro.close()


def _normalizado(celda: Any) -> Any:
    """`66` y `66.0` son la misma celda (ver la docstring)."""
    if isinstance(celda, float) and celda.is_integer():
        return int(celda)
    return celda


def _tipado(hojas: dict[str, list[list[Any]]]) -> dict[str, list[list[tuple[type, Any]]]]:
    return {
        nombre: [[(type(_normalizado(c)), _normalizado(c)) for c in fila] for fila in filas]
        for nombre, filas in hojas.items()
    }


def test_los_dos_excel_coinciden_celda_por_celda(tmp_path: Path) -> None:
    for par in _pares():
        salidas = _ejecutar(par, tmp_path)
        for papel, nombre in _SALIDAS.items():
            assert nombre in salidas, f"par {par.id}: el módulo no produjo {nombre}"
            obtenido = _celdas(salidas[nombre])
            esperado = _celdas(_directorio() / par.archivos[papel])
            assert list(obtenido) == list(esperado), f"par {par.id}/{nombre}: hojas distintas"
            for hoja in esperado:
                assert len(obtenido[hoja]) == len(esperado[hoja]), (
                    f"par {par.id}/{nombre}/{hoja}: {len(obtenido[hoja])} filas, "
                    f"se esperaban {len(esperado[hoja])}"
                )
            assert _tipado(obtenido) == _tipado(esperado), (
                f"par {par.id}/{nombre}: el Excel no coincide. Procedencia del esperado: "
                f"{par.procedencia.generado_por} el {par.procedencia.generado_el} en "
                f"{par.procedencia.generado_en}."
            )


def test_ningun_par_real_tiene_observaciones(tmp_path: Path) -> None:
    """Un par con observaciones significa que la salida esperada codifica una
    pérdida silenciosa que la migración vino a hacer visible: hay que decidir
    cuál requisito manda, y esta prueba existe para que esa decisión no se
    tome por omisión."""
    for par in _pares():
        salidas = _ejecutar(par, tmp_path)
        assert NOMBRE_OBSERVACIONES not in salidas, (
            f"par {par.id}: el módulo reporta observaciones sobre una entrada real. Revisar "
            f"{salidas.get(NOMBRE_OBSERVACIONES)}."
        )


def test_cada_par_declara_sus_dos_entradas_y_sus_dos_salidas() -> None:
    """El manifiesto es el contrato: un par a medias no es un par."""
    for par in _pares():
        assert {*_ENTRADAS, *_SALIDAS} <= set(par.archivos), f"par {par.id}: incompleto"
