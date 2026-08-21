"""Paridad de `Contado_Carga` contra los 3 pares reales (ítem #16; ADR 0015).

**Estado: 1 par de los 3 que exige el PRD.** El par 01 se dio de alta el
2026-08-20 a partir de un extracto real y la suite pasa contra él; faltan el
02 y el 03. Con `FIXTURES_PARIDAD` sin definir —el caso normal en una máquina
recién clonada— todo esto saltea con aviso, y falla en CI.

**El riesgo abierto del TECH-DESIGN quedó cerrado para el par 01, no en
general.** Ese riesgo es real: si un par contiene una fila que el script viejo
descartó en silencio, el módulo endurecido —haciendo exactamente lo correcto—
produce una salida distinta a la esperada, y ahí hay que decidir cuál de los
dos requisitos manda. Sobre el par 01 no pasa: 31 registros, cero descartes.
`test_ningun_par_real_tiene_descartes` mantiene la guardia para los que
vengan, y el `descartes.txt` de la ejecución es la evidencia de esa
conversación si algún día se dispara.

**El TXT se compara byte a byte y el Excel por contenido normalizado**, tal
como fija ADR 0015: el TXT es el que carga la codificación y el fin de línea,
que es lo que la paridad tiene que proteger.
"""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import load_workbook

from app.core.tipos import ArchivoEntrada
from app.procesadores.contado_carga.modulo import NOMBRE_TXT, NOMBRE_XLSX, ContadoCarga
from tests.paridad.manifiesto import (
    VARIABLE_DE_UBICACION,
    ParDeParidad,
    exigir_pares,
    resolver_raiz,
)

PROCESADOR = "contado_carga"


def _pares() -> list[ParDeParidad]:
    """Compuerta de ADR 0015. Corta la prueba —fallando o salteando— si falta algo."""
    return exigir_pares(PROCESADOR)


def _directorio() -> Path:
    raiz = resolver_raiz()
    assert raiz is not None, f"{VARIABLE_DE_UBICACION} sin definir pese a haber pasado la compuerta"
    return raiz / PROCESADOR


def _ejecutar(par: ParDeParidad, destino: Path) -> dict[str, Path]:
    """Corre el módulo sobre la entrada del par, en un directorio de trabajo propio.

    Se llama a `procesar` directo y no a través del pipeline: lo que la
    paridad verifica son las reglas de negocio, no la plomería del proceso
    hijo, que ya tiene sus propias pruebas.
    """
    origen = _directorio() / par.archivos["entrada"]
    trabajo = destino / par.id
    trabajo.mkdir(parents=True)
    copia = trabajo / origen.name
    copia.write_bytes(origen.read_bytes())

    entrada = ArchivoEntrada(
        nombre_original=origen.name,
        ruta_temporal=copia,
        tamano_comprimido=copia.stat().st_size,
        formato="xlsx",
    )
    return {s.nombre_propuesto: s.ruta_temporal for s in ContadoCarga().procesar([entrada])}


def _celdas(ruta: Path) -> list[list[object]]:
    libro = load_workbook(ruta, data_only=True)
    try:
        return [list(fila) for fila in libro.worksheets[0].iter_rows(values_only=True)]
    finally:
        libro.close()


def test_el_txt_coincide_byte_a_byte(tmp_path: Path) -> None:
    """Codificación y fin de línea incluidos. Es el punto entero de ADR 0015."""
    for par in _pares():
        salidas = _ejecutar(par, tmp_path)
        esperado = (_directorio() / par.archivos["salida_txt"]).read_bytes()
        obtenido = salidas[NOMBRE_TXT].read_bytes()
        assert obtenido == esperado, (
            f"par {par.id}: el TXT no coincide byte a byte. Procedencia del esperado: "
            f"{par.procedencia.generado_por} el {par.procedencia.generado_el} en "
            f"{par.procedencia.generado_en}, fin de línea {par.procedencia.fin_de_linea}."
        )


def test_el_excel_coincide_por_contenido(tmp_path: Path) -> None:
    for par in _pares():
        salidas = _ejecutar(par, tmp_path)
        esperado = _celdas(_directorio() / par.archivos["salida_xlsx"])
        assert _celdas(salidas[NOMBRE_XLSX]) == esperado, f"par {par.id}: el Excel no coincide"


def test_ningun_par_real_tiene_descartes(tmp_path: Path) -> None:
    """El riesgo abierto del TECH-DESIGN, convertido en comprobación.

    Un par con descartes significa que la salida esperada codifica el
    comportamiento que la migración vino a corregir. No es un fallo del
    módulo: es la señal de que hay que decidir cuál requisito manda, y esta
    prueba existe para que esa decisión no se tome por omisión.
    """
    from app.procesadores.contado_carga.modulo import NOMBRE_DESCARTES

    for par in _pares():
        salidas = _ejecutar(par, tmp_path)
        assert NOMBRE_DESCARTES not in salidas, (
            f"par {par.id}: el módulo endurecido descarta filas que el script viejo perdía en "
            f"silencio. Revisar {salidas.get(NOMBRE_DESCARTES)} y decidir explícitamente si "
            "manda la paridad byte a byte o el endurecimiento obligatorio (TECH-DESIGN, "
            "riesgos abiertos)."
        )


@pytest.mark.parametrize("papel", ["entrada", "salida_xlsx", "salida_txt"])
def test_cada_par_declara_los_tres_archivos(papel: str) -> None:
    """El manifiesto es el contrato: un par a medias no es un par."""
    for par in _pares():
        assert papel in par.archivos, f"par {par.id}: falta '{papel}' en el manifiesto"
