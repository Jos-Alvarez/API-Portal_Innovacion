"""Pruebas del procesador `flujo-caja-ingresos`.

Los libros reales no están en este repositorio y no pueden estarlo (ADR 0015),
así que estas pruebas arman libros sintéticos con `openpyxl` —con cada hoja en
el layout de la muestra real: encabezado en su fila, datos desde
`fila_inicio`— y los guardan con nombres SIN extensión (`entrada_0`) como los
que escribe `app/recepcion.py`. Prueban las **reglas**; la paridad contra las
salidas históricas es `tests/paridad/test_flujo_caja_ingresos.py`.
"""

from __future__ import annotations

import io
import zipfile
from collections.abc import Mapping, Sequence
from datetime import datetime
from pathlib import Path
from typing import Any

import pytest
from openpyxl import Workbook, load_workbook

from app.core.errores import ErrorContenido
from app.core.tipos import ArchivoEntrada
from app.procesadores.flujo_caja_ingresos.modulo import (
    CLAVE,
    COLUMNAS_BANCO_NACION,
    COLUMNAS_FIDEICOMISO,
    COLUMNAS_OPERATIVAS,
    FIN_DE_LINEA,
    HOJA_BANCO_NACION,
    HOJAS_FIDEICOMISO,
    HOJAS_OPERATIVAS,
    MOTIVO_FECHA,
    MOTIVO_FECHA_TEXTO,
    MOTIVO_HOJA_AUSENTE,
    MOTIVO_HOJA_SIN_FILAS,
    MOTIVO_IMPORTE,
    NOMBRE_BANCO_NACION,
    NOMBRE_CONSOLIDADO,
    NOMBRE_OBSERVACIONES,
    FlujoCajaIngresos,
    _ConfigHoja,
    _orden_de_pandas,
    calcular_semana,
    calcular_semanas_mes,
    categorizar,
    limpiar_valor_monetario,
)

_CONFIGS: dict[str, _ConfigHoja] = {
    c.nombre: c for c in (*HOJAS_FIDEICOMISO, *HOJAS_OPERATIVAS, HOJA_BANCO_NACION)
}
Datos = Mapping[str, Sequence[Sequence[object]]]


def _d(mes: int, dia: int, anio: int = 2025, hora: int = 0) -> datetime:
    return datetime(anio, mes, dia, hora)


def _hoja(config: _ConfigHoja, filas: Sequence[Sequence[object]]) -> list[list[object]]:
    """La hoja con el layout real: vacía hasta el encabezado, que va en su fila
    y con sus rótulos sobre las columnas configuradas, y los datos a
    continuación. Cada fila de datos trae un valor por columna configurada."""
    ancho = max(config.columnas) + 1
    prefijo = max(config.fila_inicio, config.fila_encabezado + 1)
    matriz: list[list[object]] = [[None] * ancho for _ in range(prefijo)]
    for columna, rotulo in zip(config.columnas, config.encabezados, strict=True):
        matriz[config.fila_encabezado][columna - 1] = rotulo
    for datos in filas:
        fila: list[object] = [None] * ancho
        for columna, valor in zip(config.columnas, datos, strict=True):
            fila[columna - 1] = valor
        matriz.append(fila)
    return matriz


def _bytes_de_libro(datos: Datos, crudas: Datos | None = None) -> bytes:
    """`datos` pasa por `_hoja`; `crudas` se escribe tal cual (hojas ajenas)."""
    libro = Workbook()
    libro.remove(libro["Sheet"])
    for nombre, filas in datos.items():
        hoja = libro.create_sheet(nombre)
        for fila in _hoja(_CONFIGS[nombre], filas):
            hoja.append(fila)
    for nombre, filas in (crudas or {}).items():
        hoja = libro.create_sheet(nombre)
        for cruda in filas:
            hoja.append(list(cruda))
    memoria = io.BytesIO()
    libro.save(memoria)
    return memoria.getvalue()


def _fideicomisas(
    **reemplazos: Sequence[Sequence[object]],
) -> dict[str, Sequence[Sequence[object]]]:
    """Las seis hojas de Fideicomiso con un abono cada una; `reemplazos` por nombre."""
    datos: dict[str, Sequence[Sequence[object]]] = {
        c.nombre: [(_d(1, 2), f"ABONO {c.nombre}", 100)] for c in HOJAS_FIDEICOMISO
    }
    datos.update(reemplazos)
    return datos


def _operativas() -> dict[str, Sequence[Sequence[object]]]:
    """Las cuatro hojas de Operativas y BCO NACION, cada una con un movimiento."""
    return {
        "IBK OPE 223": [(_d(1, 2), "ABONO REMESAS", None, 100)],
        "IBK OPE $ 230": [(_d(1, 2), "ITF", -1.5, None)],
        "BBVA 43332 PEN": [(_d(1, 3), "ITF", -2)],
        "BBVA 43340 USD": [(_d(1, 3), "ABONO RECEP. TRANSF.", 50)],
        "BCO NACION": [("2025.01.02", None, 167)],
    }


class _Fabrica:
    """Escribe libros como `entrada_N`, sin extensión, igual que la recepción real."""

    def __init__(self, directorio: Path) -> None:
        self.directorio = directorio
        self.contador = 0

    def entrada(
        self, datos: Datos, nombre: str = "libro.xlsx", crudas: Datos | None = None
    ) -> ArchivoEntrada:
        ruta = self.directorio / f"entrada_{self.contador}"
        self.contador += 1
        ruta.write_bytes(_bytes_de_libro(datos, crudas))
        return ArchivoEntrada(
            nombre_original=nombre,
            ruta_temporal=ruta,
            tamano_comprimido=ruta.stat().st_size,
            formato="xlsx",
        )


@pytest.fixture
def fabrica(tmp_path: Path) -> _Fabrica:
    return _Fabrica(tmp_path)


@pytest.fixture
def procesador() -> FlujoCajaIngresos:
    return FlujoCajaIngresos()


def _procesar(procesador: FlujoCajaIngresos, *entradas: ArchivoEntrada) -> dict[str, Path]:
    assert procesador.validar(list(entradas)) is None
    return {s.nombre_propuesto: s.ruta_temporal for s in procesador.procesar(list(entradas))}


def _leer(ruta: Path) -> dict[str, list[list[Any]]]:
    libro = load_workbook(ruta)
    try:
        return {h.title: [list(f) for f in h.iter_rows(values_only=True)] for h in libro.worksheets}
    finally:
        libro.close()


def _error(procesador: FlujoCajaIngresos, *entradas: ArchivoEntrada) -> ErrorContenido:
    error = procesador.validar(list(entradas))
    assert isinstance(error, ErrorContenido)
    return error


def _observaciones(ruta: Path) -> list[list[str]]:
    """Las líneas de detalle de `observaciones.txt`, partidas por tabulador."""
    lineas = ruta.read_bytes().decode("utf-8").split(FIN_DE_LINEA)
    return [linea.split("\t") for linea in lineas[3:] if linea]


# --- Detección del tipo de libro -------------------------------------------------


def test_el_tipo_se_detecta_por_las_hojas_y_no_por_el_nombre(
    procesador: FlujoCajaIngresos, fabrica: _Fabrica
) -> None:
    """Los nombres no dicen nada; las hojas, todo. Salen las dos salidas."""
    fid = fabrica.entrada(_fideicomisas(), nombre="uno.xlsx")
    ope = fabrica.entrada(_operativas(), nombre="dos.xlsx")
    salidas = _procesar(procesador, ope, fid)
    assert list(salidas) == [NOMBRE_CONSOLIDADO, NOMBRE_BANCO_NACION]
    consolidado = _leer(salidas[NOMBRE_CONSOLIDADO])
    assert list(consolidado) == ["Consolidado Fideicomiso", "Consolidado Operativas"]
    assert list(_leer(salidas[NOMBRE_BANCO_NACION])) == ["Banco Nacion"]


def test_hojas_de_los_dos_tipos_en_un_libro_es_tipo_no_reconocido(
    procesador: FlujoCajaIngresos, fabrica: _Fabrica
) -> None:
    datos = {**_fideicomisas(), **_operativas()}
    error = _error(procesador, fabrica.entrada(datos, nombre="mezcla Fideicomisas.xlsx"))
    assert error.contexto == {
        "archivo": "mezcla Fideicomisas.xlsx",
        "motivo": "tipo_no_reconocido",
        "columna": None,
    }


def test_solo_bco_nacion_alcanza_para_operativas(
    procesador: FlujoCajaIngresos, fabrica: _Fabrica
) -> None:
    entrada = fabrica.entrada({"BCO NACION": [("2025.01.02", None, 167)]})
    salidas = _procesar(procesador, entrada)
    assert list(salidas) == [NOMBRE_BANCO_NACION, NOMBRE_OBSERVACIONES]
    ausentes = {o[1] for o in _observaciones(salidas[NOMBRE_OBSERVACIONES])}
    assert ausentes == {c.nombre for c in HOJAS_OPERATIVAS}


@pytest.mark.parametrize(
    ("nombre", "hoja"),
    [
        ("LAMSAC - Mov Ctas Fideicomisas 2025.xlsx", "IBK 1106 FID"),
        ("lamsac - mov ctas OPERATIVAS.xlsx", "IBK OPE 223"),
    ],
)
def test_sin_hojas_configuradas_decide_el_nombre_y_falta_la_hoja(
    procesador: FlujoCajaIngresos, fabrica: _Fabrica, nombre: str, hoja: str
) -> None:
    """La regla del `glob` de los scripts, como respaldo: el error dice qué se esperaba."""
    entrada = fabrica.entrada({}, nombre=nombre, crudas={"Hoja1": [["nada"]]})
    error = _error(procesador, entrada)
    assert error.contexto == {"archivo": nombre, "motivo": "hoja_faltante", "columna": hoja}


@pytest.mark.parametrize("nombre", ["cualquiera.xlsx", "Fideicomisas y Operativas.xlsx"])
def test_sin_hojas_ni_nombre_util_es_tipo_no_reconocido(
    procesador: FlujoCajaIngresos, fabrica: _Fabrica, nombre: str
) -> None:
    entrada = fabrica.entrada({}, nombre=nombre, crudas={"Hoja1": [["nada"]]})
    assert _error(procesador, entrada).contexto["motivo"] == "tipo_no_reconocido"


def test_dos_libros_del_mismo_tipo_es_tipo_duplicado(
    procesador: FlujoCajaIngresos, fabrica: _Fabrica
) -> None:
    a = fabrica.entrada(_fideicomisas(), nombre="a.xlsx")
    b = fabrica.entrada(_fideicomisas(), nombre="b.xlsx")
    esperado = {"archivo": "b.xlsx", "motivo": "tipo_duplicado", "columna": None}
    assert _error(procesador, a, b).contexto == esperado
    with pytest.raises(ErrorContenido) as capturado:
        procesador.procesar([a, b])
    assert capturado.value.contexto == esperado


# --- Las tres combinaciones de entrada -----------------------------------------


def test_solo_operativas_da_consolidado_de_operativas_y_banco_nacion(
    procesador: FlujoCajaIngresos, fabrica: _Fabrica
) -> None:
    salidas = _procesar(procesador, fabrica.entrada(_operativas()))
    assert list(salidas) == [NOMBRE_CONSOLIDADO, NOMBRE_BANCO_NACION]
    assert list(_leer(salidas[NOMBRE_CONSOLIDADO])) == ["Consolidado Operativas"]


def test_solo_fideicomisas_da_solo_el_consolidado_de_fideicomiso(
    procesador: FlujoCajaIngresos, fabrica: _Fabrica
) -> None:
    salidas = _procesar(procesador, fabrica.entrada(_fideicomisas()))
    assert list(salidas) == [NOMBRE_CONSOLIDADO]
    hojas = _leer(salidas[NOMBRE_CONSOLIDADO])
    assert list(hojas) == ["Consolidado Fideicomiso"]
    assert hojas["Consolidado Fideicomiso"][0] == list(COLUMNAS_FIDEICOMISO)
    assert len(hojas["Consolidado Fideicomiso"]) == 1 + len(HOJAS_FIDEICOMISO)


# --- Fideicomiso ---------------------------------------------------------------


def test_fideicomiso_filtros_columnas_totales_y_orden(
    procesador: FlujoCajaIngresos, fabrica: _Fabrica
) -> None:
    """Sólo importes > 0 con fecha; lo ilegible se reporta; el blanco no."""
    datos: Datos = {
        "IBK 1106 FID": [
            (_d(1, 3), "A", 10.5),  # fila 17 de Excel
            (_d(1, 2), "B", -5),  # negativo: fuera
            (_d(1, 2), "C", 0),  # cero: fuera
            (None, "Saldo Final 02-01-2025", 999),  # sin fecha: fuera, sin reporte
            (_d(1, 2), "D", "1,234"),  # fila 21: ilegible, fuera y reportado
            (_d(1, 2), "E", "\xa0"),  # blanco: fuera, sin reporte
            (_d(1, 2), "F", " 12 "),  # texto numérico: `to_numeric` lo lee
        ],
        "BCP 306-0-76 FID REC": [(_d(1, 2), None, 7)],
    }
    salidas = _procesar(procesador, fabrica.entrada(datos, nombre="fid.xlsx"))
    filas = _leer(salidas[NOMBRE_CONSOLIDADO])["Consolidado Fideicomiso"]

    assert filas[0] == list(COLUMNAS_FIDEICOMISO)
    ibk, bcp = "IBK 1106 FID", "BCP 306-0-76 FID REC"
    assert filas[1:] == [
        [_d(1, 2), "F", 12, "PEN", "ABONO", "FIDEICOMISO", ibk, "Ene", "Semana 1", 22.5, 29.5],
        [_d(1, 2), None, 7, "PEN", "ABONO", "FIDEICOMISO", bcp, "Ene", "Semana 1", 7, 29.5],
        [_d(1, 3), "A", 10.5, "PEN", "ABONO", "FIDEICOMISO", ibk, "Ene", "Semana 1", 22.5, 29.5],
    ]
    observaciones = _observaciones(salidas[NOMBRE_OBSERVACIONES])
    assert ["fid.xlsx", ibk, "21", MOTIVO_IMPORTE, "1,234"] in observaciones
    assert sorted(o[1] for o in observaciones if o[3] == MOTIVO_HOJA_AUSENTE) == sorted(
        c.nombre for c in HOJAS_FIDEICOMISO if c.nombre not in datos
    )


def test_fideicomiso_fecha_con_hora_el_ultimo_dia_queda_sin_semana(
    procesador: FlujoCajaIngresos, fabrica: _Fabrica
) -> None:
    """La cuarta semana termina a las 00:00 del último día: 10:00 cae fuera."""
    datos = _fideicomisas(**{"IBK 1106 FID": [(_d(1, 31, hora=10), "X", 5)]})
    filas = _leer(_procesar(procesador, fabrica.entrada(datos))[NOMBRE_CONSOLIDADO])
    fila = next(f for f in filas["Consolidado Fideicomiso"] if f[1] == "X")
    assert fila[7:9] == ["Ene", None]  # Semana "" se escribe como celda vacía


def test_fechas_como_texto_y_numeros_en_el_consolidado(
    procesador: FlujoCajaIngresos, fabrica: _Fabrica
) -> None:
    """Texto: mes primero y reportado. Número u hoja ilegible: sin fecha y reportado."""
    datos = _fideicomisas(
        **{
            "IBK 1106 FID": [
                (_d(1, 2), "real", 1),
                ("02/01/2025", "texto", 2),  # 1 de febrero, no 2 de enero
                ("2025-01-05", "iso", 3),  # año primero: no se reporta
                ("no es fecha", "ilegible", 4),
                (45000, "numero", 5),
            ]
        }
    )
    salidas = _procesar(procesador, fabrica.entrada(datos, nombre="f.xlsx"))
    filas = {f[1]: f for f in _leer(salidas[NOMBRE_CONSOLIDADO])["Consolidado Fideicomiso"]}
    assert filas["texto"][0] == _d(2, 1)
    assert filas["texto"][7:9] == ["Feb", "Semana 1"]
    assert filas["iso"][0] == _d(1, 5)
    for sin_fecha in ("ilegible", "numero"):
        assert filas[sin_fecha][0] is None
        assert filas[sin_fecha][7:9] == [None, None]
    # Las filas sin fecha van al final, como el NaT de `sort_values`.
    orden = [f[1] for f in _leer(salidas[NOMBRE_CONSOLIDADO])["Consolidado Fideicomiso"][1:]]
    assert orden[-2:] == ["ilegible", "numero"]
    motivos = [(o[2], o[3], o[4]) for o in _observaciones(salidas[NOMBRE_OBSERVACIONES])]
    assert motivos == [
        ("18", MOTIVO_FECHA_TEXTO, "02/01/2025"),
        ("20", MOTIVO_FECHA, "no es fecha"),
        ("21", MOTIVO_FECHA, "45000"),
    ]


# --- Operativas ----------------------------------------------------------------


def test_operativas_cargos_y_abonos_por_tipo_de_hoja(
    procesador: FlujoCajaIngresos, fabrica: _Fabrica
) -> None:
    datos: Datos = {
        "IBK OPE 223": [
            (_d(1, 2), "COMISION MANTENIMIENTO", -50, None),  # fila 16
            (_d(1, 2), "ABONO REMESAS", "\xa0", "1,000.50"),  # fila 17
            (_d(1, 2), "centavos", 0.005, 0.004),  # fila 18: ninguno pasa de 0,01
            (_d(1, 2), "CASH PAGO", "abc", 10),  # fila 19: cargo ilegible
        ],
        "BBVA 43332 PEN": [
            (_d(1, 3), "ITF", -30),
            (_d(1, 3), "PAGO PROVEEDORES X", 40),
            (_d(1, 3), "casi nada", 0.01),
        ],
        "BCO NACION": [("2025.01.02", None, 1)],
    }
    salidas = _procesar(procesador, fabrica.entrada(datos, nombre="ope.xlsx"))
    filas = _leer(salidas[NOMBRE_CONSOLIDADO])["Consolidado Operativas"]

    assert filas[0] == list(COLUMNAS_OPERATIVAS)
    ibk, bbva = "IBK OPE 223", "BBVA 43332 PEN"
    assert [f[:10] for f in filas[1:]] == [
        [_d(1, 2), "COMISION MANTENIMIENTO", 50, 0, "PEN", "OPERATIVAS", ibk]
        + ["COMISIONES / MANT Y PORTES", "Ene", "Semana 1"],
        [_d(1, 2), "ABONO REMESAS", 0, 1000.5, "PEN", "OPERATIVAS", ibk, "MONEDAS"]
        + ["Ene", "Semana 1"],
        [_d(1, 2), "CASH PAGO", 0, 10, "PEN", "OPERATIVAS", ibk, "PAGO PROVEEDORES (PEN)"]
        + ["Ene", "Semana 1"],
        [_d(1, 3), "ITF", 30, 0, "PEN", "OPERATIVAS", bbva, "ITF", "Ene", "Semana 1"],
        [_d(1, 3), "PAGO PROVEEDORES X", 0, 40, "PEN", "OPERATIVAS", bbva]
        + ["PAGO PROVEEDORES (PEN)", "Ene", "Semana 1"],
    ]
    # Totales de la primera fila: hoja, moneda y categoría.
    assert filas[1][10:] == [50, 1010.5, 80, 1050.5, 50, 0]
    assert ["ope.xlsx", ibk, "19", MOTIVO_IMPORTE, "abc"] in _observaciones(
        salidas[NOMBRE_OBSERVACIONES]
    )


def test_operativas_redondea_al_exportar_y_totaliza_sin_redondear(
    procesador: FlujoCajaIngresos, fabrica: _Fabrica
) -> None:
    datos: Datos = {
        "BBVA 43340 USD": [(_d(3, 3), "ITF", -10.004), (_d(3, 4), "ITF", -10.004)],
        "BCO NACION": [("2025.01.02", None, 1)],
    }
    filas = _leer(_procesar(procesador, fabrica.entrada(datos))[NOMBRE_CONSOLIDADO])
    operativas = filas["Consolidado Operativas"]
    assert [f[2] for f in operativas[1:]] == [10.0, 10.0]
    assert operativas[1][10] == pytest.approx(20.008)
    assert operativas[1][4] == "USD"


# --- Banco de la Nación ----------------------------------------------------------


def test_banco_nacion_texto_filtros_importes_y_totales(
    procesador: FlujoCajaIngresos, fabrica: _Fabrica
) -> None:
    datos = {
        "BCO NACION": [
            ("NO HAY MVTO. 01/01", None, None),  # fila 3: letras, fuera
            ("2025.01.02", None, 167),  # fila 4
            ("2025.01.03", "\xa0", "200"),  # fila 5: cargo en blanco vale 0
            (_d(1, 7), 5, None),  # fila 6: una fecha de verdad, como texto
            ("2025.01.08", "abc", 4),  # fila 7: cargo ilegible vale 0
            ("   ", 9, 9),  # fila 8: fecha en blanco, fuera
            ("2025.02.30", 1, 1),  # fila 9: fecha imposible
            ("07.01.2025", None, 3),  # fila 10: mes primero → 1 de julio
        ]
    }
    salidas = _procesar(procesador, fabrica.entrada(datos, nombre="ope.xlsx"))
    filas = _leer(salidas[NOMBRE_BANCO_NACION])["Banco Nacion"]

    meta = ["PEN", "OPERATIVA", "BCO NACION"]
    assert filas[0] == list(COLUMNAS_BANCO_NACION)
    assert filas[1:] == [
        ["2025.01.02", 0, 167, *meta, "Ene", "Semana 1", 5, 371],
        ["2025.01.03", 0, 200, *meta, "Ene", "Semana 1", 5, 371],
        ["2025-01-07 00:00:00", 5, 0, *meta, "Ene", "Semana 1", 5, 371],
        ["2025.01.08", 0, 4, *meta, "Ene", "Semana 1", 5, 371],
        ["2025.02.30", 1, 1, *meta, None, None, 1, 1],
        ["07.01.2025", 0, 3, *meta, "Jul", "Semana 1", 0, 3],
    ]
    motivos = [(o[2], o[3], o[4]) for o in _observaciones(salidas[NOMBRE_OBSERVACIONES])]
    assert ("7", MOTIVO_IMPORTE, "abc") in motivos
    assert ("9", MOTIVO_FECHA, "2025.02.30") in motivos
    assert ("10", MOTIVO_FECHA_TEXTO, "07.01.2025") in motivos
    assert not [m for m in motivos if m[0] in {"3", "4", "5", "6", "8"}]


def test_operativas_sin_bco_nacion_no_da_banco_nacion(
    procesador: FlujoCajaIngresos, fabrica: _Fabrica
) -> None:
    datos = _operativas()
    del datos["BCO NACION"]
    salidas = _procesar(procesador, fabrica.entrada(datos, nombre="ope.xlsx"))
    assert list(salidas) == [NOMBRE_CONSOLIDADO, NOMBRE_OBSERVACIONES]
    assert _observaciones(salidas[NOMBRE_OBSERVACIONES]) == [
        ["ope.xlsx", "BCO NACION", "", MOTIVO_HOJA_AUSENTE, ""]
    ]


def test_bco_nacion_sin_filas_se_escribe_solo_con_titulos(
    procesador: FlujoCajaIngresos, fabrica: _Fabrica
) -> None:
    """Como el script: exporta el `DataFrame` vacío. Y se reporta."""
    datos = {**_operativas(), "BCO NACION": [("Fecha", "Cargo", "Abono")]}
    salidas = _procesar(procesador, fabrica.entrada(datos, nombre="ope.xlsx"))
    assert _leer(salidas[NOMBRE_BANCO_NACION])["Banco Nacion"] == [list(COLUMNAS_BANCO_NACION)]
    assert _observaciones(salidas[NOMBRE_OBSERVACIONES]) == [
        ["ope.xlsx", "BCO NACION", "", MOTIVO_HOJA_SIN_FILAS, ""]
    ]


# --- Reglas verbatim -------------------------------------------------------------


@pytest.mark.parametrize(
    ("mes", "anio", "fines"),
    [
        # Enero de 2025 empieza miércoles: la semana 1 va hasta el domingo SIGUIENTE.
        (1, 2025, [12, 19, 26, 31]),
        # Junio de 2025 empieza domingo: la semana 1 dura ocho días.
        (6, 2025, [8, 15, 22, 30]),
        # Septiembre de 2025 empieza lunes: la semana 1 termina ese domingo.
        (9, 2025, [7, 14, 21, 30]),
        # Noviembre de 2025 empieza sábado: hasta el domingo siguiente.
        (11, 2025, [9, 16, 23, 30]),
        # Febrero de 2026 empieza domingo y tiene 28 días.
        (2, 2026, [8, 15, 22, 28]),
    ],
)
def test_calcular_semanas_mes(mes: int, anio: int, fines: list[int]) -> None:
    semanas = calcular_semanas_mes(mes, anio)
    assert [fin.day for _inicio, fin in semanas] == fines
    assert semanas[0][0] == datetime(anio, mes, 1)
    assert all(
        b[0] - a[1] == datetime(2000, 1, 2) - datetime(2000, 1, 1)
        for a, b in zip(semanas, semanas[1:], strict=False)
    )


@pytest.mark.parametrize(
    ("fecha", "semana"),
    [
        (_d(1, 12), "Semana 1"),
        (_d(1, 13), "Semana 2"),
        (_d(1, 31), "Semana 4"),
        (_d(1, 31, hora=10), ""),
        (_d(6, 8), "Semana 1"),
        (_d(9, 8), "Semana 2"),
        (None, ""),
    ],
)
def test_calcular_semana(fecha: datetime | None, semana: str) -> None:
    assert calcular_semana(fecha) == semana


@pytest.mark.parametrize(
    ("concepto", "categoria"),
    [
        ("ITF COMISION", "ITF"),  # la primera regla gana
        ("COMIS.RECAUDACION", "COMISIONES / MANT Y PORTES"),
        ("com. mantenimiento", "COMISIONES / MANT Y PORTES"),
        ("ENVIO DE EXTRACTO", "COMISIONES / MANT Y PORTES"),
        ("ABONO REMESAS", "MONEDAS"),  # antes que ABONO
        ("N/D PROSEGUR", "MONEDAS"),
        ("APERTURA PLAZO 123", "TRANSFERENCIAS"),
        ("SPOT MER", "TRANSFERENCIAS"),
        ("CASH PROVEEDOR", "PAGO PROVEEDORES (PEN)"),
        ("PAGO C/PH", "PAGO PROVEEDORES (PEN)"),
        ("CASH DEFINIR", "PLANILLA"),
        ("AFPNET", "PLANILLA"),
        ("ABONO RECEP. TRANSF.", "ABONO"),
        ("FACTURA", "ABONO"),
        ("DEPOSITO", ""),
        (None, ""),
        (123, ""),
    ],
)
def test_categorizar(concepto: object, categoria: str) -> None:
    assert categorizar(concepto) == categoria


@pytest.mark.parametrize(
    ("valor", "esperado"),
    [
        ("-$1,234.56", (-1234.56, False)),
        ("($1,234.56)", (-1234.56, False)),
        ("1,234.56", (1234.56, False)),
        ("$ 1 000", (1000.0, False)),
        (5, (5.0, False)),
        (-2.5, (-2.5, False)),
        (None, (0.0, False)),
        ("", (0.0, False)),
        ("\xa0", (0.0, False)),
        ("None", (0.0, False)),
        ("abc", (0.0, True)),
        ("1.2.3", (0.0, True)),
    ],
)
def test_limpiar_valor_monetario(valor: object, esperado: tuple[float, bool]) -> None:
    assert limpiar_valor_monetario(valor) == esperado


def test_el_orden_es_el_quicksort_de_pandas_y_no_uno_estable() -> None:
    """Vector fijado con `pandas.Series.sort_values` (pandas 3.0.6, numpy 2.4.2).

    Cuarenta fechas con tres valores distintos y dos nulas: los empates salen
    en el orden del introsort de numpy, que NO es el orden de entrada.
    """
    fechas = [None if i in (5, 20) else _d(1, 1 + (i * 7) % 3) for i in range(40)]
    esperado = [0, 36, 33, 30, 27, 24, 21, 18, 15, 12, 39, 9, 3, 6, 10, 37, 1, 34, 31, 28]
    esperado += [4, 25, 19, 22, 7, 16, 13, 11, 38, 29, 8, 17, 32, 2, 35, 14, 23, 26, 5, 20]
    assert _orden_de_pandas(fechas) == esperado


# --- Errores -------------------------------------------------------------------


def test_un_encabezado_corrido_es_columna_faltante(
    procesador: FlujoCajaIngresos, fabrica: _Fabrica
) -> None:
    """Si el banco movió una columna, el script leía otra cosa; acá se detiene."""
    config = _CONFIGS["IBK 1106 FID"]
    filas = _hoja(config, [(_d(1, 2), "A", 1)])
    encabezado = filas[config.fila_encabezado]
    encabezado[7], encabezado[8] = None, "Abonos"  # corrido una columna
    datos = _fideicomisas()
    del datos["IBK 1106 FID"]
    entrada = fabrica.entrada(datos, nombre="f.xlsx", crudas={"IBK 1106 FID": filas})
    esperado = {
        "archivo": "f.xlsx",
        "motivo": "columna_faltante",
        "columna": "IBK 1106 FID: Abonos",
    }
    assert _error(procesador, entrada).contexto == esperado
    with pytest.raises(ErrorContenido) as capturado:
        procesador.procesar([entrada])
    assert capturado.value.contexto == esperado


def test_una_hoja_sin_encabezado_es_columna_faltante(
    procesador: FlujoCajaIngresos, fabrica: _Fabrica
) -> None:
    entrada = fabrica.entrada(
        {"BCO NACION": [("2025.01.02", None, 1)]}, crudas={"BBVA 43332 PEN": [["corto"]]}
    )
    assert _error(procesador, entrada).contexto == {
        "archivo": "libro.xlsx",
        "motivo": "columna_faltante",
        "columna": "BBVA 43332 PEN: F. Operación",
    }


def test_un_libro_sin_filas_es_cero_filas_y_no_escribe_nada(
    procesador: FlujoCajaIngresos, fabrica: _Fabrica, tmp_path: Path
) -> None:
    """No hay éxito parcial: el consolidado del otro libro tampoco queda escrito."""
    bueno = fabrica.entrada(_operativas(), nombre="ope.xlsx")
    vacio = fabrica.entrada(
        {c.nombre: [(_d(1, 2), "negativo", -1)] for c in HOJAS_FIDEICOMISO}, nombre="fid.xlsx"
    )
    with pytest.raises(ErrorContenido) as capturado:
        procesador.procesar([bueno, vacio])
    assert capturado.value.contexto == {
        "archivo": "fid.xlsx",
        "motivo": "cero_filas",
        "columna": None,
    }
    assert sorted(p.name for p in tmp_path.iterdir()) == ["entrada_0", "entrada_1"]


def test_un_excel_ilegible_es_cero_filas(procesador: FlujoCajaIngresos, tmp_path: Path) -> None:
    ruta = tmp_path / "entrada_0"
    ruta.write_bytes(b"esto no es un xlsx")
    entrada = ArchivoEntrada(
        nombre_original="corrupto.xlsx", ruta_temporal=ruta, tamano_comprimido=18, formato="xlsx"
    )
    assert _error(procesador, entrada).contexto["motivo"] == "cero_filas"


# --- Observaciones -------------------------------------------------------------


def test_entradas_completas_no_dejan_observaciones(
    procesador: FlujoCajaIngresos, fabrica: _Fabrica, caplog: pytest.LogCaptureFixture
) -> None:
    with caplog.at_level("INFO", logger="app"):
        salidas = _procesar(
            procesador, fabrica.entrada(_fideicomisas()), fabrica.entrada(_operativas())
        )
    assert NOMBRE_OBSERVACIONES not in salidas
    assert not [r for r in caplog.records if getattr(r, "evento", None) == "descartes"]


def test_hoja_ausente_y_hoja_sin_filas_se_reportan_con_crlf(
    procesador: FlujoCajaIngresos, fabrica: _Fabrica, caplog: pytest.LogCaptureFixture
) -> None:
    datos = _fideicomisas(**{"BBVA FID 44606": [(_d(1, 2), "solo negativos", -3)]})
    del datos["SANTANDER 470 FID"]
    with caplog.at_level("INFO", logger="app"):
        salidas = _procesar(procesador, fabrica.entrada(datos, nombre="fid.xlsx"))

    crudo = salidas[NOMBRE_OBSERVACIONES].read_bytes()
    assert crudo.endswith(b"\r\n")
    assert b"\n" not in crudo.replace(b"\r\n", b"")
    lineas = crudo.decode("utf-8").split(FIN_DE_LINEA)
    assert lineas[0] == f"Observaciones: 2 ({MOTIVO_HOJA_AUSENTE}=1, {MOTIVO_HOJA_SIN_FILAS}=1)"
    assert lineas[2] == "archivo\thoja\tfila\tmotivo\tvalor"
    assert lineas[3:5] == [
        # Las ausentes se anotan antes de recorrer las presentes.
        f"fid.xlsx\tSANTANDER 470 FID\t\t{MOTIVO_HOJA_AUSENTE}\t",
        f"fid.xlsx\tBBVA FID 44606\t\t{MOTIVO_HOJA_SIN_FILAS}\t",
    ]
    (registro,) = [r for r in caplog.records if getattr(r, "evento", None) == "descartes"]
    assert registro.clave == CLAVE  # type: ignore[attr-defined]
    assert registro.archivo == "fid.xlsx"  # type: ignore[attr-defined]
    assert registro.filas_procesadas == 4  # type: ignore[attr-defined]


# --- Coherencia con el registro ----------------------------------------------


def test_la_clave_coincide_con_el_registry() -> None:
    from app.core.contrato import obtener_contrato
    from app.registry import REGISTRY

    procesador = REGISTRY[CLAVE]
    assert procesador.clave == CLAVE == "flujo-caja-ingresos"
    assert isinstance(procesador, FlujoCajaIngresos)

    contrato = obtener_contrato(CLAVE)
    assert contrato is not None
    assert (contrato.entradas_min, contrato.entradas_max) == (1, 2)
    assert contrato.formatos_aceptados == ("xlsx",)


# --- De punta a punta, por la ruta real y el proceso hijo -------------------


class TestPuntaAPunta:
    """La tubería completa: HTTP → temporales sin extensión → `spawn` → ZIP."""

    @staticmethod
    def _subir(cliente: Any, archivos: list[tuple[str, bytes]]) -> Any:
        return cliente.post(
            f"/interno/procesadores/{CLAVE}",
            files=[
                ("archivos", (nombre, contenido, "application/octet-stream"))
                for nombre, contenido in archivos
            ],
            headers={"Authorization": "Bearer sentinela-token-de-pruebas-3f9c2a"},
        )

    def test_los_dos_libros_vuelven_como_zip_con_los_dos_excel(self, token_sentinela: str) -> None:
        from fastapi.testclient import TestClient

        from app.main import crear_app

        fid = _bytes_de_libro(_fideicomisas())
        ope = _bytes_de_libro(_operativas())
        with TestClient(crear_app()) as cliente:
            respuesta = self._subir(
                cliente,
                [("Mov Ctas Fideicomisas.xlsx", fid), ("Mov Ctas Operativas.xlsx", ope)],
            )

        assert respuesta.status_code == 200, respuesta.text
        with zipfile.ZipFile(io.BytesIO(respuesta.content)) as paquete:
            assert sorted(paquete.namelist()) == [NOMBRE_BANCO_NACION, NOMBRE_CONSOLIDADO]
            libro = load_workbook(io.BytesIO(paquete.read(NOMBRE_CONSOLIDADO)))
            assert libro.sheetnames == ["Consolidado Fideicomiso", "Consolidado Operativas"]

    def test_dos_libros_del_mismo_tipo_es_422(self, token_sentinela: str) -> None:
        from fastapi.testclient import TestClient

        from app.main import crear_app

        ope = _bytes_de_libro(_operativas())
        with TestClient(crear_app()) as cliente:
            respuesta = self._subir(cliente, [("a.xlsx", ope), ("b.xlsx", ope)])

        assert respuesta.status_code == 422
        assert respuesta.json()["contexto"] == {
            "archivo": "b.xlsx",
            "motivo": "tipo_duplicado",
            "columna": None,
        }
