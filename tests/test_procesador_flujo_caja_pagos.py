"""Pruebas del procesador `flujo-caja-pagos`.

Los libros reales no están en este repositorio y no pueden estarlo (ADR 0015),
así que estas pruebas arman libros sintéticos con `openpyxl` en el layout de
las muestras reales —encabezado en la primera fila, 28 columnas— y los
convierten en libros con macros (`_como_xlsm`: el mismo paquete que un `.xlsm`
de Excel, con `vbaProject.bin` y el tipo de contenido `macroEnabled`). Se
guardan con nombres SIN extensión (`entrada_0`) como los que escribe
`app/recepcion.py`. Prueban las **reglas**; la paridad contra la salida
histórica es `tests/paridad/test_flujo_caja_pagos.py`.
"""

from __future__ import annotations

import io
import zipfile
from collections.abc import Mapping, Sequence
from datetime import datetime
from pathlib import Path
from typing import Any

import pytest
from openpyxl import Workbook, load_workbook

from app.core.contrato import obtener_contrato
from app.core.errores import ErrorContenido, ErrorFormato
from app.core.tipos import ArchivoEntrada
from app.core.validaciones import validar_formato, validar_tamano_descomprimido
from app.procesadores.flujo_caja_pagos.modulo import (
    CABECERAS_ESTANDAR,
    CLAVE,
    COLUMNAS_GENERADAS,
    FIN_DE_LINEA,
    MIME_XLSX,
    MOTIVO_FECHA,
    MOTIVO_HOJA_AUSENTE,
    MOTIVO_IMPORTE,
    MOTIVO_MONEDA,
    MOTIVO_MONEDA_AUSENTE,
    NOMBRE_CONSOLIDADO,
    NOMBRE_OBSERVACIONES,
    FlujoCajaPagos,
    _clave_ntfs,
    adivinar_formato,
    aplicar_signo,
    calcular_semana,
    categorizar,
    limpiar_importe,
)
from app.recepcion import _formato

Registro = Mapping[str, object]
Hojas = Mapping[str, Sequence[Registro | None]]

_PREDETERMINADA: dict[str, object] = {
    "Ejercicio / mes": "2025/07",
    "Cuenta": 3000019305,
    "Fecha de documento": "18.06.2025",
    "Período contable": 7,
    "Base p plazo pago": "25.06.2025",
    "Nombre 1": "PROVEEDOR",
    "Clase de documento": 1,
    "Referencia": "01-F001-1",
    "Moneda del documento": "PEN",
    "Importe en moneda local": "1,000.00-",
    "Importe en ML3": "250.00-",
    "Texto": "SERVICIO",
    "Orden": None,
    "Fecha de pago BANCOS": datetime(2025, 7, 7),
}
"""Una fila válida con el aspecto de las reales: los importes como texto con
`-` final, `Cuenta` como número y la fecha de pago como fecha. El 7 de julio
de 2025 cae en la `Semana 2` y está a 12 días de la base."""

_LIMA_EXPRESA: tuple[str, ...] = (
    "SOLES",
    "CAJA CHICA LE",
    "REEMBOLSOS_SOL",
    "DOLARES",
    "REEMBOLSOS_DOL",
    "PLANILLA",
    "PLATAFORMA IBK",
    "PLATAFORMA BBVA",
)
_PEX: tuple[str, ...] = ("PEX SOL", "PEX DOL", "PLATAFORMA BBVA", "PLATAFORMA BCP")
_DOLARES: frozenset[str] = frozenset({"DOLARES", "REEMBOLSOS_DOL", "PEX DOL"})


def _fila(**cambios: object) -> dict[str, object]:
    """Una fila; `cambios` usa los títulos con `_` en lugar de espacios."""
    fila = dict(_PREDETERMINADA)
    for clave, valor in cambios.items():
        titulo = next(t for t in CABECERAS_ESTANDAR if t.replace(" ", "_") == clave)
        fila[titulo] = valor
    return fila


def _como_xlsm(contenido: bytes) -> bytes:
    """El paquete de un libro con macros: `vbaProject.bin`, su relación y el
    tipo de contenido `macroEnabled` del libro, como en las muestras reales."""
    origen = zipfile.ZipFile(io.BytesIO(contenido))
    destino = io.BytesIO()
    with zipfile.ZipFile(destino, "w", zipfile.ZIP_DEFLATED) as paquete:
        for info in origen.infolist():
            datos = origen.read(info.filename)
            if info.filename == "[Content_Types].xml":
                datos = datos.replace(
                    b"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml",
                    b"application/vnd.ms-excel.sheet.macroEnabled.main+xml",
                ).replace(
                    b"</Types>",
                    b'<Override PartName="/xl/vbaProject.bin" '
                    b'ContentType="application/vnd.ms-office.vbaProject" /></Types>',
                )
            elif info.filename == "xl/_rels/workbook.xml.rels":
                datos = datos.replace(
                    b"</Relationships>",
                    b'<Relationship Type="http://schemas.microsoft.com/office/2006/'
                    b'relationships/vbaProject" Target="vbaProject.bin" Id="rIdVba" />'
                    b"</Relationships>",
                )
            paquete.writestr(info, datos)
        paquete.writestr("xl/vbaProject.bin", b"\xd0\xcf\x11\xe0" + b"\x00" * 60)
    return destino.getvalue()


def _bytes_de_libro(
    hojas: Hojas,
    *,
    encabezado: Sequence[object] = CABECERAS_ESTANDAR,
    crudas: Mapping[str, Sequence[Sequence[object]]] | None = None,
) -> bytes:
    """Cada hoja con el encabezado en la primera fila y una fila por registro
    (`None` es una fila vacía). `crudas` se escribe tal cual."""
    libro = Workbook()
    libro.remove(libro["Sheet"])
    for nombre, registros in hojas.items():
        hoja = libro.create_sheet(nombre)
        hoja.append(list(encabezado))
        for registro in registros:
            hoja.append([] if registro is None else [registro.get(t) for t in CABECERAS_ESTANDAR])
    for nombre, filas in (crudas or {}).items():
        hoja = libro.create_sheet(nombre)
        for fila in filas:
            hoja.append(list(fila))
    memoria = io.BytesIO()
    libro.save(memoria)
    return _como_xlsm(memoria.getvalue())


def _lima_expresa(**reemplazos: Sequence[Registro | None]) -> dict[str, Sequence[Registro | None]]:
    """Las ocho hojas de LIMA EXPRESA con una fila cada una; las "por moneda"
    en soles. `reemplazos` por nombre de hoja, con `_` en lugar de espacios."""
    hojas: dict[str, Sequence[Registro | None]] = {
        nombre: [_fila(Moneda_del_documento="USD" if nombre in _DOLARES else "PEN")]
        for nombre in _LIMA_EXPRESA
    }
    for clave, registros in reemplazos.items():
        hojas[clave.replace("_", " ") if clave.replace("_", " ") in hojas else clave] = registros
    return hojas


def _pex() -> dict[str, Sequence[Registro | None]]:
    return {
        nombre: [_fila(Moneda_del_documento="USD" if nombre in _DOLARES else "PEN")]
        for nombre in _PEX
    }


class _Fabrica:
    """Escribe libros como `entrada_N`, sin extensión, igual que la recepción real."""

    def __init__(self, directorio: Path) -> None:
        self.directorio = directorio
        self.contador = 0

    def entrada(self, contenido: bytes | Hojas, nombre: str = "libro.xlsm") -> ArchivoEntrada:
        ruta = self.directorio / f"entrada_{self.contador}"
        self.contador += 1
        ruta.write_bytes(contenido if isinstance(contenido, bytes) else _bytes_de_libro(contenido))
        return ArchivoEntrada(
            nombre_original=nombre,
            ruta_temporal=ruta,
            tamano_comprimido=ruta.stat().st_size,
            formato=_formato(nombre),
        )


@pytest.fixture
def fabrica(tmp_path: Path) -> _Fabrica:
    return _Fabrica(tmp_path)


@pytest.fixture
def procesador() -> FlujoCajaPagos:
    return FlujoCajaPagos()


def _procesar(procesador: FlujoCajaPagos, *entradas: ArchivoEntrada) -> dict[str, Path]:
    assert procesador.validar(list(entradas)) is None
    return {s.nombre_propuesto: s.ruta_temporal for s in procesador.procesar(list(entradas))}


def _leer(ruta: Path) -> dict[str, list[dict[str, Any]]]:
    """Cada hoja como lista de filas por título (la primera fila son los títulos)."""
    libro = load_workbook(ruta)
    try:
        hojas: dict[str, list[dict[str, Any]]] = {}
        for hoja in libro.worksheets:
            filas: list[list[Any]] = [list(f) for f in hoja.iter_rows(values_only=True)]
            hojas[hoja.title] = [dict(zip(filas[0], fila, strict=True)) for fila in filas[1:]]
        return hojas
    finally:
        libro.close()


def _titulos(ruta: Path, hoja: str) -> list[Any]:
    libro = load_workbook(ruta, read_only=True)
    try:
        return list(next(libro[hoja].iter_rows(values_only=True)))
    finally:
        libro.close()


def _error(procesador: FlujoCajaPagos, *entradas: ArchivoEntrada) -> ErrorContenido:
    error = procesador.validar(list(entradas))
    assert isinstance(error, ErrorContenido)
    return error


def _observaciones(ruta: Path) -> list[list[str]]:
    """Las líneas de detalle de `observaciones.txt`, partidas por tabulador."""
    lineas = ruta.read_bytes().decode("utf-8").split(FIN_DE_LINEA)
    return [linea.split("\t") for linea in lineas[3:] if linea]


# --- El formato .xlsm en la validación de contrato ------------------------------


def test_el_libro_sintetico_es_un_xlsm_de_verdad(fabrica: _Fabrica) -> None:
    entrada = fabrica.entrada(_lima_expresa())
    with zipfile.ZipFile(entrada.ruta_temporal) as paquete:
        assert "xl/vbaProject.bin" in paquete.namelist()
        tipos = paquete.read("[Content_Types].xml")
        assert b'"application/vnd.ms-excel.sheet.macroEnabled.main+xml"' in tipos


def test_el_contrato_acepta_xlsm_y_xlsx_sin_importar_mayusculas(fabrica: _Fabrica) -> None:
    contrato = obtener_contrato(CLAVE)
    assert contrato is not None
    for nombre in ("Pagos LIMA EXPRESA.xlsm", "Pagos PEX.XLSM", "Pagos PEX.xlsx"):
        validar_formato(nombre_original=nombre, formato=_formato(nombre), contrato=contrato)
    with pytest.raises(ErrorFormato):
        validar_formato(nombre_original="macro.xls", formato="xls", contrato=contrato)
    # La inspección del directorio central del ZIP también lo acepta.
    validar_tamano_descomprimido(fabrica.entrada(_lima_expresa()))


def test_un_xlsm_real_se_lee_por_descriptor_sin_extension(
    procesador: FlujoCajaPagos, fabrica: _Fabrica
) -> None:
    """La ruta temporal no tiene extensión: `load_workbook` recibe el descriptor."""
    entrada = fabrica.entrada(_lima_expresa(), nombre="01.-Pagos LIMA EXPRESA.xlsm")
    assert entrada.ruta_temporal.suffix == ""
    salidas = _procesar(procesador, entrada)
    assert list(salidas) == [NOMBRE_CONSOLIDADO]


# --- Detección del tipo y orden de los libros ---------------------------------


def test_el_tipo_se_detecta_por_las_hojas_y_no_por_el_nombre(
    procesador: FlujoCajaPagos, fabrica: _Fabrica
) -> None:
    salidas = _procesar(
        procesador,
        fabrica.entrada(_lima_expresa(), nombre="PEX engañoso.xlsm"),
        fabrica.entrada(_pex(), nombre="LIMA EXPRESA engañoso.xlsm"),
    )
    soles = _leer(salidas[NOMBRE_CONSOLIDADO])["Soles"]
    por_excel = {(f["Excel"], f["Sociedad"]) for f in soles}
    assert por_excel == {
        ("PEX engañoso.xlsm", "LIMA EXPRESA"),
        ("LIMA EXPRESA engañoso.xlsm", "PEX"),
    }


def test_plataforma_bcp_y_bbva_no_deciden_el_tipo(
    procesador: FlujoCajaPagos, fabrica: _Fabrica
) -> None:
    """Los libros LIMA EXPRESA reales también traen `PLATAFORMA BCP`."""
    hojas = _lima_expresa()
    hojas = {**hojas, "PLATAFORMA BCP": [_fila()]}
    salidas = _procesar(procesador, fabrica.entrada(hojas, nombre="semana.xlsm"))
    fuentes = {f["Fuente"] for f in _leer(salidas[NOMBRE_CONSOLIDADO])["Soles"]}
    assert "PLATAFORMA BCP" not in fuentes
    assert {f["Sociedad"] for f in _leer(salidas[NOMBRE_CONSOLIDADO])["Soles"]} == {"LIMA EXPRESA"}


def test_sin_hojas_distintivas_decide_el_nombre(
    procesador: FlujoCajaPagos, fabrica: _Fabrica
) -> None:
    entrada = fabrica.entrada({"PLATAFORMA BBVA": [_fila()]}, nombre="Pagos pex semana 3.xlsm")
    soles = _leer(_procesar(procesador, entrada)[NOMBRE_CONSOLIDADO])["Soles"]
    assert [(f["Sociedad"], f["Fuente"]) for f in soles] == [("PEX", "PLATAFORMA BBVA")]


@pytest.mark.parametrize(
    "nombre", ["cualquiera.xlsm", "LIMA EXPRESA y PEX.xlsm", "carpeta/PEX/otro.xlsm"]
)
def test_sin_hojas_distintivas_ni_nombre_util_es_tipo_no_reconocido(
    procesador: FlujoCajaPagos, fabrica: _Fabrica, nombre: str
) -> None:
    """El script procesaba un nombre con las dos subcadenas como PEX en la
    posición de los LIMA EXPRESA; acá no se adivina. Sólo cuenta el nombre
    base, no las carpetas."""
    entrada = fabrica.entrada({"PLATAFORMA BBVA": [_fila()]}, nombre=nombre)
    assert _error(procesador, entrada).contexto == {
        "archivo": nombre,
        "motivo": "tipo_no_reconocido",
        "columna": None,
    }


def test_hojas_de_los_dos_tipos_es_tipo_no_reconocido(
    procesador: FlujoCajaPagos, fabrica: _Fabrica
) -> None:
    entrada = fabrica.entrada({"SOLES": [_fila()], "PEX DOL": [_fila()]})
    assert _error(procesador, entrada).contexto["motivo"] == "tipo_no_reconocido"


def test_tipo_por_nombre_sin_ninguna_hoja_configurada_es_hoja_faltante(
    procesador: FlujoCajaPagos, fabrica: _Fabrica
) -> None:
    entrada = fabrica.entrada({"OTRA": [_fila()]}, nombre="LIMA EXPRESA.xlsm")
    assert _error(procesador, entrada).contexto == {
        "archivo": "LIMA EXPRESA.xlsm",
        "motivo": "hoja_faltante",
        "columna": "SOLES",
    }


def test_los_libros_van_lima_expresa_primero_y_en_el_orden_de_ntfs(
    procesador: FlujoCajaPagos, fabrica: _Fabrica
) -> None:
    """Varios del mismo tipo son lo normal. El orden de subida no importa: el
    de NTFS compara en mayúsculas, así que `_` (0x5F) va después de las letras
    y `b` antes que `C`."""
    subidos = [
        ("a PEX.xlsm", _pex()),
        ("_LIMA EXPRESA.xlsm", _lima_expresa()),
        ("Cuarta LIMA EXPRESA.xlsm", _lima_expresa()),
        ("b LIMA EXPRESA.xlsm", _lima_expresa()),
        ("A lima expresa.xlsm", _lima_expresa()),
    ]
    salidas = _procesar(procesador, *(fabrica.entrada(h, nombre=n) for n, h in subidos))
    soles = _leer(salidas[NOMBRE_CONSOLIDADO])["Soles"]
    orden = list(dict.fromkeys(f["Excel"] for f in soles))
    assert orden == [
        "A lima expresa.xlsm",
        "b LIMA EXPRESA.xlsm",
        "Cuarta LIMA EXPRESA.xlsm",
        "_LIMA EXPRESA.xlsm",
        "a PEX.xlsm",
    ]


def test_la_clave_ntfs_compara_en_mayusculas_por_unidades_utf16() -> None:
    nombres = ["_x", "b", "A", "ñ", "Z", "é"]
    assert sorted(nombres, key=_clave_ntfs) == ["A", "b", "Z", "_x", "é", "ñ"]


# --- Reglas de negocio --------------------------------------------------------


@pytest.mark.parametrize(
    ("orden", "categoria"),
    [
        (None, "OPEX"),
        ("  ", "OPEX"),
        ("404587", "OPEX"),
        ("203882", "CAPEX"),
        (" 3001", "CAPEX"),
        ("5", ""),
        ("404587.5", "OPEX"),
    ],
)
def test_categorizar(orden: str | None, categoria: str) -> None:
    assert categorizar(orden) == categoria


@pytest.mark.parametrize(
    ("texto", "importe"),
    [
        ("12,500.00-", 12500.0),
        ("-5", -5),
        ("1,000", 1000),
        ("42060.64", 42060.64),
        ("1e3", 1000.0),
        ("5--", None),
        ("abc", None),
        (None, None),
    ],
)
def test_limpiar_importe_quita_un_guion_final_y_las_comas(
    texto: str | None, importe: float | None
) -> None:
    assert limpiar_importe(texto) == importe


@pytest.mark.parametrize(
    ("importe", "clase", "esperado"),
    [
        (100, "7", -100),
        (-100, "7 ", -100),
        (-5, "1", 5),
        (5.5, None, 5.5),
        (None, "7", None),
    ],
)
def test_el_signo_lo_decide_solo_la_clase_7(
    importe: float | None, clase: str | None, esperado: float | None
) -> None:
    assert aplicar_signo(importe, clase) == esperado


@pytest.mark.parametrize(
    ("texto", "dayfirst", "formato"),
    [
        ("13-06-2025", True, "%d-%m-%Y"),
        ("05-06-2025", True, "%d-%m-%Y"),
        ("05-06-2025", False, "%m-%d-%Y"),
        ("13.07.2025", False, "%d.%m.%Y"),
        ("2025-07-02 00:00:00", False, "%Y-%m-%d %H:%M:%S"),
        # dateutil con `dayfirst` lee el año primero como año/día/mes.
        ("2025-06-01", True, "%Y-%d-%m"),
        ("2025-06-13", True, "%Y-%m-%d"),
        (" 13-06-2025", True, " %d-%m-%Y"),
        ("01-07-25", True, None),
        ("31-02-2025", True, None),
        ("45839", False, None),
    ],
)
def test_adivinar_formato_como_pandas(texto: str, dayfirst: bool, formato: str | None) -> None:
    assert adivinar_formato(texto, dayfirst=dayfirst) == formato


def test_reglas_de_una_fila_de_soles(procesador: FlujoCajaPagos, fabrica: _Fabrica) -> None:
    hojas = _lima_expresa(
        SOLES=[
            _fila(Cuenta=3000019305, Orden=203882, Clase_de_documento=7),
            _fila(Base_p_plazo_pago="05.06.2025", Importe_en_moneda_local="-5"),
        ]
    )
    salidas = _procesar(procesador, fabrica.entrada(hojas, nombre="Semana LIMA EXPRESA.xlsm"))
    soles = _leer(salidas[NOMBRE_CONSOLIDADO])["Soles"]
    primera, segunda = soles[0], soles[1]
    # `dtype=str`: los números de las columnas no convertidas salen como texto.
    assert primera["Cuenta"] == "3000019305"
    assert primera["Período contable"] == "7"
    assert primera["Clase de documento"] == "7"
    assert primera["Imptebase descuento"] is None
    assert primera["Importe en moneda local"] == -1000
    assert primera["Importe en ML3"] == -250
    assert primera["Categoria"] == "CAPEX"
    assert primera["Sociedad"] == "LIMA EXPRESA"
    assert primera["Excel"] == "Semana LIMA EXPRESA.xlsm"
    assert primera["Fuente"] == "SOLES"
    assert primera["Fecha de pago BANCOS"] == datetime(2025, 7, 7)
    assert (primera["PPP"], primera["Mes"], primera["Semana"]) == (12, "Jul", "Semana 2")
    # Base con el día primero: 5 de junio, 32 días antes del 7 de julio.
    assert segunda["PPP"] == 32
    assert segunda["Importe en moneda local"] == 5
    assert segunda["Categoria"] == "OPEX"


def test_titulos_y_orden_de_columnas(procesador: FlujoCajaPagos, fabrica: _Fabrica) -> None:
    salidas = _procesar(procesador, fabrica.entrada(_lima_expresa()))
    esperados = [*CABECERAS_ESTANDAR, *COLUMNAS_GENERADAS]
    assert _titulos(salidas[NOMBRE_CONSOLIDADO], "Soles") == esperados
    assert _titulos(salidas[NOMBRE_CONSOLIDADO], "Dolares") == esperados


def test_filas_sin_cuenta_y_primera_fila_se_descartan(
    procesador: FlujoCajaPagos, fabrica: _Fabrica
) -> None:
    hojas = _lima_expresa(
        SOLES=[
            None,
            _fila(Cuenta=None, Texto="subtotal"),
            _fila(Cuenta="  ", Texto="espacios"),
            _fila(Cuenta="NA", Texto="nulo de pandas"),
            _fila(Texto="buena"),
        ]
    )
    salidas = _procesar(procesador, fabrica.entrada(hojas))
    soles = _leer(salidas[NOMBRE_CONSOLIDADO])["Soles"]
    assert [f["Texto"] for f in soles if f["Fuente"] == "SOLES"] == ["buena"]


def test_hojas_por_moneda_se_reparten_y_la_moneda_ajena_se_reporta(
    procesador: FlujoCajaPagos, fabrica: _Fabrica
) -> None:
    hojas = _lima_expresa(
        PLANILLA=[
            _fila(Moneda_del_documento="USD", Texto="dolar"),
            _fila(Moneda_del_documento="PEN", Texto="sol"),
            _fila(Moneda_del_documento="EUR", Texto="euro"),
            _fila(Moneda_del_documento=" PEN", Texto="sol con espacio"),
        ]
    )
    salidas = _procesar(procesador, fabrica.entrada(hojas, nombre="le.xlsm"))
    consolidado = _leer(salidas[NOMBRE_CONSOLIDADO])
    assert [f["Texto"] for f in consolidado["Soles"] if f["Fuente"] == "PLANILLA"] == ["sol"]
    assert [f["Texto"] for f in consolidado["Dolares"] if f["Fuente"] == "PLANILLA"] == ["dolar"]
    assert _observaciones(salidas[NOMBRE_OBSERVACIONES]) == [
        ["le.xlsm", "PLANILLA", "4", MOTIVO_MONEDA, "EUR"],
        ["le.xlsm", "PLANILLA", "5", MOTIVO_MONEDA, " PEN"],
    ]


def test_el_orden_de_filas_es_el_de_concatenacion(
    procesador: FlujoCajaPagos, fabrica: _Fabrica
) -> None:
    """Por libro: Soles, Dolares y las "por moneda" repartidas en su lugar."""
    salidas = _procesar(procesador, fabrica.entrada(_lima_expresa()))
    consolidado = _leer(salidas[NOMBRE_CONSOLIDADO])
    assert [f["Fuente"] for f in consolidado["Soles"]] == [
        "SOLES",
        "CAJA CHICA LE",
        "REEMBOLSOS_SOL",
        "PLANILLA",
        "PLATAFORMA IBK",
        "PLATAFORMA BBVA",
    ]
    assert [f["Fuente"] for f in consolidado["Dolares"]] == ["DOLARES", "REEMBOLSOS_DOL"]


def test_totales_por_semana_sin_excel_en_la_agrupacion(
    procesador: FlujoCajaPagos, fabrica: _Fabrica
) -> None:
    julio_14 = datetime(2025, 7, 14)  # Semana 3
    uno = _lima_expresa(
        SOLES=[
            _fila(Importe_en_moneda_local="100.00-"),
            _fila(Importe_en_moneda_local="30.00-", Clase_de_documento=7),
            _fila(Importe_en_moneda_local="1.00-", Orden=203882),
            _fila(Importe_en_moneda_local="9.00-", Fecha_de_pago_BANCOS=julio_14),
        ],
        CAJA_CHICA_LE=[_fila(Importe_en_moneda_local="1,000.00-")],
    )
    dos = _lima_expresa(SOLES=[_fila(Importe_en_moneda_local="0.5")])
    salidas = _procesar(
        procesador, fabrica.entrada(uno, nombre="1 LE.xlsm"), fabrica.entrada(dos, nombre="2.xlsm")
    )
    soles = _leer(salidas[NOMBRE_CONSOLIDADO])["Soles"]
    totales = [
        (
            f["Excel"][0],
            f["Fuente"],
            f["Total Hoja por Semana"],
            f["Total Mes por Semana"],
            f["Total Categoria por Semana"],
        )
        for f in soles
        if f["Fuente"] in ("SOLES", "CAJA CHICA LE")
    ]
    # Semana 2 de SOLES: 100 - 30 + 1 (uno) + 0.5 (dos) = 71.5, sin separar
    # por libro. Mes-semana suma además, de los dos libros, CAJA CHICA LE,
    # REEMBOLSOS_SOL y las tres "por moneda" en soles (5 x 2 x 1000). OPEX de
    # la semana 2: todo menos el 1 de CAPEX.
    mes = 71.5 + 5 * 2000
    assert totales == [
        ("1", "SOLES", 71.5, mes, mes - 1),
        ("1", "SOLES", 71.5, mes, mes - 1),
        ("1", "SOLES", 71.5, mes, 1),
        ("1", "SOLES", 9, 9, 9),
        ("1", "CAJA CHICA LE", 2000, mes, mes - 1),
        ("2", "SOLES", 71.5, mes, mes - 1),
        ("2", "CAJA CHICA LE", 2000, mes, mes - 1),
    ]


def test_sin_fecha_de_pago_quedan_vacios_ppp_mes_y_semana(
    procesador: FlujoCajaPagos, fabrica: _Fabrica
) -> None:
    hojas = _lima_expresa(
        SOLES=[_fila(Fecha_de_pago_BANCOS=None), _fila(Fecha_de_pago_BANCOS="pendiente")]
    )
    salidas = _procesar(procesador, fabrica.entrada(hojas, nombre="le.xlsm"))
    soles = [f for f in _leer(salidas[NOMBRE_CONSOLIDADO])["Soles"] if f["Fuente"] == "SOLES"]
    assert [(f["PPP"], f["Mes"], f["Semana"]) for f in soles] == [(None, None, None)] * 2
    # La vacía no se reporta; la escrita e ilegible, sí.
    assert _observaciones(salidas[NOMBRE_OBSERVACIONES]) == [
        ["le.xlsm", "SOLES", "3", MOTIVO_FECHA, "pendiente"]
    ]


def test_fecha_de_pago_como_texto_define_el_formato_de_la_columna(
    procesador: FlujoCajaPagos, fabrica: _Fabrica
) -> None:
    """Sin `dayfirst`: `07.03.2025` es el 3 de julio, y una fecha de verdad en
    la misma columna ya no calza con ese formato."""
    hojas = _lima_expresa(
        SOLES=[
            _fila(Fecha_de_pago_BANCOS="07.03.2025"),
            _fila(Fecha_de_pago_BANCOS=datetime(2025, 7, 7)),
        ]
    )
    salidas = _procesar(procesador, fabrica.entrada(hojas, nombre="le.xlsm"))
    soles = [f for f in _leer(salidas[NOMBRE_CONSOLIDADO])["Soles"] if f["Fuente"] == "SOLES"]
    assert soles[0]["Fecha de pago BANCOS"] == datetime(2025, 7, 3)
    assert soles[0]["Semana"] == "Semana 1"
    assert soles[1]["Fecha de pago BANCOS"] is None
    assert _observaciones(salidas[NOMBRE_OBSERVACIONES]) == [
        ["le.xlsm", "SOLES", "3", MOTIVO_FECHA, "2025-07-07 00:00:00"]
    ]


@pytest.mark.parametrize(
    ("fecha", "semana"),
    [
        (datetime(2025, 7, 1), "Semana 1"),
        (datetime(2025, 7, 6), "Semana 1"),
        (datetime(2025, 7, 7), "Semana 2"),
        (datetime(2025, 7, 31), "Semana 4"),
        (datetime(2025, 7, 31, 12), ""),
        (None, ""),
    ],
)
def test_calcular_semana(fecha: datetime | None, semana: str) -> None:
    assert calcular_semana(fecha) == semana


def test_importe_ilegible_queda_vacio_y_se_reporta(
    procesador: FlujoCajaPagos, fabrica: _Fabrica
) -> None:
    hojas = _lima_expresa(SOLES=[_fila(Importe_en_moneda_local="12,500.00 PEN")])
    salidas = _procesar(procesador, fabrica.entrada(hojas, nombre="le.xlsm"))
    fila = next(f for f in _leer(salidas[NOMBRE_CONSOLIDADO])["Soles"] if f["Fuente"] == "SOLES")
    assert fila["Importe en moneda local"] is None
    assert fila["Total Hoja por Semana"] == 0
    assert _observaciones(salidas[NOMBRE_OBSERVACIONES]) == [
        ["le.xlsm", "SOLES", "2", MOTIVO_IMPORTE, "12,500.00 PEN"]
    ]


def test_sin_filas_en_dolares_sale_solo_la_hoja_soles(
    procesador: FlujoCajaPagos, fabrica: _Fabrica
) -> None:
    hojas = _lima_expresa(DOLARES=[], REEMBOLSOS_DOL=[])
    salidas = _procesar(procesador, fabrica.entrada(hojas, nombre="le.xlsm"))
    libro = load_workbook(salidas[NOMBRE_CONSOLIDADO], read_only=True)
    try:
        assert libro.sheetnames == ["Soles"]
    finally:
        libro.close()


# --- Encabezados ---------------------------------------------------------------


def test_variantes_de_puntuacion_del_encabezado_se_aceptan(
    procesador: FlujoCajaPagos, fabrica: _Fabrica
) -> None:
    """Las de las muestras reales: `Base p, plazo pago`, columnas sin rótulo."""
    encabezado: list[object] = list(CABECERAS_ESTANDAR)
    encabezado[4] = "Base p/ plazo pago"
    encabezado[15] = None
    encabezado[17] = "Vía de pago"
    contenido = _bytes_de_libro(_lima_expresa(), encabezado=encabezado)
    assert procesador.validar([fabrica.entrada(contenido)]) is None


def test_un_encabezado_corrido_es_columna_faltante(
    procesador: FlujoCajaPagos, fabrica: _Fabrica
) -> None:
    contenido = _bytes_de_libro(_lima_expresa(), encabezado=[None, *CABECERAS_ESTANDAR])
    assert _error(procesador, fabrica.entrada(contenido)).contexto == {
        "archivo": "libro.xlsm",
        "motivo": "columna_faltante",
        "columna": "SOLES: Cuenta",
    }


def test_hoja_con_datos_y_sin_encabezado_es_columna_faltante(
    procesador: FlujoCajaPagos, fabrica: _Fabrica
) -> None:
    """Con la primera fila en blanco `validar` no puede decidir; `procesar` sí."""
    hojas = _lima_expresa()
    del hojas["SOLES"]
    fila = [_PREDETERMINADA.get(t) for t in CABECERAS_ESTANDAR]
    contenido = _bytes_de_libro(hojas, crudas={"SOLES": [[], fila]})
    entrada = fabrica.entrada(contenido)
    assert procesador.validar([entrada]) is None
    with pytest.raises(ErrorContenido) as capturado:
        procesador.procesar([entrada])
    assert capturado.value.contexto["columna"] == "SOLES: Cuenta"  # type: ignore[typeddict-item]


# --- Errores y observaciones -----------------------------------------------------


def test_sin_filas_en_ninguna_hoja_es_cero_filas_y_no_escribe_nada(
    procesador: FlujoCajaPagos, fabrica: _Fabrica, tmp_path: Path
) -> None:
    vacio: Hojas = {nombre: [] for nombre in _LIMA_EXPRESA}
    entradas = [
        fabrica.entrada(vacio, nombre="uno LIMA EXPRESA.xlsm"),
        fabrica.entrada(vacio, nombre="dos LIMA EXPRESA.xlsm"),
    ]
    with pytest.raises(ErrorContenido) as capturado:
        procesador.procesar(entradas)
    assert capturado.value.contexto == {
        "archivo": "dos LIMA EXPRESA.xlsm, uno LIMA EXPRESA.xlsm",
        "motivo": "cero_filas",
        "columna": None,
    }
    assert sorted(p.name for p in tmp_path.iterdir()) == ["entrada_0", "entrada_1"]


def test_un_archivo_ilegible_es_cero_filas(procesador: FlujoCajaPagos, fabrica: _Fabrica) -> None:
    entrada = fabrica.entrada(b"esto no es un xlsm", nombre="corrupto.xlsm")
    assert _error(procesador, entrada).contexto["motivo"] == "cero_filas"


def test_entradas_completas_no_dejan_observaciones(
    procesador: FlujoCajaPagos, fabrica: _Fabrica, caplog: pytest.LogCaptureFixture
) -> None:
    with caplog.at_level("INFO", logger="app"):
        salidas = _procesar(procesador, fabrica.entrada(_lima_expresa()), fabrica.entrada(_pex()))
    assert list(salidas) == [NOMBRE_CONSOLIDADO]
    assert not [r for r in caplog.records if getattr(r, "evento", None) == "descartes"]


def test_hoja_ausente_y_sin_moneda_se_reportan_y_la_vacia_no(
    procesador: FlujoCajaPagos, fabrica: _Fabrica, caplog: pytest.LogCaptureFixture
) -> None:
    hojas = _lima_expresa(CAJA_CHICA_LE=[_fila(Cuenta=None)])
    del hojas["REEMBOLSOS_DOL"]
    del hojas["PLATAFORMA IBK"]
    angosta: list[list[object]] = [["Ejercicio / mes", "Cuenta"], ["2025/07", 3000019305]]
    contenido = _bytes_de_libro(hojas, crudas={"PLATAFORMA IBK": angosta})
    with caplog.at_level("INFO", logger="app"):
        salidas = _procesar(procesador, fabrica.entrada(contenido, nombre="le.xlsm"))

    crudo = salidas[NOMBRE_OBSERVACIONES].read_bytes()
    assert crudo.endswith(b"\r\n")
    assert b"\n" not in crudo.replace(b"\r\n", b"")
    lineas = crudo.decode("utf-8").split(FIN_DE_LINEA)
    assert lineas[0] == (f"Observaciones: 2 ({MOTIVO_HOJA_AUSENTE}=1, {MOTIVO_MONEDA_AUSENTE}=1)")
    assert lineas[2] == "archivo\thoja\tfila\tmotivo\tvalor"
    # CAJA CHICA LE no aporta filas: operación normal, no se reporta.
    assert _observaciones(salidas[NOMBRE_OBSERVACIONES]) == [
        ["le.xlsm", "REEMBOLSOS_DOL", "", MOTIVO_HOJA_AUSENTE, ""],
        ["le.xlsm", "PLATAFORMA IBK", "", MOTIVO_MONEDA_AUSENTE, ""],
    ]
    (registro,) = [r for r in caplog.records if getattr(r, "evento", None) == "descartes"]
    assert registro.clave == CLAVE  # type: ignore[attr-defined]
    assert registro.archivo == "le.xlsm"  # type: ignore[attr-defined]
    assert registro.filas_procesadas == 5  # type: ignore[attr-defined]


def test_una_hoja_angosta_no_tiene_total_de_categoria(
    procesador: FlujoCajaPagos, fabrica: _Fabrica
) -> None:
    """Sin la columna del importe, dos totales valen 0 y el tercero no existe."""
    angosta: list[list[object]] = [
        ["Ejercicio / mes", "Cuenta", "Fecha de documento"],
        ["2025/07", 1, "x"],
    ]
    contenido = _bytes_de_libro({}, crudas={"SOLES": angosta})
    salidas = _procesar(procesador, fabrica.entrada(contenido, nombre="LIMA EXPRESA.xlsm"))
    assert _titulos(salidas[NOMBRE_CONSOLIDADO], "Soles") == [
        "Ejercicio / mes",
        "Cuenta",
        "Fecha de documento",
        "Sociedad",
        "Excel",
        "Fuente",
        "Categoria",
        "PPP",
        "Mes",
        "Semana",
        "Total Hoja por Semana",
        "Total Mes por Semana",
    ]
    (fila,) = _leer(salidas[NOMBRE_CONSOLIDADO])["Soles"]
    assert (fila["Categoria"], fila["Total Hoja por Semana"], fila["Mes"]) == ("OPEX", 0, None)


# --- Coherencia con el registro ----------------------------------------------


def test_la_clave_coincide_con_el_registry() -> None:
    from app.registry import REGISTRY

    procesador = REGISTRY[CLAVE]
    assert procesador.clave == CLAVE == "flujo-caja-pagos"
    assert isinstance(procesador, FlujoCajaPagos)

    contrato = obtener_contrato(CLAVE)
    assert contrato is not None
    assert (contrato.entradas_min, contrato.entradas_max) == (1, 10)
    assert contrato.formatos_aceptados == ("xlsm", "xlsx")


# --- De punta a punta, por la ruta real y el proceso hijo -------------------


class TestPuntaAPunta:
    """La tubería completa: HTTP → temporales sin extensión → `spawn` → salida."""

    @staticmethod
    def _subir(cliente: Any, archivos: list[tuple[str, bytes]]) -> Any:
        return cliente.post(
            f"/interno/procesadores/{CLAVE}",
            files=[
                ("archivos", (nombre, contenido, "application/octet-stream"))
                for nombre, contenido in archivos
            ],
            headers={"Authorization": "Bearer sentinela-token-de-pruebas-3f9c2a"},
        )

    def test_dos_xlsm_vuelven_como_el_xlsx_sin_zip(self, token_sentinela: str) -> None:
        from fastapi.testclient import TestClient

        from app.main import crear_app

        uno = _bytes_de_libro(_lima_expresa())
        dos = _bytes_de_libro(_lima_expresa(SOLES=[_fila(Texto="segunda semana")]))
        with TestClient(crear_app()) as cliente:
            respuesta = self._subir(
                cliente,
                [
                    ("02.-Pagos Programados LIMA EXPRESA.xlsm", dos),
                    ("01.-Pagos Programados LIMA EXPRESA.xlsm", uno),
                ],
            )

        assert respuesta.status_code == 200, respuesta.text
        assert NOMBRE_CONSOLIDADO in respuesta.headers["content-disposition"]
        assert respuesta.headers["content-type"] == MIME_XLSX
        libro = load_workbook(io.BytesIO(respuesta.content))
        assert libro.sheetnames == ["Soles", "Dolares"]
        excel = [fila[29] for fila in libro["Soles"].iter_rows(min_row=2, values_only=True)]
        assert list(dict.fromkeys(excel)) == [
            "01.-Pagos Programados LIMA EXPRESA.xlsm",
            "02.-Pagos Programados LIMA EXPRESA.xlsm",
        ]

    def test_un_tipo_no_reconocido_es_422(self, token_sentinela: str) -> None:
        from fastapi.testclient import TestClient

        from app.main import crear_app

        with TestClient(crear_app()) as cliente:
            respuesta = self._subir(
                cliente, [("otro.xlsm", _bytes_de_libro({"PLATAFORMA BBVA": [_fila()]}))]
            )

        assert respuesta.status_code == 422
        assert respuesta.json()["contexto"] == {
            "archivo": "otro.xlsm",
            "motivo": "tipo_no_reconocido",
            "columna": None,
        }
