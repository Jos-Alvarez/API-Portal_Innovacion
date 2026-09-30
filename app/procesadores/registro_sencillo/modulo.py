"""Migración de `Registro_Sencillo` (transferencias entre cuentas de peaje sencillo).

Traduce `Registro_Sencillo.py` —el script manual de la raíz del repositorio—
a un `Procesador` de este servicio, siguiendo el precedente inmediato de
`app/procesadores/prepago_carga`, del que el script es casi una copia. Las
reglas de negocio son las del script y esta migración **no las
reinterpreta**: el layout de las tres líneas, la regla de la glosa, el
redondeo de importes y el formateo célula por célula del TXT salen de ahí tal
cual.

**Rarezas del script que se conservan a propósito**, porque cambiarlas
cambiaría la salida:

- El **año** de la cabecera sale de FECHA DOCUMENTO y el **mes** de FECHA
  CONTABLE. Con fechas del mismo mes da igual; en un cruce de año no.
- Las dos líneas de detalle —la 50 del banco de salida y la 40 del banco de
  entrada— llevan **el mismo importe: IMPORTE SALIDA**. El script nunca lee
  IMPORTE ENTRADA; este módulo tampoco, y ni siquiera exige esa columna.
- La glosa se decide por el **primer** caso que coincide, en este orden:
  banco de salida `1101015005`, banco de salida `1101015055`, banco de
  entrada `1101015006`, banco de entrada `1101015056`; si ninguno coincide,
  `TRANSFERENCIA ENTRE CUENTAS`. La regla de salida gana sobre la de
  entrada. La comparación es **igualdad numérica**, como en el script: un
  banco escrito como TEXTO `"1101015005"` en la celda no coincide (pandas lo
  dejaba como `str` y `"1101015005" == 1101015005` es falso), y la fila sale
  con la glosa por defecto. Se replica tal cual.
- Un banco `float` se vuelca como `int`; cualquier otro valor, tal cual.
- Un importe que no parsea vale cero, sin aviso (`limpiar_monto`).

Lo que sí cambia, por los mismos requisitos del PRD que fijó `contado_carga`:

1. **Columnas por nombre, no por posición.** El script lee `iloc[0]` a
   `iloc[4]`: si alguien inserta una columna, todo se corre un lugar sin
   fallar. Acá falta una columna obligatoria y la ejecución muere con
   `ErrorContenido.columna_faltante`.
2. **Fin de línea fijado explícitamente** (`FIN_DE_LINEA`), nunca heredado de
   la plataforma.
3. **Los descartes se cuentan y se reportan.** El script tiene tres
   descartes silenciosos —BANCO ENTRADA vacío con un `continue` pelado, y
   cualquiera de las dos fechas inválida con un `print`— y acá los tres salen
   en `descartes.txt`, con motivo, fila y valor. El cuarto `continue` del
   script, el de BANCO SALIDA nulo, es **inalcanzable**: el filtro previo ya
   quitó esas filas. No se migra.
4. **Cero filas procesadas es `ErrorContenido.cero_filas`**, nunca un éxito
   con las manos vacías.
5. **Ninguna excepción se traga.** El `try/except: print; continue` por fila
   desaparece; una fila que rompa de una forma no prevista levanta y
   `app/core/ejecucion.py` la convierte en `FalloDelModulo`.

**Limitación conocida, aceptada a propósito:** el extracto de `prepago-carga`
trae las mismas cinco columnas que este módulo exige (más `TIPO`, `COMISON`,
etc.), así que un archivo de prepago subido por error a este procesador **se
procesa sin error** y produce un asiento de peaje sencillo que ignora la
comisión. No se agrega una validación para detectarlo: se decidió cubrirlo
con la descripción del procesador en el portal.

Este módulo importa la biblioteca estándar, `openpyxl` y —de `core/`— la
interfaz, los tipos, los errores y el registro operativo. **No importa
`app.core.db`** ni nada que lo alcance (ADR 0013), no crea procesos ni hilos,
y no tiene efecto de importación alguno.
"""

from __future__ import annotations

import math
from collections.abc import Sequence
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any, Final

from openpyxl import load_workbook

from app.core.errores import ErrorContenido, ErrorTipificado
from app.core.interfaz import Procesador
from app.core.registro import configurar_logging, registrar_descartes
from app.core.tipos import ArchivoEntrada, ArchivoSalida

CLAVE: Final[str] = "registro-sencillo"


@dataclass(frozen=True, slots=True)
class _Columnas:
    """Los cinco encabezados que el módulo necesita. Ver la nota del módulo."""

    banco_salida: str
    banco_entrada: str
    fecha_documento: str
    fecha_contable: str
    importe_salida: str


# CONFIRMADOS contra la fila de encabezado del extracto real ("TRANSFERENCIAS
# ENTRE CUENTAS.xlsx", 6 columnas, hoja "Hoja1"). Las posiciones coinciden con
# los `iloc` del script: 0 a 4. La sexta columna, IMPORTE ENTRADA, el script
# no la lee y por eso no figura acá.
COLUMNAS: Final[_Columnas] = _Columnas(
    banco_salida="BANCO SALIDA",
    banco_entrada="BANCO ENTRADA",
    fecha_documento="FECHA DOCUMENTO",
    fecha_contable="FECHA CONTABLE",
    importe_salida="IMPORTE SALIDA",
)

FIN_DE_LINEA: Final[str] = "\r\n"
"""CRLF, fijado explícitamente (PRD, endurecimiento obligatorio).

**Observado, no confirmado**, igual que en `prepago_carga`: la salida
histórica del script trae 40 CRLF y cero LF sueltos, así que CRLF es lo que
el sistema destino viene recibiendo. Nadie del lado consumidor confirmó que
lo *exija*."""

CODIFICACION: Final[str] = "utf-8"
COLUMNAS_POR_FILA: Final[int] = 26
UMBRAL_DE_IDENTIFICADOR: Final[float] = 1000000000
"""Por encima de este valor, un `float` entero se vuelca sin decimales: el
script asume que es un identificador y no un monto. Se migra verbatim."""

TEXTO_CABECERA: Final[str] = "PEAJE SENCILLO"

GLOSA_POR_DEFECTO: Final[str] = "TRANSFERENCIA ENTRE CUENTAS"
GLOSA_IBK: Final[str] = "REMESAS IBK 223"
GLOSA_BBVA: Final[str] = "REMESAS BBVA 332"

REGLAS_BANCO_SALIDA: Final[tuple[tuple[int, str], ...]] = (
    (1101015005, GLOSA_IBK),
    (1101015055, GLOSA_BBVA),
)
"""Se evalúan ANTES que las del banco de entrada: el primer caso que coincide gana."""

REGLAS_BANCO_ENTRADA: Final[tuple[tuple[int, str], ...]] = (
    (1101015006, GLOSA_IBK),
    (1101015056, GLOSA_BBVA),
)

NOMBRE_TXT: Final[str] = "resultado_transferencias.txt"
NOMBRE_DESCARTES: Final[str] = "descartes.txt"

MIME_TEXTO: Final[str] = "text/plain; charset=utf-8"

_NULOS: Final[frozenset[str]] = frozenset(
    {
        "",
        "#N/A",
        "#N/A N/A",
        "#NA",
        "-1.#IND",
        "-1.#QNAN",
        "-NaN",
        "-nan",
        "1.#IND",
        "1.#QNAN",
        "<NA>",
        "N/A",
        "NA",
        "NULL",
        "NaN",
        "None",
        "n/a",
        "nan",
        "null",
    }
)
"""Los textos que `pd.read_excel` convierte en NaN al leer (`na_values` por
defecto). Comparados EXACTOS, sin `strip()`, como hace pandas. Deciden qué es
"vacío" para el filtro previo, para BANCO ENTRADA y para el importe."""

_FORMATOS_DE_FECHA: Final[tuple[str, ...]] = (
    "%d.%m.%Y",
    "%d/%m/%Y",
    "%d-%m-%Y",
    "%d.%m.%y",
    "%d/%m/%y",
    "%d-%m-%y",
    "%Y-%m-%d",
    "%Y/%m/%d",
    "%Y%m%d",
    "%d.%m.%Y %H:%M:%S",
    "%d/%m/%Y %H:%M:%S",
    "%Y-%m-%d %H:%M:%S",
)
"""Formatos de texto aceptados, **día primero**, replicando el `dayfirst=True`
del script. El extracto real trae las dos fechas como TEXTO `01.09.2026`, así
que el primer formato no es un caso de borde: es el caso normal."""

MOTIVO_BANCO_ENTRADA: Final[str] = "banco_entrada_vacio"
MOTIVO_FECHA_DOCUMENTO: Final[str] = "fecha_documento_invalida"
MOTIVO_FECHA_CONTABLE: Final[str] = "fecha_contable_invalida"


@dataclass(frozen=True, slots=True)
class Descarte:
    """Una fila que no llegó a la salida, con por qué y con qué valor."""

    fila: int
    motivo: str
    valor: str


def _es_nulo(valor: object) -> bool:
    """Lo que pandas habría leído como NaN."""
    if valor is None:
        return True
    if isinstance(valor, float):
        return math.isnan(valor)
    return isinstance(valor, str) and valor in _NULOS


def _a_fecha(valor: object) -> date | None:
    """Equivalente de `pd.to_datetime(valor, dayfirst=True, errors='coerce')`.

    Devuelve `None` en vez de `NaT`; quien llama lo trata como el descarte que
    el script hacía con un `continue`. Un valor numérico **no es una fecha**,
    con el mismo criterio que los procesadores anteriores: `pd.to_datetime`
    lo leía como nanosegundos desde 1970 y producía un asiento con fecha
    `19700101`. Esa basura se reporta en vez de migrarse.
    """
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    if not isinstance(valor, str) or _es_nulo(valor):
        return None
    texto = valor.strip()
    for formato in _FORMATOS_DE_FECHA:
        try:
            return datetime.strptime(texto, formato).date()
        except ValueError:
            continue
    return None


def _a_importe(valor: object) -> float:
    """Equivalente de `limpiar_monto`: nulo es 0, quita comas, lo que no parsea es 0.

    Siempre `float`, como el script: un importe entero `63000` pasa a
    `63000.0`, y eso es lo que el TXT vuelca (ver `_celda_de_txt`).
    """
    if _es_nulo(valor):
        return 0.0
    try:
        return float(str(valor).replace(",", "").strip())
    except ValueError:
        return 0.0


def _a_banco(valor: object) -> object:
    """Un `float` se vuelca como `int`; cualquier otro valor, tal cual."""
    return int(valor) if isinstance(valor, float) else valor


def _a_texto(valor: object) -> str:
    return "" if valor is None else str(valor)


def _glosa(banco_salida: object, banco_entrada: object) -> str:
    """La glosa de la columna 25, con el mismo orden de precedencia que el script.

    Igualdad numérica pelada (`==` contra un `int`), como el script: un banco
    de TEXTO no coincide nunca y cae en `GLOSA_POR_DEFECTO`. Ver la nota del
    módulo.
    """
    for cuenta, glosa in REGLAS_BANCO_SALIDA:
        if banco_salida == cuenta:
            return glosa
    for cuenta, glosa in REGLAS_BANCO_ENTRADA:
        if banco_entrada == cuenta:
            return glosa
    return GLOSA_POR_DEFECTO


def _celda_de_txt(celda: object) -> str:
    """Formateo célula por célula del volcado, migrado verbatim del script.

    Un `float` entero por encima de `UMBRAL_DE_IDENTIFICADOR` pierde los
    decimales (el script asume identificador), y **cualquier otro `float` sale
    con `str()` tal cual**: un importe de `63000` se vuelca `"63000.0"`. Los
    bancos son `int` y esquivan la rama. "Arreglar" esto rompe la paridad
    byte a byte.
    """
    if celda == "":
        return ""
    if isinstance(celda, float):
        if celda == int(celda) and celda > UMBRAL_DE_IDENTIFICADOR:
            return str(int(celda))
        return str(celda)
    return str(celda)


def _indices_de_columnas(encabezado: Sequence[object], archivo: str) -> dict[str, int]:
    """Mapa nombre → índice de las cinco columnas obligatorias.

    Falla con `ErrorContenido.columna_faltante` en la PRIMERA que falte, en el
    orden en que están declaradas en `COLUMNAS`.
    """
    posiciones = {
        str(nombre).strip(): indice
        for indice, nombre in enumerate(encabezado)
        if nombre is not None
    }
    indices: dict[str, int] = {}
    for campo in (
        COLUMNAS.banco_salida,
        COLUMNAS.banco_entrada,
        COLUMNAS.fecha_documento,
        COLUMNAS.fecha_contable,
        COLUMNAS.importe_salida,
    ):
        if campo not in posiciones:
            raise ErrorContenido.columna_faltante(archivo=archivo, columna=campo)
        indices[campo] = posiciones[campo]
    return indices


def _celda(fila: list[Any], indice: int) -> object:
    return fila[indice] if indice < len(fila) else None


def _leer_filas(ruta: Path, archivo: str) -> tuple[list[list[Any]], dict[str, int]]:
    """Encabezado y filas de datos de la primera hoja, en modo sólo lectura.

    La primera hoja, como el `read_excel` sin `sheet_name` del script.
    **Se abre un descriptor y se le pasa el objeto, nunca la ruta**: los
    temporales de `app/recepcion.py` se llaman `entrada_0`, sin extensión, y
    `load_workbook` rechaza una ruta por su extensión antes de leer un byte
    (ver `contado_carga._leer_filas`, donde se descubrió). Un `.xlsx` ilegible
    o sin ninguna fila es error de contenido, no una excepción cruda.
    """
    try:
        with ruta.open("rb") as descriptor:
            libro = load_workbook(descriptor, read_only=True, data_only=True)
            try:
                hoja = libro.worksheets[0]
                filas: list[list[Any]] = [list(fila) for fila in hoja.iter_rows(values_only=True)]
            finally:
                # Antes de cerrar el descriptor: en modo `read_only` el libro
                # lee perezosamente y sigue necesitándolo mientras se itera.
                libro.close()
    except Exception as error:
        raise ErrorContenido.cero_filas(archivo=archivo) from error

    if not filas:
        raise ErrorContenido.cero_filas(archivo=archivo)
    return filas[1:], _indices_de_columnas(filas[0], archivo)


def _linea_detalle(clave_contable: int, cuenta: object, importe: float, glosa: str) -> list[Any]:
    linea: list[Any] = [""] * COLUMNAS_POR_FILA
    linea[0] = 2
    linea[12] = clave_contable
    linea[13] = cuenta
    linea[17] = round(abs(importe), 2)
    linea[25] = glosa
    return linea


def _procesar_fila(fila: list[Any], indices: dict[str, int]) -> list[list[Any]] | Descarte:
    """Las filas de salida de un registro, o su descarte (con `fila=0`: la
    numeración la pone quien llama)."""
    banco_entrada = _celda(fila, indices[COLUMNAS.banco_entrada])
    if _es_nulo(banco_entrada):
        return Descarte(fila=0, motivo=MOTIVO_BANCO_ENTRADA, valor=_a_texto(banco_entrada))

    crudo_documento = _celda(fila, indices[COLUMNAS.fecha_documento])
    fecha_documento = _a_fecha(crudo_documento)
    if fecha_documento is None:
        return Descarte(fila=0, motivo=MOTIVO_FECHA_DOCUMENTO, valor=_a_texto(crudo_documento))

    crudo_contable = _celda(fila, indices[COLUMNAS.fecha_contable])
    fecha_contable = _a_fecha(crudo_contable)
    if fecha_contable is None:
        return Descarte(fila=0, motivo=MOTIVO_FECHA_CONTABLE, valor=_a_texto(crudo_contable))

    importe = _a_importe(_celda(fila, indices[COLUMNAS.importe_salida]))
    banco_salida = _a_banco(_celda(fila, indices[COLUMNAS.banco_salida]))
    banco_entrada = _a_banco(banco_entrada)
    glosa = _glosa(banco_salida, banco_entrada)

    cabecera: list[Any] = [""] * COLUMNAS_POR_FILA
    cabecera[0] = 1
    cabecera[2] = "VPR1"
    cabecera[4] = fecha_documento.year  # año del DOCUMENTO...
    cabecera[5] = "SA"
    cabecera[6] = fecha_documento.strftime("%Y%m%d")
    cabecera[7] = fecha_contable.strftime("%Y%m%d")
    cabecera[8] = fecha_contable.strftime("%m")  # ...y mes CONTABLE. Verbatim.
    cabecera[10] = TEXTO_CABECERA
    cabecera[11] = "PEN"

    # El MISMO importe (IMPORTE SALIDA) en las dos líneas, como el script.
    salida = _linea_detalle(50, banco_salida, importe, glosa)
    entrada = _linea_detalle(40, banco_entrada, importe, glosa)

    return [cabecera, salida, entrada, [""] * COLUMNAS_POR_FILA]


def _escribir_txt(filas: list[list[Any]], ruta: Path) -> None:
    """Vuelca las filas separadas por tabulaciones, con `FIN_DE_LINEA` explícito.

    `newline=""` es obligatorio: sin él, Python traduce cada `\\n` según la
    plataforma y el fin de línea volvería a heredarse.
    """
    with ruta.open("w", encoding=CODIFICACION, newline="") as salida:
        for fila in filas:
            salida.write("\t".join(_celda_de_txt(celda) for celda in fila) + FIN_DE_LINEA)


def _escribir_descartes(descartes: list[Descarte], ruta: Path) -> None:
    """Reporte de descartes, una fila por línea, con el mismo fin de línea."""
    conteo: dict[str, int] = {}
    for descarte in descartes:
        conteo[descarte.motivo] = conteo.get(descarte.motivo, 0) + 1
    resumen = ", ".join(f"{motivo}={total}" for motivo, total in sorted(conteo.items()))

    lineas = [
        f"Filas descartadas: {len(descartes)} ({resumen})",
        "",
        "\t".join(("fila", "motivo", "valor")),
    ]
    lineas.extend(
        "\t".join((str(d.fila), d.motivo, d.valor.replace("\t", " ").replace("\n", " ")))
        for d in descartes
    )
    with ruta.open("w", encoding=CODIFICACION, newline="") as salida:
        salida.write(FIN_DE_LINEA.join(lineas) + FIN_DE_LINEA)


class RegistroSencillo(Procesador):
    """El procesador real: un Excel de transferencias entra, un TXT sale.

    Más un `descartes.txt` cuando —y sólo cuando— hubo filas descartadas: con
    una entrada limpia la salida es un único archivo, como la del script.
    `app.core.empaquetado` decide si eso viaja suelto o comprimido; este
    módulo nunca arma un ZIP (ADR 0006).
    """

    def __init__(self, clave: str = CLAVE) -> None:
        self.clave = clave

    def validar(self, archivos: list[ArchivoEntrada]) -> ErrorTipificado | None:
        """Comprueba que el Excel se abra y traiga las cinco columnas obligatorias.

        Devuelve el error en vez de levantarlo, como fija `Procesador`.
        """
        for entrada in archivos:
            try:
                _leer_filas(entrada.ruta_temporal, entrada.nombre_original)
            except ErrorContenido as error:
                return error
        return None

    def procesar(self, archivos: list[ArchivoEntrada]) -> list[ArchivoSalida]:
        """Transforma la entrada y devuelve uno o dos archivos.

        Cero filas procesadas levanta `ErrorContenido.cero_filas` **antes** de
        escribir nada. Ninguna excepción se traga; una fila que rompa de una
        forma no prevista se propaga y `app/core/ejecucion.py` la convierte en
        `FalloDelModulo`.
        """
        entrada = archivos[0]
        directorio = entrada.ruta_temporal.parent
        filas_crudas, indices = _leer_filas(entrada.ruta_temporal, entrada.nombre_original)

        salida_filas: list[list[Any]] = []
        descartes: list[Descarte] = []
        registros = 0
        for numero, fila in enumerate(filas_crudas, start=2):  # 1 es el encabezado
            if _es_nulo(_celda(fila, indices[COLUMNAS.banco_salida])):
                continue  # el script filtra estas filas ANTES de contar nada

            resultado = _procesar_fila(fila, indices)
            if isinstance(resultado, Descarte):
                descartes.append(
                    Descarte(fila=numero, motivo=resultado.motivo, valor=resultado.valor)
                )
                continue
            salida_filas.extend(resultado)
            registros += 1

        if not salida_filas:
            raise ErrorContenido.cero_filas(archivo=entrada.nombre_original)

        ruta_txt = directorio / NOMBRE_TXT
        _escribir_txt(salida_filas, ruta_txt)
        salidas = [
            ArchivoSalida(nombre_propuesto=NOMBRE_TXT, ruta_temporal=ruta_txt, tipo_mime=MIME_TEXTO)
        ]

        if descartes:
            ruta_descartes = directorio / NOMBRE_DESCARTES
            _escribir_descartes(descartes, ruta_descartes)
            salidas.append(
                ArchivoSalida(
                    nombre_propuesto=NOMBRE_DESCARTES,
                    ruta_temporal=ruta_descartes,
                    tipo_mime=MIME_TEXTO,
                )
            )
            # Se configura ACÁ y no al importar: este código corre en el hijo
            # de `spawn`, que no tiene handler (ver `contado_carga.procesar`).
            configurar_logging()
            registrar_descartes(
                clave=self.clave,
                archivo=entrada.nombre_original,
                filas_procesadas=registros,
                descartes=[(d.motivo, d.fila) for d in descartes],
            )

        return salidas
