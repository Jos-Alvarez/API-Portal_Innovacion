"""Migración de `flujoCaja_consolida_pagos`.

Traduce el script manual de la raíz del repositorio que consolida los libros
semanales de "Pagos Programados" (`.xlsm`) de dos tipos —**LIMA EXPRESA** y
**PEX**— en `Consolidado_Pagos.xlsx`, con las hojas `Soles` y `Dolares` (cada
una sólo si tiene filas). Las reglas de negocio son las del script y esta
migración **no las reinterpreta**: grupos de hojas por tipo, las 28 cabeceras
estándar asignadas por posición, el filtro de `Cuenta`, `categorizar`, la
limpieza de importes, el signo por `Clase de documento`, el `PPP`, el cálculo
de semanas, los tres totales por semana, el orden de columnas y el orden de
filas salen de ahí tal cual.

**Entradas y salidas.** Una ejecución recibe de uno a diez libros; varios del
MISMO tipo es lo normal (un libro por semana). Sale un solo
`Consolidado_Pagos.xlsx`, más un `observaciones.txt` cuando —y sólo cuando—
pasó algo que el script salteaba o perdía en silencio.
`app.core.empaquetado` decide si todo eso viaja suelto o en ZIP; este módulo
nunca arma un ZIP (ADR 0006).

**Orden de los libros** (`_ordenar`). El script tomaba primero TODOS los
libros LIMA EXPRESA y después TODOS los PEX, cada grupo en el orden en que
`glob.glob` los devolvía. En Windows `glob` usa `os.scandir`, que devuelve el
orden del índice del directorio de NTFS: los nombres comparados unidad UTF-16
por unidad UTF-16 después de pasarlos a mayúsculas con la tabla `$UpCase` del
volumen (`_clave_ntfs`). Acá se ordena igual, por `nombre_original`. Ese orden
es el de concatenación, así que decide el orden de las filas de la salida.

Lo que cambia respecto del script, por el PRD ("Endurecimiento obligatorio"):

1. **El tipo de libro se detecta por las HOJAS** que sólo tiene ese tipo
   (`_HOJAS_DISTINTIVAS`); el nombre del archivo (`*LIMA EXPRESA*`, `*PEX*`,
   como el `glob` del script, sin distinguir mayúsculas) es sólo el respaldo.
   Hojas de los dos tipos, o ninguna hoja y un nombre que no decide, es
   `tipo_no_reconocido`. El script, con un nombre que contenía LAS DOS
   subcadenas, procesaba el libro UNA sola vez como PEX pero en la posición
   de los LIMA EXPRESA (la segunda asignación del diccionario pisaba la
   primera sin moverla); acá ese nombre, sin hojas que decidan, no se
   adivina.
2. **Los encabezados se validan.** El script asigna las 28 cabeceras por
   POSICIÓN y descarta la primera fila, sea lo que sea. Las hojas reales sí
   tienen el encabezado en la primera fila, con variantes de puntuación entre
   hojas (`Base p plazo pago`, `Base p, plazo pago`, `Base p/ plazo pago`) y
   columnas sin rótulo que ninguna regla lee. Se sigue leyendo por posición,
   pero antes se comprueba que el rótulo de cada columna que alguna regla
   LEE (`_ENCABEZADOS_VALIDADOS`) diga lo esperado, comparando sólo letras y
   dígitos. Si el layout cambió, la ejecución muere con `columna_faltante`.
3. **Ninguna excepción se traga.** Desaparecen los `try/except: print`.
4. **Hoja configurada ausente: se sigue salteando**, como el script, pero se
   reporta (`hoja_ausente`). Una hoja presente que no aporta ninguna fila
   **no** se reporta: toda semana real trae alguna (caja chica, planilla) y es
   operación normal, no una pérdida; reportarla convertía cada ejecución
   normal en un ZIP. Ninguna fila en toda la ejecución es `cero_filas`.
5. **Las pérdidas silenciosas se reportan**: importes ilegibles
   (`importe_invalido`), fechas ilegibles que dejan sin `PPP`, `Mes` o
   `Semana` (`fecha_invalida`), una hoja "por moneda" sin columna de moneda
   (`moneda_ausente`, el script la descartaba entera) y las filas de esas
   hojas cuya moneda no es `PEN` ni `USD` (`moneda_no_reconocida`, el script
   las descartaba).

**Rarezas del script migradas VERBATIM** —parecen errores, pero cambiarlas
cambia la salida:

- El `-` final de un importe (`"12,500.00-"`) se QUITA, no lo vuelve
  negativo: el signo lo decide sólo `Clase de documento == '7'`, que lo pone
  negativo; toda otra clase lo deja positivo (`abs`).
- `Base p plazo pago` se interpreta con el día primero (`dayfirst=True`)
  tras cambiar `.` por `-`; `Fecha de pago BANCOS` SIN `dayfirst`.
- `calcular_semanas_mes`: cuatro semanas, la cuarta hasta el fin de mes (ver
  la migración de `flujo-caja-ingresos`, que comparte la regla).
- Si la hoja no llega a la columna del importe, `Total Hoja por Semana` y
  `Total Mes por Semana` valen 0 y `Total Categoria por Semana` no existe.

**Emulación de pandas.** Este módulo no importa pandas (el hijo de `spawn` lo
reimportaría en cada petición; ver el README de la migración de
`contado_carga`), así que reproduce lo observable:

- `read_excel(header=None, dtype=str)`: un entero guardado como `float` se lee
  como `int` (`_celda`) y después TODO valor no nulo se vuelve `str()`
  (`_texto`): una fecha guardada como fecha se vuelve `"2025-07-02 00:00:00"`
  y `Cuenta` sale como TEXTO. Los textos de `_NULOS` son NaN. Las filas
  conservan su posición de Excel, incluidas las vacías, y el ancho del
  `DataFrame` es el de la fila más larga sin contar celdas vacías al final.
- `pd.to_numeric(errors='coerce')` en `_numero_de_texto`.
- `pd.to_datetime(errors='coerce')` sobre texto (`_fechas_de_columna`): pandas
  INFIERE un formato a partir del primer valor no nulo de la columna y lo
  aplica a todos; lo que no calza es `NaT`.
- `groupby(...).sum()` + `merge(how='left')`: suma compensada de Kahan por
  grupo, en el orden de las filas, saltando NaN (`_suma_de_grupo`). El
  `merge` a izquierda conserva el orden de las filas y, como las claves son
  únicas del lado derecho, sólo pega el total de su grupo a cada fila.
  Ninguna clave es NaN: `Mes`, `Semana` y `Categoria` valen `""` sin dato.
- `pd.concat` alinea columnas por nombre; las que una hoja no tiene quedan
  NaN.

**Decisiones de esta migración que el script no tomaba:**

- **Un número en `Fecha de pago BANCOS`** ya no puede llegar como número: con
  `dtype=str` es texto (`"45839"`), no calza con ningún formato y queda sin
  fecha; se reporta como `fecha_invalida`. Mismo resultado que el script.
- **Primer valor de fecha con una forma que este módulo no reconoce**: pandas
  cae a `dateutil` valor por valor; acá cada valor se interpreta por separado
  con las mismas formas (ISO, día/mes/año, año/mes/día). Las muestras reales
  no llegan a este caso.
- **`Excel` es el nombre base de `nombre_original`** (el `os.path.basename`
  del script), sin carpetas si el navegador las mandó.
- **Tipo detectado por nombre y ninguna hoja configurada** es
  `hoja_faltante` con la primera hoja de la configuración.

Este módulo importa la biblioteca estándar, `openpyxl` y —de `core/`— la
interfaz, los tipos, los errores y el registro operativo. **No importa
`app.core.db`** ni nada que lo alcance (ADR 0013), no crea procesos ni hilos,
y no tiene efecto de importación alguno.
"""

from __future__ import annotations

import math
import re
import unicodedata
from collections.abc import Mapping, Sequence
from dataclasses import dataclass, field
from datetime import datetime, timedelta
from pathlib import Path
from types import MappingProxyType
from typing import Final

from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import Cell
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.worksheet._read_only import ReadOnlyWorksheet
from openpyxl.worksheet._write_only import WriteOnlyWorksheet

from app.core.errores import ErrorContenido, ErrorTipificado
from app.core.interfaz import Procesador
from app.core.registro import configurar_logging, registrar_descartes
from app.core.tipos import ArchivoEntrada, ArchivoSalida

CLAVE: Final[str] = "flujo-caja-pagos"

TIPO_LIMA_EXPRESA: Final[str] = "LIMA EXPRESA"
TIPO_PEX: Final[str] = "PEX"
TIPOS: Final[tuple[str, ...]] = (TIPO_LIMA_EXPRESA, TIPO_PEX)
"""Los dos tipos, en el orden en que el script los procesaba, con la subcadena
que su `glob` buscaba en el nombre (`*LIMA EXPRESA*`, `*PEX*`). También es el
valor de la columna `Sociedad`."""

NOMBRE_CONSOLIDADO: Final[str] = "Consolidado_Pagos.xlsx"
NOMBRE_OBSERVACIONES: Final[str] = "observaciones.txt"

HOJA_SOLES: Final[str] = "Soles"
HOJA_DOLARES: Final[str] = "Dolares"
GRUPO_POR_MONEDA: Final[str] = "Por moneda"
MONEDA_SOLES: Final[str] = "PEN"
MONEDA_DOLARES: Final[str] = "USD"

FIN_DE_LINEA: Final[str] = "\r\n"
"""El de `observaciones.txt` de los demás procesadores: fijado, no heredado de
la plataforma."""
CODIFICACION: Final[str] = "utf-8"

MIME_XLSX: Final[str] = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MIME_TEXTO: Final[str] = "text/plain; charset=utf-8"

MOTIVO_HOJA_AUSENTE: Final[str] = "hoja_ausente"
MOTIVO_IMPORTE: Final[str] = "importe_invalido"
MOTIVO_FECHA: Final[str] = "fecha_invalida"
MOTIVO_MONEDA_AUSENTE: Final[str] = "moneda_ausente"
MOTIVO_MONEDA: Final[str] = "moneda_no_reconocida"

MESES_ES: Final[tuple[str, ...]] = (
    "Ene",
    "Feb",
    "Mar",
    "Abr",
    "May",
    "Jun",
    "Jul",
    "Ago",
    "Sep",
    "Oct",
    "Nov",
    "Dic",
)

FORMATO_FECHA_SALIDA: Final[str] = "YYYY-MM-DD HH:MM:SS"
"""El `datetime_format` por defecto con que `to_excel` escribía las fechas."""


# --- Configuración ------------------------------------------------------------

CABECERAS_ESTANDAR: Final[tuple[str, ...]] = (
    "Ejercicio / mes",
    "Cuenta",
    "Fecha de documento",
    "Período contable",
    "Base p plazo pago",
    "Condiciones de pago",
    "Vencimiento neto",
    "Nombre 1",
    "Clase de documento",
    "Referencia",
    "Nº documento",
    "Moneda del documento",
    "Imptebase descuento",
    "Importe en moneda local",
    "Importe en ML3",
    "Retencion",
    "Fecha de entrada",
    "Via de pago",
    "Asignación",
    "Fecha de pago",
    "Fecha compensación",
    "Doccompensación",
    "Texto",
    "Nombre del usuario",
    "Documento compras",
    "Centro de coste",
    "Orden",
    "Fecha de pago BANCOS",
)
"""`CABECERAS_ESTANDAR` del script, verbatim: los títulos de la salida."""

NUM_COLUMNAS_PERMITIDAS: Final[int] = len(CABECERAS_ESTANDAR)

COL_CUENTA: Final[str] = "Cuenta"
COL_BASE: Final[str] = "Base p plazo pago"
COL_CLASE: Final[str] = "Clase de documento"
COL_MONEDA: Final[str] = "Moneda del documento"
COL_IMPORTE_LOCAL: Final[str] = "Importe en moneda local"
COL_IMPORTE_ML3: Final[str] = "Importe en ML3"
COL_ORDEN: Final[str] = "Orden"
COL_FECHA_PAGO: Final[str] = "Fecha de pago BANCOS"

COLUMNAS_IMPORTE: Final[tuple[str, ...]] = (COL_IMPORTE_LOCAL, COL_IMPORTE_ML3)

_ENCABEZADOS_VALIDADOS: Final[tuple[int, ...]] = tuple(
    CABECERAS_ESTANDAR.index(c)
    for c in (
        COL_CUENTA,
        COL_BASE,
        COL_CLASE,
        COL_MONEDA,
        COL_IMPORTE_LOCAL,
        COL_IMPORTE_ML3,
        COL_ORDEN,
        COL_FECHA_PAGO,
    )
)
"""Posiciones (desde 0) de las columnas que alguna regla lee: son las únicas
cuyo rótulo se valida (ver la docstring del módulo)."""

COLUMNAS_GENERADAS: Final[tuple[str, ...]] = (
    "Sociedad",
    "Excel",
    "Fuente",
    "Categoria",
    "PPP",
    "Mes",
    "Semana",
    "Total Hoja por Semana",
    "Total Mes por Semana",
    "Total Categoria por Semana",
)
"""`columnas_finales` del script: van al final, en este orden, las que haya."""

_TOTAL_HOJA: Final[str] = "Total Hoja por Semana"
_TOTAL_MES: Final[str] = "Total Mes por Semana"
_TOTAL_CATEGORIA: Final[str] = "Total Categoria por Semana"

GRUPOS_POR_TIPO: Final[Mapping[str, tuple[tuple[str, tuple[str, ...]], ...]]] = MappingProxyType(
    {
        TIPO_LIMA_EXPRESA: (
            (HOJA_SOLES, ("SOLES", "CAJA CHICA LE", "REEMBOLSOS_SOL")),
            (HOJA_DOLARES, ("DOLARES", "REEMBOLSOS_DOL")),
            (GRUPO_POR_MONEDA, ("PLANILLA", "PLATAFORMA IBK", "PLATAFORMA BBVA")),
        ),
        TIPO_PEX: (
            (HOJA_SOLES, ("PEX SOL",)),
            (HOJA_DOLARES, ("PEX DOL",)),
            (GRUPO_POR_MONEDA, ("PLATAFORMA BBVA", "PLATAFORMA BCP")),
        ),
    }
)
"""`obtener_grupos_hojas` del script, verbatim y en su orden (que es el de
concatenación)."""


def _hojas_de_tipo(tipo: str) -> tuple[str, ...]:
    return tuple(hoja for _grupo, hojas in GRUPOS_POR_TIPO[tipo] for hoja in hojas)


_HOJAS_COMPARTIDAS: Final[frozenset[str]] = frozenset({"PLATAFORMA BBVA", "PLATAFORMA BCP"})
"""Hojas que no deciden el tipo. `PLATAFORMA BBVA` está configurada para los
dos; `PLATAFORMA BCP` sólo para PEX, pero los tres libros LIMA EXPRESA reales
también la traen (sin que el script la lea para ellos)."""

_HOJAS_DISTINTIVAS: Final[Mapping[str, frozenset[str]]] = MappingProxyType(
    {tipo: frozenset(_hojas_de_tipo(tipo)) - _HOJAS_COMPARTIDAS for tipo in TIPOS}
)
"""Las hojas configuradas que deciden el tipo: `SOLES`, `CAJA CHICA LE`,
`REEMBOLSOS_SOL`, `DOLARES`, `REEMBOLSOS_DOL`, `PLANILLA` y `PLATAFORMA IBK`
para LIMA EXPRESA; `PEX SOL` y `PEX DOL` para PEX."""


# --- Conversiones que emulan a pandas -----------------------------------------

_NULOS: Final[frozenset[str]] = frozenset(
    {
        "",
        "#N/A",
        "#N/A N/A",
        "#NA",
        "-1.#IND",
        "-1.#QNAN",
        "-NaN",
        "-nan",
        "1.#IND",
        "1.#QNAN",
        "<NA>",
        "N/A",
        "NA",
        "NULL",
        "NaN",
        "None",
        "n/a",
        "nan",
        "null",
        # Errores de Excel: `read_excel` los lee como NaN, y openpyxl en modo
        # `values_only` los entrega como este texto.
        "#DIV/0!",
        "#VALUE!",
        "#REF!",
        "#NAME?",
        "#NUM!",
        "#NULL!",
    }
)
"""Lo que `pd.read_excel` convierte en NaN al leer. Comparado EXACTO, sin
`strip()`, como hace pandas."""

_ESPACIOS_ASCII: Final[str] = " \t\n\r\f\v"
_ENTERO: Final[re.Pattern[str]] = re.compile(r"[+-]?\d+")
_GUION_FINAL: Final[re.Pattern[str]] = re.compile(r"-$")
"""El `str.replace(r'-$', '', regex=True)` del script (`re.sub`: `$` también
calza antes de un salto de línea final)."""

Importe = int | float


def _es_nulo(valor: object) -> bool:
    """Lo que pandas habría leído como NaN."""
    if valor is None:
        return True
    if isinstance(valor, float):
        return math.isnan(valor)
    return isinstance(valor, str) and valor in _NULOS


def _celda(valor: object) -> object:
    """El `_convert_cell` del lector openpyxl de pandas: `66.0` se lee `66`."""
    if isinstance(valor, float) and math.isfinite(valor) and valor.is_integer():
        return int(valor)
    return valor


def _texto(valor: object) -> str | None:
    """`read_excel(dtype=str)`: `None` por NaN; si no, `str()` de lo leído."""
    convertido = _celda(valor)
    return None if _es_nulo(convertido) else str(convertido)


def _numero_de_texto(texto: str) -> Importe | None:
    """`pd.to_numeric(texto, errors='coerce')`: `None` por NaN.

    Tolera espacios ASCII alrededor (`" 12 "` es 12) pero no el espacio duro,
    ni separadores de miles, ni el `_` que Python sí acepta.
    """
    limpio = texto.strip(_ESPACIOS_ASCII)
    if not limpio or not limpio.isascii() or "_" in limpio:
        return None
    if _ENTERO.fullmatch(limpio):
        return int(limpio)
    try:
        numero = float(limpio)
    except ValueError:
        return None
    return None if math.isnan(numero) else numero


def limpiar_importe(texto: str | None) -> Importe | None:
    """La limpieza de importes del script: quita UN `-` final y las comas, y
    `pd.to_numeric(errors='coerce')`. El `-` final NO vuelve negativo el
    importe (ver `aplicar_signo`)."""
    if texto is None:
        return None
    return _numero_de_texto(_GUION_FINAL.sub("", texto).replace(",", ""))


def aplicar_signo(importe: Importe | None, clase: str | None) -> Importe | None:
    """`-abs(x)` si `str(clase).strip() == '7'`; `abs(x)` en otro caso. NaN sigue NaN."""
    if importe is None:
        return None
    clase_str = "nan" if clase is None else clase
    return -abs(importe) if clase_str.strip() == "7" else abs(importe)


def categorizar(orden: str | None) -> str:
    """`categorizar` del script, verbatim, sobre el `Orden` leído como texto."""
    if orden is None or orden.strip() == "":
        return "OPEX"
    orden_str = orden.strip()
    if orden_str.startswith("4"):
        return "OPEX"
    if orden_str.startswith(("2", "3")):
        return "CAPEX"
    return ""


_NAT: Final[frozenset[str]] = frozenset({"NaT", "nat", "NAT", "nan", "NaN", "NAN"})
"""Textos que `pd.to_datetime` lee como `NaT` y que salta al buscar el primer
valor con el cual inferir el formato."""

_HORA: Final[str] = r"(?P<hora>[ T]\d{1,2}:\d{2}(?::\d{2}(?P<fraccion>\.\d{1,6})?)?)?"
_ANIO_PRIMERO: Final[re.Pattern[str]] = re.compile(
    r"\d{4}(?P<sep>[-/.])\d{1,2}(?P=sep)\d{1,2}" + _HORA
)
_ANIO_ULTIMO: Final[re.Pattern[str]] = re.compile(
    r"\d{1,2}(?P<sep>[-/.])\d{1,2}(?P=sep)(?P<anio>\d{4}|\d{2})" + _HORA
)


def _prueba(texto: str, formato: str) -> bool:
    try:
        datetime.strptime(texto, formato)
    except ValueError:
        return False
    return True


def _formato_de_hora(forma: re.Match[str]) -> str:
    hora = forma.group("hora")
    if not hora:
        return ""
    formato = f"{hora[0]}%H:%M" + (":%S" if hora.count(":") == 2 else "")
    return formato + (".%f" if forma.group("fraccion") else "")


def adivinar_formato(texto: str, *, dayfirst: bool, anio_corto: bool = False) -> str | None:
    """`guess_datetime_format` de pandas (que usa dateutil) para las formas que
    este módulo conoce: año primero o año último, con `-`, `/` o `.`, y hora
    opcional. `None` si no es una de ellas o no es una fecha válida.

    Con `dayfirst=True` se prueba primero el DÍA en la posición que dateutil
    le da, también con el año primero (`"2025-06-01"` → `%Y-%d-%m`, el 6 de
    enero, verificado contra pandas 3.0.6); sin él, primero el mes. El otro
    orden, sólo si el primero no es una fecha (`13-06-2025` no tiene mes 13).
    pandas no infiere formatos con año de dos cifras (`anio_corto=False`),
    pero dateutil sí los lee valor por valor (`anio_corto=True`). Los espacios
    alrededor pasan al formato, donde `strptime` exige al menos uno.
    """
    nucleo = texto.strip()
    if not nucleo:
        return None
    antes = texto[: len(texto) - len(texto.lstrip())]
    despues = texto[len(texto.rstrip()) :]
    ordenes = (("%d", "%m"), ("%m", "%d")) if dayfirst else (("%m", "%d"), ("%d", "%m"))
    candidatos: list[str] = []
    forma = _ANIO_PRIMERO.fullmatch(nucleo)
    if forma:
        sep = forma.group("sep")
        candidatos = [f"%Y{sep}{a}{sep}{b}{_formato_de_hora(forma)}" for a, b in ordenes]
    forma = _ANIO_ULTIMO.fullmatch(nucleo)
    if forma and (anio_corto or len(forma.group("anio")) == 4):
        sep = forma.group("sep")
        anio = "%Y" if len(forma.group("anio")) == 4 else "%y"
        candidatos = [f"{a}{sep}{b}{sep}{anio}{_formato_de_hora(forma)}" for a, b in ordenes]
    for candidato in candidatos:
        if _prueba(nucleo, candidato):
            return f"{antes}{candidato}{despues}"
    return None


def _a_fecha(texto: str, formato: str | None) -> datetime | None:
    if formato is None:
        return None
    try:
        return datetime.strptime(texto, formato)
    except ValueError:
        return None


def _fechas_de_columna(textos: Sequence[str | None], *, dayfirst: bool) -> list[datetime | None]:
    """`pd.to_datetime(columna_de_texto, errors='coerce', dayfirst=...)`.

    pandas infiere UN formato a partir del primer valor no nulo (saltando los
    textos de `_NAT`) y lo exige a todos con `strptime`: lo que no calza es
    `NaT`. Si del primer valor no se infiere nada, cada valor se interpreta
    por separado, como el respaldo de dateutil (ver la docstring del módulo).
    """
    primero = next((t for t in textos if t is not None and t and t not in _NAT), None)
    formato = None if primero is None else adivinar_formato(primero, dayfirst=dayfirst)
    fechas: list[datetime | None] = []
    for texto in textos:
        if texto is None:
            fechas.append(None)
        elif formato is not None:
            fechas.append(_a_fecha(texto, formato))
        else:
            individual = adivinar_formato(texto, dayfirst=dayfirst, anio_corto=True)
            fechas.append(_a_fecha(texto, individual))
    return fechas


def calcular_semanas_mes(mes: int, anio: int) -> tuple[tuple[datetime, datetime], ...]:
    """`calcular_semanas_mes` del script, verbatim."""
    primer_dia_mes = datetime(anio, mes, 1)
    ultimo_dia_mes = (primer_dia_mes.replace(day=28) + timedelta(days=4)).replace(
        day=1
    ) - timedelta(days=1)
    dia_semana = primer_dia_mes.weekday()
    if dia_semana == 6:  # domingo
        fin_sem1 = primer_dia_mes + timedelta(days=7)
    elif dia_semana in (2, 3, 4, 5):  # miércoles a sábado
        fin_sem1 = primer_dia_mes + timedelta(days=13 - dia_semana)
    else:  # lunes y martes
        fin_sem1 = primer_dia_mes + timedelta(days=6 - dia_semana)
    semanas = [(primer_dia_mes, fin_sem1)]
    inicio = fin_sem1 + timedelta(days=1)
    for i in range(2, 5):
        if i < 4:
            fin = inicio + timedelta(days=6)
        else:
            fin = inicio + timedelta(days=(ultimo_dia_mes - inicio).days)
        semanas.append((inicio, fin))
        inicio = fin + timedelta(days=1)
    return tuple(semanas)


def calcular_semana(fecha: datetime | None) -> str:
    """`calcular_semana`: `"Semana N"`, o `""` si la fecha es nula o no cae en ninguna."""
    if fecha is None:
        return ""
    for numero, (inicio, fin) in enumerate(calcular_semanas_mes(fecha.month, fecha.year), 1):
        if inicio <= fecha <= fin:
            return f"Semana {numero}"
    return ""


def _mes(fecha: datetime | None) -> str:
    return "" if fecha is None else MESES_ES[fecha.month - 1]


def _suma_de_grupo(valores: Sequence[Importe | None]) -> Importe:
    """`groupby(...).sum()`: suma compensada de Kahan que salta NaN, como
    `group_sum` de pandas. Un grupo sin ningún número suma 0.

    Con enteros la suma es exacta y se devuelve `int` (mismo número).
    """
    numeros = [v for v in valores if v is not None and not (isinstance(v, float) and v != v)]
    if all(isinstance(v, int) for v in numeros):
        return sum(numeros)
    total = 0.0
    compensacion = 0.0
    for valor in numeros:
        y = valor - compensacion
        t = total + y
        compensacion = t - total - y
        if compensacion != compensacion:  # NaN: pandas la reinicia igual
            compensacion = 0.0
        total = t
    return total


def _totales(claves: Sequence[tuple[str, ...]], valores: Sequence[Importe | None]) -> list[Importe]:
    """`groupby(claves)[valor].sum()` + `merge(how='left')`: el total del grupo
    en cada fila, en el orden de las filas."""
    grupos: dict[tuple[str, ...], list[Importe | None]] = {}
    for clave, valor in zip(claves, valores, strict=True):
        grupos.setdefault(clave, []).append(valor)
    sumas = {clave: _suma_de_grupo(grupo) for clave, grupo in grupos.items()}
    return [sumas[clave] for clave in claves]


def _clave_ntfs(nombre: str) -> bytes:
    """Clave de orden del índice de directorio de NTFS (lo que devolvía `glob`).

    NTFS compara los nombres como secuencias de unidades UTF-16 después de
    pasar cada unidad a mayúsculas con la tabla `$UpCase` del volumen, que es
    la conversión simple de Unicode (un carácter por un carácter). Los
    caracteres fuera del plano básico no se convierten. Como UTF-16 big
    endian se compara byte a byte igual que unidad a unidad, la clave es esa
    codificación.
    """
    mayusculas = "".join(
        c.upper() if ord(c) <= 0xFFFF and len(c.upper()) == 1 else c for c in nombre
    )
    return mayusculas.encode("utf-16-be", "surrogatepass")


def _nombre_base(nombre_original: str) -> str:
    """`os.path.basename` de Windows: lo que sigue a la última `/` o `\\`."""
    return nombre_original.replace("\\", "/").rpartition("/")[2]


# --- Lectura ------------------------------------------------------------------


@dataclass(frozen=True, slots=True)
class _Fila:
    """Una fila de datos con `Cuenta` no vacía: su número en Excel (desde 1) y
    sus primeras columnas como las deja `dtype=str`."""

    numero: int
    textos: tuple[str | None, ...]


@dataclass(frozen=True, slots=True)
class _HojaLeida:
    """`df_raw.iloc[1:, :28]` ya filtrado por `Cuenta`. `columnas` es la
    cantidad de columnas del `DataFrame` (hasta 28)."""

    nombre: str
    columnas: int
    filas: tuple[_Fila, ...]


@dataclass(frozen=True, slots=True)
class _Libro:
    archivo: str
    excel: str
    tipo: str
    hojas: Mapping[str, _HojaLeida]


def _normalizar(texto: object) -> str:
    """Rótulo comparable: sólo letras y dígitos, sin tildes, sin mayúsculas."""
    descompuesto = unicodedata.normalize("NFKD", str(texto))
    return "".join(c for c in descompuesto if c.isalnum()).casefold()


def _validar_encabezado(fila: Sequence[object], ancho: int, hoja: str, archivo: str) -> None:
    """Falla con `columna_faltante` en la primera columna leída cuyo rótulo no
    coincide. La columna viaja como `"<hoja>: <rótulo esperado>"`.

    Sólo se validan las posiciones que la hoja alcanza (`ancho`): una hoja más
    angosta que las 28 columnas la procesa el script igual, sin las que le
    faltan (ver `calcular_totales_y_reordenar` y `moneda_ausente`).
    """
    for indice in _ENCABEZADOS_VALIDADOS:
        if indice >= ancho:
            continue
        esperado = CABECERAS_ESTANDAR[indice]
        real = fila[indice] if indice < len(fila) else None
        if real is None or _normalizar(real) != _normalizar(esperado):
            raise ErrorContenido.columna_faltante(archivo=archivo, columna=f"{hoja}: {esperado}")


def _es_vacia(valor: object) -> bool:
    """Lo que el lector de pandas convierte en `""` y recorta al final de la fila."""
    return valor is None or valor == ""


def _leer_hoja(
    hoja: ReadOnlyWorksheet, nombre: str, archivo: str, *, con_filas: bool
) -> _HojaLeida:
    """Valida el encabezado (la primera fila) y devuelve las filas con `Cuenta`.

    `reset_dimensions()` y el recorrido de `hoja.values` son los de pandas en
    modo sólo lectura: las filas vacías intermedias llegan como tuplas vacías,
    así que la posición de cada fila es la de Excel. El encabezado se valida
    contra el ancho de la hoja entera; sin `con_filas` sólo se lee el
    encabezado, así que se valida contra su propio ancho, y una primera fila
    en blanco no se valida mientras no aparezcan datos debajo (con datos, es
    `columna_faltante`).
    """
    hoja.reset_dimensions()
    ancho = 0
    crudas: list[tuple[int, tuple[object, ...]]] = []
    encabezado: tuple[object, ...] = ()
    for posicion, fila in enumerate(hoja.values):
        largo = len(fila)
        while largo and _es_vacia(fila[largo - 1]):
            largo -= 1
        ancho = max(ancho, largo)
        if posicion == 0:
            encabezado = tuple(fila[:largo])
            if not con_filas:
                break
            continue
        if largo < 2:
            continue
        cuenta = _texto(fila[1])
        if cuenta is None or cuenta.strip() == "":
            continue
        crudas.append((posicion + 1, tuple(fila[:NUM_COLUMNAS_PERMITIDAS])))
    if encabezado or crudas:
        _validar_encabezado(encabezado, ancho, nombre, archivo)
    columnas = min(ancho, NUM_COLUMNAS_PERMITIDAS)
    filas = tuple(
        _Fila(
            numero=numero,
            textos=tuple(_texto(valores[i]) if i < len(valores) else None for i in range(columnas)),
        )
        for numero, valores in crudas
    )
    return _HojaLeida(nombre=nombre, columnas=columnas, filas=filas)


def _detectar_tipo(hojas: Sequence[str], archivo: str) -> str:
    """El tipo del libro, decidido en este orden:

    1. Por las hojas distintivas presentes (nombres EXACTOS, como los leía
       `pd.read_excel(sheet_name=...)`): de un solo tipo → ese tipo; de los
       dos → `tipo_no_reconocido`.
    2. Sin ninguna, la regla del script: el nombre contiene "LIMA EXPRESA" o
       "PEX" (sin distinguir mayúsculas, como el `glob` de Windows). Ninguno
       o los dos → `tipo_no_reconocido`.
    """
    por_hojas = [t for t in TIPOS if any(h in _HOJAS_DISTINTIVAS[t] for h in hojas)]
    if len(por_hojas) == 1:
        return por_hojas[0]
    if por_hojas:
        raise ErrorContenido.tipo_no_reconocido(archivo=archivo)
    nombre = _nombre_base(archivo).casefold()
    por_nombre = [t for t in TIPOS if t.casefold() in nombre]
    if len(por_nombre) == 1:
        return por_nombre[0]
    raise ErrorContenido.tipo_no_reconocido(archivo=archivo)


def _leer_libro(entrada: ArchivoEntrada, *, con_filas: bool) -> _Libro:
    """Detecta el tipo, valida los encabezados y, si se pide, lee las filas.

    **Se abre un descriptor y se le pasa el objeto, nunca la ruta**: los
    temporales de `app/recepcion.py` no tienen extensión y `load_workbook`
    rechaza la ruta por eso (con un objeto no mira la extensión, y un `.xlsm`
    se lee igual que un `.xlsx`). Un archivo que ni siquiera abre como Excel
    es `cero_filas`, como en los demás procesadores; una falla DESPUÉS de
    abrirlo no está prevista y se propaga.
    """
    archivo = entrada.nombre_original
    with entrada.ruta_temporal.open("rb") as descriptor:
        try:
            libro = load_workbook(descriptor, read_only=True, data_only=True)
        except Exception as error:
            raise ErrorContenido.cero_filas(archivo=archivo) from error
        try:
            # Antes de cerrar el descriptor: en modo `read_only` el libro lee
            # perezosamente y lo sigue necesitando mientras se itera.
            tipo = _detectar_tipo(libro.sheetnames, archivo)
            configuradas = _hojas_de_tipo(tipo)
            hojas: dict[str, _HojaLeida] = {}
            for nombre in configuradas:
                if nombre in hojas:
                    continue
                hoja = libro[nombre] if nombre in libro.sheetnames else None
                if isinstance(hoja, ReadOnlyWorksheet):
                    # Una hoja de gráfico con el nombre esperado cuenta como ausente.
                    hojas[nombre] = _leer_hoja(hoja, nombre, archivo, con_filas=con_filas)
        finally:
            libro.close()
    if not hojas:
        raise ErrorContenido.hoja_faltante(archivo=archivo, hoja=configuradas[0])
    return _Libro(
        archivo=archivo,
        excel=_nombre_base(archivo),
        tipo=tipo,
        hojas=MappingProxyType(hojas),
    )


def _ordenar(libros: Sequence[_Libro]) -> list[_Libro]:
    """Primero todos los LIMA EXPRESA y después todos los PEX, cada grupo en el
    orden de NTFS (`_clave_ntfs`); a igual nombre, en el orden de subida."""
    return sorted(libros, key=lambda libro: (TIPOS.index(libro.tipo), _clave_ntfs(libro.excel)))


# --- Procesamiento ------------------------------------------------------------


@dataclass(frozen=True, slots=True)
class Observacion:
    """Algo que el script habría hecho en silencio, con dónde y con qué valor.

    `fila` es la de Excel (desde 1), o `None` si la observación es de la hoja
    entera.
    """

    archivo: str
    hoja: str
    fila: int | None
    motivo: str
    valor: str


@dataclass(slots=True)
class _Contexto:
    archivo: str
    observaciones: list[Observacion] = field(default_factory=list)

    def observar(self, hoja: str, fila: int | None, motivo: str, valor: object) -> None:
        self.observaciones.append(
            Observacion(
                archivo=self.archivo,
                hoja=hoja,
                fila=fila,
                motivo=motivo,
                valor="" if valor is None else str(valor),
            )
        )


Registro = dict[str, object]
"""Una fila de la salida, por nombre de columna. Una columna ausente es NaN."""


@dataclass(frozen=True, slots=True)
class _Tramo:
    """El `DataFrame` que `procesar_hoja` devolvía para una hoja."""

    columnas: int
    registros: list[Registro]
    numeros: list[int]


def _es_blanco(texto: str | None) -> bool:
    return texto is None or not texto.strip()


def procesar_hoja(hoja: _HojaLeida, libro: _Libro, contexto: _Contexto) -> _Tramo | None:
    """`procesar_hoja` + `generar_columnas_mes_semana` del script.

    Devuelve `None` donde el script devolvía `None`: menos de dos columnas o
    ninguna fila con `Cuenta`. No se reporta; ver el punto 4 del módulo.
    """
    columnas = CABECERAS_ESTANDAR[: hoja.columnas]
    if hoja.columnas <= 1 or not hoja.filas:
        return None

    registros: list[Registro] = []
    for fila in hoja.filas:
        registro: Registro = dict(zip(columnas, fila.textos, strict=True))
        textos = dict(zip(columnas, fila.textos, strict=True))
        registro["Categoria"] = categorizar(textos[COL_ORDEN]) if COL_ORDEN in textos else "OPEX"
        registro["Sociedad"] = libro.tipo
        registro["Excel"] = libro.excel
        registro["Fuente"] = hoja.nombre
        for columna in COLUMNAS_IMPORTE:
            if columna not in textos:
                continue
            importe = limpiar_importe(textos[columna])
            if importe is None and not _es_blanco(textos[columna]):
                contexto.observar(hoja.nombre, fila.numero, MOTIVO_IMPORTE, textos[columna])
            if COL_CLASE in textos:
                importe = aplicar_signo(importe, textos[COL_CLASE])
            registro[columna] = importe
        registros.append(registro)

    numeros = [fila.numero for fila in hoja.filas]
    if COL_FECHA_PAGO in columnas:
        textos_pago = [f.textos[columnas.index(COL_FECHA_PAGO)] for f in hoja.filas]
        pagos = _fechas_de_columna(textos_pago, dayfirst=False)
        _reportar_fechas(hoja.nombre, numeros, textos_pago, pagos, contexto)
    else:
        pagos = [None] * len(registros)

    if COL_FECHA_PAGO in columnas and COL_BASE in columnas:
        textos_base = [f.textos[columnas.index(COL_BASE)] for f in hoja.filas]
        bases = _fechas_de_columna(
            [None if t is None else t.replace(".", "-") for t in textos_base], dayfirst=True
        )
        _reportar_fechas(hoja.nombre, numeros, textos_base, bases, contexto)
        ppp: list[int | None] = [
            None if pago is None or base is None else (pago - base).days
            for pago, base in zip(pagos, bases, strict=True)
        ]
    else:
        ppp = [None] * len(registros)

    for registro, pago, dias in zip(registros, pagos, ppp, strict=True):
        registro["PPP"] = dias
        if COL_FECHA_PAGO in columnas:
            registro[COL_FECHA_PAGO] = pago
        registro["Mes"] = _mes(pago)
        registro["Semana"] = calcular_semana(pago)
    return _Tramo(columnas=hoja.columnas, registros=registros, numeros=numeros)


def _reportar_fechas(
    hoja: str,
    numeros: Sequence[int],
    textos: Sequence[str | None],
    fechas: Sequence[datetime | None],
    contexto: _Contexto,
) -> None:
    """Una fecha escrita que no se pudo leer deja sin `PPP`, `Mes` o `Semana`."""
    for numero, texto, fecha in zip(numeros, textos, fechas, strict=True):
        if fecha is None and not _es_blanco(texto):
            contexto.observar(hoja, numero, MOTIVO_FECHA, texto)


@dataclass(slots=True)
class _Consolidado:
    """Las listas `consolidados_finales` del script, más quién aportó cada fila."""

    soles: list[_Tramo] = field(default_factory=list)
    dolares: list[_Tramo] = field(default_factory=list)
    filas_por_libro: dict[int, int] = field(default_factory=dict)

    def agregar(self, destino: str, tramo: _Tramo, libro: int) -> None:
        (self.soles if destino == HOJA_SOLES else self.dolares).append(tramo)
        self.filas_por_libro[libro] = self.filas_por_libro.get(libro, 0) + len(tramo.registros)


def _separar_por_moneda(
    tramo: _Tramo, hoja: str, contexto: _Contexto
) -> tuple[_Tramo | None, _Tramo | None]:
    """El `df[df['Moneda del documento'] == 'PEN' / 'USD']` del script (igualdad
    EXACTA). Las filas de otra moneda se pierden y se reportan."""
    if CABECERAS_ESTANDAR.index(COL_MONEDA) >= tramo.columnas:
        contexto.observar(hoja, None, MOTIVO_MONEDA_AUSENTE, "")
        return None, None
    partes: dict[str, tuple[list[Registro], list[int]]] = {
        MONEDA_SOLES: ([], []),
        MONEDA_DOLARES: ([], []),
    }
    for registro, numero in zip(tramo.registros, tramo.numeros, strict=True):
        moneda = registro[COL_MONEDA]
        if isinstance(moneda, str) and moneda in partes:
            partes[moneda][0].append(registro)
            partes[moneda][1].append(numero)
        else:
            contexto.observar(hoja, numero, MOTIVO_MONEDA, moneda)
    soles, dolares = (
        _Tramo(columnas=tramo.columnas, registros=registros, numeros=numeros) if registros else None
        for registros, numeros in (partes[MONEDA_SOLES], partes[MONEDA_DOLARES])
    )
    return soles, dolares


def _consolidar(libros: Sequence[_Libro], contextos: Sequence[_Contexto]) -> _Consolidado:
    """El recorrido `for archivo ... for nombre_hoja, lista_hojas ...` del script.

    `contextos` va alineado con `libros` (por posición: dos subidas pueden
    llamarse igual).
    """
    consolidado = _Consolidado()
    for indice, (libro, contexto) in enumerate(zip(libros, contextos, strict=True)):
        for grupo, nombres in GRUPOS_POR_TIPO[libro.tipo]:
            for nombre in nombres:
                hoja = libro.hojas.get(nombre)
                if hoja is None:
                    contexto.observar(nombre, None, MOTIVO_HOJA_AUSENTE, "")
                    continue
                tramo = procesar_hoja(hoja, libro, contexto)
                if tramo is None:
                    continue
                if grupo != GRUPO_POR_MONEDA:
                    consolidado.agregar(grupo, tramo, indice)
                    continue
                soles, dolares = _separar_por_moneda(tramo, nombre, contexto)
                if soles is not None:
                    consolidado.agregar(HOJA_SOLES, soles, indice)
                if dolares is not None:
                    consolidado.agregar(HOJA_DOLARES, dolares, indice)
    return consolidado


def calcular_totales_y_reordenar(
    tramos: Sequence[_Tramo], nombre_hoja: str
) -> tuple[tuple[str, ...], list[list[object]]]:
    """`pd.concat` + `calcular_totales_semana_y_reordenar` del script.

    Devuelve los títulos y las filas de la hoja de salida. Las columnas
    estándar son las de la hoja más ancha (el `concat` las une en su orden);
    después van las de `COLUMNAS_GENERADAS` que existan.
    """
    registros = [r for tramo in tramos for r in tramo.registros]
    estandar = CABECERAS_ESTANDAR[: max(tramo.columnas for tramo in tramos)]
    columna_importe = COL_IMPORTE_LOCAL if nombre_hoja == HOJA_SOLES else COL_IMPORTE_ML3
    totales: dict[str, list[Importe]] = {}
    if columna_importe not in estandar:
        totales[_TOTAL_HOJA] = [0] * len(registros)
        totales[_TOTAL_MES] = [0] * len(registros)
    else:
        importes = [_importe_o_nulo(r.get(columna_importe)) for r in registros]
        claves = [
            {c: str(r[c]) for c in ("Sociedad", "Fuente", "Categoria", "Mes", "Semana")}
            for r in registros
        ]
        totales[_TOTAL_HOJA] = _totales(
            [(c["Sociedad"], c["Fuente"], c["Mes"], c["Semana"]) for c in claves], importes
        )
        totales[_TOTAL_MES] = _totales(
            [(c["Sociedad"], c["Mes"], c["Semana"]) for c in claves], importes
        )
        totales[_TOTAL_CATEGORIA] = _totales(
            [(c["Sociedad"], c["Categoria"], c["Mes"], c["Semana"]) for c in claves], importes
        )
    generadas = tuple(c for c in COLUMNAS_GENERADAS if c in totales or not c.startswith("Total"))
    titulos = (*estandar, *generadas)
    filas: list[list[object]] = []
    for i, registro in enumerate(registros):
        filas.append(
            [
                totales[titulo][i] if titulo in totales else registro.get(titulo)
                for titulo in titulos
            ]
        )
    return titulos, filas


def _importe_o_nulo(valor: object) -> Importe | None:
    """Estrecha el tipo de un importe ya calculado por este módulo."""
    if valor is None or isinstance(valor, int | float):
        return valor
    raise TypeError(f"importe no numérico: {valor!r}")


# --- Escritura ----------------------------------------------------------------

_BORDE_FINO: Final[Side] = Side(style="thin")


def _celdas_de_salida(hoja: WriteOnlyWorksheet, fila: Sequence[object]) -> list[object]:
    """Las fechas llevan el formato con que `to_excel` las escribía."""
    celdas: list[object] = []
    for valor in fila:
        if isinstance(valor, datetime):
            celda = WriteOnlyCell(hoja, value=valor)
            celda.number_format = FORMATO_FECHA_SALIDA
            celdas.append(celda)
        else:
            celdas.append(valor)
    return celdas


def _encabezado(hoja: WriteOnlyWorksheet, columnas: Sequence[str]) -> list[Cell]:
    """La fila de títulos con el estilo de `to_excel`: negrita, borde fino, centrada."""
    celdas: list[Cell] = []
    for titulo in columnas:
        celda = WriteOnlyCell(hoja, value=titulo)
        celda.font = Font(bold=True)
        celda.border = Border(
            left=_BORDE_FINO, right=_BORDE_FINO, top=_BORDE_FINO, bottom=_BORDE_FINO
        )
        celda.alignment = Alignment(horizontal="center", vertical="top")
        celdas.append(celda)
    return celdas


_HojaDeSalida = tuple[str, Sequence[str], Sequence[Sequence[object]]]
"""Nombre de la hoja, títulos de columna y filas."""


def _escribir_xlsx(hojas: Sequence[_HojaDeSalida], ruta: Path) -> None:
    """Una hoja por salida, con fila de títulos y sin índice (`index=False`).

    `""` se escribe como `""`, igual que pandas: openpyxl lo emite como celda
    vacía. `None` (NaN, NaT) también queda vacía.
    """
    libro = Workbook(write_only=True)
    for nombre, columnas, filas in hojas:
        hoja = libro.create_sheet(title=nombre)
        hoja.append(_encabezado(hoja, columnas))
        for fila in filas:
            hoja.append(_celdas_de_salida(hoja, fila))
    libro.save(ruta)


def _limpio(texto: str) -> str:
    return texto.replace("\t", " ").replace("\r", " ").replace("\n", " ")


def _escribir_observaciones(observaciones: Sequence[Observacion], ruta: Path) -> None:
    """Reporte de observaciones, una por línea, con `FIN_DE_LINEA` explícito."""
    conteo: dict[str, int] = {}
    for observacion in observaciones:
        conteo[observacion.motivo] = conteo.get(observacion.motivo, 0) + 1
    resumen = ", ".join(f"{motivo}={total}" for motivo, total in sorted(conteo.items()))

    lineas = [
        f"Observaciones: {len(observaciones)} ({resumen})",
        "",
        "\t".join(("archivo", "hoja", "fila", "motivo", "valor")),
    ]
    lineas.extend(
        "\t".join(
            (
                _limpio(o.archivo),
                _limpio(o.hoja),
                "" if o.fila is None else str(o.fila),
                o.motivo,
                _limpio(o.valor),
            )
        )
        for o in observaciones
    )
    with ruta.open("w", encoding=CODIFICACION, newline="") as salida:
        salida.write(FIN_DE_LINEA.join(lineas) + FIN_DE_LINEA)


class FlujoCajaPagos(Procesador):
    """El procesador real: de uno a diez libros de pagos programados entran;
    `Consolidado_Pagos.xlsx` sale.

    Más un `observaciones.txt` cuando —y sólo cuando— hubo algo que reportar.
    `app.core.empaquetado` decide si eso viaja suelto o comprimido; este módulo
    nunca arma un ZIP (ADR 0006).
    """

    def __init__(self, clave: str = CLAVE) -> None:
        self.clave = clave

    def validar(self, archivos: list[ArchivoEntrada]) -> ErrorTipificado | None:
        """Tipo y encabezados de cada libro, leyendo SÓLO la primera fila de cada hoja.

        Devuelve el error en vez de levantarlo, como fija `Procesador`. Varios
        libros del mismo tipo son lo normal: no hay `tipo_duplicado`.
        """
        try:
            for entrada in archivos:
                _leer_libro(entrada, con_filas=False)
        except ErrorContenido as error:
            return error
        return None

    def procesar(self, archivos: list[ArchivoEntrada]) -> list[ArchivoSalida]:
        """Calcula todo antes de escribir el primer byte, y devuelve las salidas.

        Ninguna excepción se traga; lo no previsto se propaga y
        `app/core/ejecucion.py` lo convierte en `FalloDelModulo`.
        """
        libros = _ordenar([_leer_libro(entrada, con_filas=True) for entrada in archivos])
        contextos = [_Contexto(archivo=libro.archivo) for libro in libros]
        consolidado = _consolidar(libros, contextos)

        hojas: list[_HojaDeSalida] = []
        destinos = ((HOJA_SOLES, consolidado.soles), (HOJA_DOLARES, consolidado.dolares))
        for nombre, tramos in destinos:
            if tramos:
                titulos, filas = calcular_totales_y_reordenar(tramos, nombre)
                hojas.append((nombre, titulos, filas))
        if not hojas:
            raise ErrorContenido.cero_filas(archivo=", ".join(libro.archivo for libro in libros))

        directorio = archivos[0].ruta_temporal.parent
        ruta_consolidado = directorio / NOMBRE_CONSOLIDADO
        _escribir_xlsx(hojas, ruta_consolidado)
        salidas = [
            ArchivoSalida(
                nombre_propuesto=NOMBRE_CONSOLIDADO,
                ruta_temporal=ruta_consolidado,
                tipo_mime=MIME_XLSX,
            )
        ]

        observaciones = [o for contexto in contextos for o in contexto.observaciones]
        if observaciones:
            ruta_observaciones = directorio / NOMBRE_OBSERVACIONES
            _escribir_observaciones(observaciones, ruta_observaciones)
            salidas.append(
                ArchivoSalida(
                    nombre_propuesto=NOMBRE_OBSERVACIONES,
                    ruta_temporal=ruta_observaciones,
                    tipo_mime=MIME_TEXTO,
                )
            )
            # El logging se configura ACÁ y no al importar: este código corre
            # en el hijo de `spawn`, que no hereda la configuración del padre
            # (ver `contado_carga`). Una línea por archivo con observaciones.
            configurar_logging()
            for indice, (libro, contexto) in enumerate(zip(libros, contextos, strict=True)):
                if contexto.observaciones:
                    registrar_descartes(
                        clave=self.clave,
                        archivo=libro.archivo,
                        filas_procesadas=consolidado.filas_por_libro.get(indice, 0),
                        descartes=[(o.motivo, o.fila or 0) for o in contexto.observaciones],
                    )
        return salidas
