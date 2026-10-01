"""Migración de `flujoCaja_BancoNacion` y `flujoCaja_consolida_Ingresos`.

Traduce los DOS scripts manuales de la raíz del repositorio que el usuario
corría uno detrás del otro sobre la misma carpeta de entrada:

- `flujoCaja_BancoNacion.py` lee la hoja `BCO NACION` del libro de cuentas
  **Operativas** y escribe `Banco_Nacion.xlsx` (hoja `Banco Nacion`).
- `flujoCaja_consolida_Ingresos.py` lee seis hojas del libro de cuentas
  **Fideicomisas** y cuatro del de **Operativas** y escribe
  `Consolidado_Ingresos.xlsx` (hojas `Consolidado Fideicomiso` y
  `Consolidado Operativas`, cada una sólo si tiene filas).

Las dos lógicas siguen SEPARADAS dentro de este módulo —`_banco_nacion` por un
lado; `_fideicomiso` y `_operativas` por el otro— y no comparten reglas: sólo
comparten la ventana del portal. Las reglas de negocio son las de los scripts
y esta migración **no las reinterpreta**: configuración de hojas, posiciones
de columna, `fila_inicio`, filtros, `categorizar`, `limpiar_valor_monetario`,
el cálculo de semanas, los meses abreviados, los totales por grupo, el orden
de columnas, el orden de filas y el redondeo salen de ahí tal cual.

**Entradas y salidas.** Una ejecución recibe uno o dos Excel: a lo sumo uno
de cada tipo (ver `_detectar_tipo`).

- Fideicomisas + Operativas → `Consolidado_Ingresos.xlsx` (las dos hojas) y
  `Banco_Nacion.xlsx`.
- Sólo Operativas → `Consolidado_Ingresos.xlsx` (hoja de Operativas) y
  `Banco_Nacion.xlsx`.
- Sólo Fideicomisas → `Consolidado_Ingresos.xlsx` (hoja de Fideicomiso).

Más un `observaciones.txt` cuando —y sólo cuando— pasó algo que los scripts
perdían o salteaban en silencio. `app.core.empaquetado` decide si todo eso
viaja suelto o en ZIP; este módulo nunca arma un ZIP (ADR 0006).

Lo que cambia respecto de los scripts, por el PRD ("Endurecimiento
obligatorio"):

1. **El tipo de libro se detecta por las HOJAS**, no por el nombre del
   archivo; el nombre (`*Fideicomisas*`, `*Operativas*`, como el `glob` de los
   scripts) es sólo el respaldo. Dos libros del mismo tipo son
   `tipo_duplicado`.
2. **Los encabezados se validan.** Los scripts leen columnas por POSICIÓN a
   partir de una fila fija. Las hojas reales sí tienen una fila de encabezado
   —en una fila fija de cada hoja, relevada de las muestras (`fila_encabezado`
   en `_ConfigHoja`)—, así que se sigue leyendo por posición, como el script,
   pero antes se comprueba que el encabezado de CADA columna leída diga lo
   esperado. Si el extracto cambió de layout, la ejecución muere con
   `columna_faltante` nombrando hoja y columna, en vez de leer otra cosa.
3. **Ninguna excepción se traga.** Desaparecen los `try/except: print` por
   hoja y los `return None`; una falla no prevista se propaga y
   `app/core/ejecucion.py` la convierte en `FalloDelModulo`.
4. **Hoja configurada ausente: se sigue salteando**, como los scripts (es
   normal que un mes no tenga un banco), pero se reporta en
   `observaciones.txt` (`hoja_ausente`), igual que una hoja presente que no
   aporta ninguna fila (`hoja_sin_filas`). Un libro que no aporta ninguna
   fila a ninguna salida es `cero_filas`.
5. **Las pérdidas silenciosas se reportan**: importes ilegibles que los
   scripts convertían en 0 o descartaban (`importe_invalido`), fechas
   ilegibles (`fecha_invalida`) y fechas escritas como texto que se
   interpretan con el mes primero (`fecha_texto`).

**Rarezas de los scripts migradas VERBATIM** —parecen errores, pero
cambiarlas cambia la salida:

- `calcular_semanas_mes`: sólo hay cuatro semanas; la primera termina el
  domingo siguiente si el mes empieza lunes o martes, el domingo de la semana
  SIGUIENTE si empieza de miércoles a sábado, y ocho días después si empieza
  en domingo; la cuarta llega hasta el último día del mes. Una fecha con hora
  el último día cae fuera de toda semana (`""`): se compara el instante, no
  el día.
- Fideicomiso sólo conserva importes estrictamente positivos; Operativas
  descarta los movimientos de hasta 0,01 en valor absoluto.
- Banco de la Nación se lee con `dtype=str`: `F. Operación` sale como TEXTO
  (`"2025.01.02"`), el filtro de letras es ASCII (`[A-Za-z]`) y la fecha se
  interpreta tras cambiar `.` por `-`. `Excel` vale `"OPERATIVA"`, en
  singular; en el consolidado vale `"OPERATIVAS"`.
- Los totales de Operativas se calculan con los importes SIN redondear y
  `Cargos` / `Abonos` se redondean a dos decimales recién al exportar.

**Emulación de pandas.** Este módulo no importa pandas (el hijo de `spawn` lo
reimportaría en cada petición; ver el README de la migración de
`contado_carga`), así que reproduce lo observable:

- `read_excel`: un número entero guardado como `float` se lee como `int`
  (`_celda`), los textos de `_NULOS` y los errores de Excel son NaN, y las
  filas conservan su posición (`iloc[fila_inicio:]` es "desde la fila de
  Excel `fila_inicio + 1`").
- `dtype=str` de Banco de la Nación: `str()` de lo leído (`_como_texto`), de
  modo que una fecha guardada como fecha se vuelve `"2025-01-07 00:00:00"`.
- `pd.to_numeric(errors='coerce')` en `_numero_de_texto`, y
  `pd.to_datetime` SIN `dayfirst` en `_texto_a_fecha` (mes primero; día
  primero sólo cuando el "mes" pasa de 12, como dateutil).
- `groupby().transform('sum')`: suma compensada de Kahan por grupo, en el
  orden de las filas (`_suma_de_grupo`). Ninguna clave de grupo es NaN en
  estos scripts: `Mes`, `Semana` y `Categoría` valen `""` cuando no hay dato,
  así que ninguna fila se queda sin total.
- `sort_values('F. Operación')`: el `quicksort` por defecto NO es estable.
  Sobre `datetime64` numpy usa su introsort genérico (sin la variante SIMD
  de los enteros), que es determinista y no depende de la CPU; se reproduce
  paso a paso en `_orden_de_pandas`, con las fechas nulas al final. Es lo
  que ordena los empates igual que las salidas históricas.
- `Series.round(2)` de numpy: multiplicar por 100, redondear al par y dividir
  (`_redondeo_numpy`), que no es lo mismo que `round(x, 2)` de Python.

**Decisiones de esta migración que los scripts no tomaban:**

- **Fecha ilegible en Banco de la Nación**: el script calculaba el mes con
  `meses_es[m - 1]` sobre una columna que, con un solo `NaT`, pasaba a
  `float`, y moría sin escribir `Banco_Nacion.xlsx`. Acá la fila se conserva
  con `Mes` y `Semana` vacíos —lo mismo que hace el consolidado con una fecha
  ilegible— y se reporta como `fecha_invalida`.
- **Un número en una columna de fecha del consolidado es fecha ilegible.**
  `pd.to_datetime` lo leía como nanosegundos desde 1970 (mes `Ene`,
  `Semana 1` de 1970); acá queda sin fecha y se reporta. Mismo criterio que
  `contado_carga` y `asientos_contables`.
- **Fechas como texto en el consolidado**: pandas infiere un formato para
  toda la columna a partir del PRIMER valor si ese valor es texto; acá cada
  texto se interpreta por separado, mes primero, que es lo que pandas hace
  cuando el primer valor es una fecha de verdad (el caso de las muestras).
  Toda fecha tomada de un texto no ISO se reporta como `fecha_texto`.
- **Las filas sin fecha no se reportan.** Son las líneas de saldo de los
  extractos (`Saldo Final 02-01-2025`, cientos por libro), que el filtro de
  fecha descarta a propósito.
- **Tipo detectado por nombre y ninguna hoja configurada** es
  `hoja_faltante` con la primera hoja de la configuración: el libro no puede
  aportar nada y el error dice qué se esperaba.

Este módulo importa la biblioteca estándar, `openpyxl` y —de `core/`— la
interfaz, los tipos, los errores y el registro operativo. **No importa
`app.core.db`** ni nada que lo alcance (ADR 0013), no crea procesos ni hilos,
y no tiene efecto de importación alguno.
"""

from __future__ import annotations

import math
import re
import unicodedata
from collections.abc import Mapping, Sequence
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from pathlib import Path
from types import MappingProxyType
from typing import Final

from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import Cell
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.worksheet._read_only import ReadOnlyWorksheet
from openpyxl.worksheet._write_only import WriteOnlyWorksheet

from app.core.errores import ErrorContenido, ErrorTipificado
from app.core.interfaz import Procesador
from app.core.registro import configurar_logging, registrar_descartes
from app.core.tipos import ArchivoEntrada, ArchivoSalida

CLAVE: Final[str] = "flujo-caja-ingresos"

TIPO_FIDEICOMISAS: Final[str] = "Fideicomisas"
TIPO_OPERATIVAS: Final[str] = "Operativas"
TIPOS: Final[tuple[str, ...]] = (TIPO_FIDEICOMISAS, TIPO_OPERATIVAS)
"""Los dos tipos de libro, con la subcadena que el `glob` de los scripts
buscaba en el nombre del archivo (`*Fideicomisas*.xlsx`, `*Operativas*.xlsx`)."""

NOMBRE_CONSOLIDADO: Final[str] = "Consolidado_Ingresos.xlsx"
NOMBRE_BANCO_NACION: Final[str] = "Banco_Nacion.xlsx"
NOMBRE_OBSERVACIONES: Final[str] = "observaciones.txt"

HOJA_SALIDA_FIDEICOMISO: Final[str] = "Consolidado Fideicomiso"
HOJA_SALIDA_OPERATIVAS: Final[str] = "Consolidado Operativas"
HOJA_SALIDA_BANCO_NACION: Final[str] = "Banco Nacion"

FIN_DE_LINEA: Final[str] = "\r\n"
"""El de `observaciones.txt` de `asientos_contables`: fijado, no heredado de la
plataforma."""
CODIFICACION: Final[str] = "utf-8"

MIME_XLSX: Final[str] = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MIME_TEXTO: Final[str] = "text/plain; charset=utf-8"

MOTIVO_HOJA_AUSENTE: Final[str] = "hoja_ausente"
MOTIVO_HOJA_SIN_FILAS: Final[str] = "hoja_sin_filas"
MOTIVO_IMPORTE: Final[str] = "importe_invalido"
MOTIVO_FECHA: Final[str] = "fecha_invalida"
MOTIVO_FECHA_TEXTO: Final[str] = "fecha_texto"

MESES_ES: Final[tuple[str, ...]] = (
    "Ene",
    "Feb",
    "Mar",
    "Abr",
    "May",
    "Jun",
    "Jul",
    "Ago",
    "Sep",
    "Oct",
    "Nov",
    "Dic",
)

FORMATO_FECHA_SALIDA: Final[str] = "YYYY-MM-DD HH:MM:SS"
"""El `datetime_format` por defecto con que `to_excel` escribía las fechas."""


# --- Configuración de hojas ---------------------------------------------------


@dataclass(frozen=True, slots=True)
class _ConfigHoja:
    """Una hoja de entrada: dónde leer y qué agregarle a cada fila.

    `columnas` son las del script, contadas desde 1. `fila_inicio` es el
    `iloc[fila_inicio:]` del script (índice desde 0 de la fila de Excel).
    `fila_encabezado` (también desde 0) y `encabezados` los agrega esta
    migración: son la fila y los rótulos que las muestras reales tienen sobre
    cada columna leída, y los valida `_validar_encabezado`.
    """

    nombre: str
    columnas: tuple[int, ...]
    fila_inicio: int
    fila_encabezado: int
    encabezados: tuple[str, ...]
    moneda: str
    excel: str
    categoria: str = ""
    procesamiento: str = ""


_ENCABEZADO_IBK_FID: Final[tuple[str, ...]] = ("Fecha de Op.", "Detalle", "Abonos")
_ENCABEZADO_BCP: Final[tuple[str, ...]] = ("Fecha", "Descripción operación", "Monto")
_ENCABEZADO_BBVA: Final[tuple[str, ...]] = ("F. Operación", "Concepto", "Importe")
_ENCABEZADO_SANTANDER: Final[tuple[str, ...]] = ("Fecha", "Descripcion", "Importe")
_ENCABEZADO_IBK_OPE: Final[tuple[str, ...]] = ("Fecha de Op.", "Detalle", "Cargos", "Abonos")

CARGOS_ABONOS_SEPARADOS: Final[str] = "cargos_abonos_separados"
IMPORTE_UNICO: Final[str] = "importe_unico"

# `obtener_configuracion_fideicomiso` del script, verbatim y en su orden (que es
# el orden de concatenación). Relevado de "LAMSAC - Mov Ctas Fideicomisas 2025
# FL.xlsx": el encabezado está siempre dos filas antes de `fila_inicio`.
HOJAS_FIDEICOMISO: Final[tuple[_ConfigHoja, ...]] = (
    _ConfigHoja(
        "IBK 1106 FID", (1, 4, 8), 16, 14, _ENCABEZADO_IBK_FID, "PEN", "FIDEICOMISO", "ABONO"
    ),
    _ConfigHoja(
        "BCP 306-0-76 FID REC", (1, 3, 4), 6, 4, _ENCABEZADO_BCP, "PEN", "FIDEICOMISO", "ABONO"
    ),
    _ConfigHoja(
        "BCP 921-0-13 FID PT", (1, 3, 4), 6, 4, _ENCABEZADO_BCP, "PEN", "FIDEICOMISO", "ABONO"
    ),
    _ConfigHoja(
        "BBVA FID 43898", (1, 5, 6), 12, 10, _ENCABEZADO_BBVA, "PEN", "FIDEICOMISO", "ABONO"
    ),
    _ConfigHoja(
        "BBVA FID 44606", (1, 5, 6), 12, 10, _ENCABEZADO_BBVA, "PEN", "FIDEICOMISO", "ABONO"
    ),
    _ConfigHoja(
        "SANTANDER 470 FID", (1, 5, 6), 7, 5, _ENCABEZADO_SANTANDER, "PEN", "FIDEICOMISO", "ABONO"
    ),
)

# `obtener_configuracion_operativas`, verbatim. Los comentarios del script dicen
# "Encabezado en fila 7/9", pero en "LAMSAC - Mov Ctas Operativas 2025
# FL_ffff.xlsx" el encabezado está DENTRO del rango leído (índice 14 en IBK,
# 10 en BBVA): el script lo leía como una fila más y la descartaba porque sus
# importes ("Cargos", "Importe") valen 0. Acá esa fila se valida y se saltea,
# que da el mismo resultado.
HOJAS_OPERATIVAS: Final[tuple[_ConfigHoja, ...]] = (
    _ConfigHoja(
        "IBK OPE 223",
        (1, 4, 7, 8),
        7,
        14,
        _ENCABEZADO_IBK_OPE,
        "PEN",
        "OPERATIVAS",
        procesamiento=CARGOS_ABONOS_SEPARADOS,
    ),
    _ConfigHoja(
        "IBK OPE $ 230",
        (1, 4, 7, 8),
        7,
        14,
        _ENCABEZADO_IBK_OPE,
        "USD",
        "OPERATIVAS",
        procesamiento=CARGOS_ABONOS_SEPARADOS,
    ),
    _ConfigHoja(
        "BBVA 43332 PEN",
        (1, 5, 6),
        9,
        10,
        _ENCABEZADO_BBVA,
        "PEN",
        "OPERATIVAS",
        procesamiento=IMPORTE_UNICO,
    ),
    _ConfigHoja(
        "BBVA 43340 USD",
        (1, 5, 6),
        9,
        10,
        _ENCABEZADO_BBVA,
        "USD",
        "OPERATIVAS",
        procesamiento=IMPORTE_UNICO,
    ),
)

# `flujoCaja_BancoNacion.py`: `header=1` (índice 1 de la fila de Excel) e
# `iloc[:, [1, 6, 7]]`, o sea las columnas 2, 7 y 8 contadas desde 1. Los datos
# empiezan en la fila siguiente al encabezado.
HOJA_BANCO_NACION: Final[_ConfigHoja] = _ConfigHoja(
    "BCO NACION", (2, 7, 8), 2, 1, ("Fecha", "Cargo", "Abono"), "PEN", "OPERATIVA"
)

_HOJAS_POR_TIPO: Final[Mapping[str, tuple[_ConfigHoja, ...]]] = MappingProxyType(
    {
        TIPO_FIDEICOMISAS: HOJAS_FIDEICOMISO,
        TIPO_OPERATIVAS: (*HOJAS_OPERATIVAS, HOJA_BANCO_NACION),
    }
)

COLUMNAS_FIDEICOMISO: Final[tuple[str, ...]] = (
    "F. Operación",
    "Concepto",
    "Importe",
    "Moneda",
    "Categoría",
    "Excel",
    "Hoja",
    "Mes",
    "Semana",
    "Total Hoja por Semana",
    "Total Mes por Semana",
)
COLUMNAS_OPERATIVAS: Final[tuple[str, ...]] = (
    "F. Operación",
    "Concepto",
    "Cargos",
    "Abonos",
    "Moneda",
    "Excel",
    "Hoja",
    "Categoría",
    "Mes",
    "Semana",
    "Total Cargos, Hoja por Semana",
    "Total Abonos, Hoja por Semana",
    "Total Cargos, Mes por Semana",
    "Total Abonos, Mes por Semana",
    "Total Cargos, Categoría por Hoja Mes Semana",
    "Total Abonos, Categoría por Hoja Mes Semana",
)
COLUMNAS_BANCO_NACION: Final[tuple[str, ...]] = (
    "F. Operación",
    "Cargos",
    "Abonos",
    "Moneda",
    "Excel",
    "Hoja",
    "Mes",
    "Semana",
    "Total Cargo por Mes, Semana",
    "Total Abono por Mes, Semana",
)


# --- Conversiones que emulan a pandas -----------------------------------------

_NULOS: Final[frozenset[str]] = frozenset(
    {
        "",
        "#N/A",
        "#N/A N/A",
        "#NA",
        "-1.#IND",
        "-1.#QNAN",
        "-NaN",
        "-nan",
        "1.#IND",
        "1.#QNAN",
        "<NA>",
        "N/A",
        "NA",
        "NULL",
        "NaN",
        "None",
        "n/a",
        "nan",
        "null",
        # Errores de Excel: `read_excel` los lee como NaN (`TYPE_ERROR`), y
        # openpyxl en modo `values_only` los entrega como este texto.
        "#DIV/0!",
        "#VALUE!",
        "#REF!",
        "#NAME?",
        "#NUM!",
        "#NULL!",
    }
)
"""Lo que `pd.read_excel` convierte en NaN al leer. Comparado EXACTO, sin
`strip()`, como hace pandas."""

_ESPACIOS_ASCII: Final[str] = " \t\n\r\f\v"
_ENTERO: Final[re.Pattern[str]] = re.compile(r"[+-]?\d+")
_LETRAS: Final[re.Pattern[str]] = re.compile(r"[A-Za-z]")
"""El `str.contains(r'[A-Za-z]')` de Banco de la Nación: sólo letras ASCII."""
_MONEDA: Final[re.Pattern[str]] = re.compile(r"[$\s,]")
"""El `re.sub(r'[$\\s,]', '', ...)` de `limpiar_valor_monetario`: `\\s` es
Unicode, así que también quita el espacio duro."""

Importe = int | float


def _es_nulo(valor: object) -> bool:
    """Lo que pandas habría leído como NaN."""
    if valor is None:
        return True
    if isinstance(valor, float):
        return math.isnan(valor)
    return isinstance(valor, str) and valor in _NULOS


def _celda(valor: object) -> object:
    """El `_convert_cell` del lector openpyxl de pandas: `66.0` se lee `66`."""
    if isinstance(valor, float) and math.isfinite(valor) and valor.is_integer():
        return int(valor)
    return valor


def _como_texto(valor: object) -> str:
    """`dtype=str` sobre un valor NO nulo: `str()` de lo leído."""
    return str(valor)


def _numero_de_texto(texto: str) -> Importe | None:
    """`pd.to_numeric(texto, errors='coerce')`: `None` por NaN.

    Tolera espacios ASCII alrededor (`" 12 "` es 12) pero no el espacio duro,
    ni separadores de miles, ni el `_` que Python sí acepta.
    """
    limpio = texto.strip(_ESPACIOS_ASCII)
    if not limpio or not limpio.isascii() or "_" in limpio:
        return None
    if _ENTERO.fullmatch(limpio):
        return int(limpio)
    try:
        numero = float(limpio)
    except ValueError:
        return None
    return None if math.isnan(numero) else numero


def _es_blanco(valor: object) -> bool:
    """Nulo, o texto hecho sólo de espacios (incluido el espacio duro)."""
    return _es_nulo(valor) or (isinstance(valor, str) and not valor.strip())


def _a_numero(valor: object) -> Importe | None:
    """`pd.to_numeric(errors='coerce')` sobre lo leído por `read_excel`."""
    if _es_nulo(valor):
        return None
    if isinstance(valor, bool):
        return int(valor)
    if isinstance(valor, int | float):
        return valor
    if isinstance(valor, str):
        return _numero_de_texto(valor)
    return None


def _a_numero_de_texto(valor: object) -> Importe | None:
    """`pd.to_numeric(errors='coerce')` sobre una columna leída con `dtype=str`."""
    if _es_nulo(valor) or isinstance(valor, bool):
        return None
    if isinstance(valor, int | float):
        return valor  # `str()` y de vuelta: el mismo número
    return _numero_de_texto(_como_texto(valor))


def limpiar_valor_monetario(valor: object) -> tuple[float, bool]:
    """`limpiar_valor_monetario` del script, verbatim, más si NO pudo convertir.

    El segundo elemento es `True` exactamente donde el script imprimía
    "No se pudo convertir ... Usando 0.": es la pérdida que se reporta.
    """
    if _es_nulo(valor) or valor == "":
        return 0.0, False
    if isinstance(valor, int | float):
        return float(valor), False
    valor_str = str(valor).strip()
    if not valor_str or valor_str.lower() in ("nan", "none", ""):
        return 0.0, False
    es_negativo = False
    if valor_str.startswith("(") and valor_str.endswith(")"):
        es_negativo = True
        valor_str = valor_str[1:-1]
    elif valor_str.startswith("-"):
        es_negativo = True
        valor_str = valor_str[1:]
    valor_limpio = _MONEDA.sub("", valor_str)
    try:
        numero = float(valor_limpio)
    except ValueError:
        return 0.0, True
    return (-numero if es_negativo else numero), False


def categorizar(concepto: object) -> str:
    """`categorizar` del script, verbatim. El orden de las reglas importa."""
    c = ("nan" if _es_nulo(concepto) else str(concepto)).upper()
    if "ITF" in c:
        return "ITF"
    if any(x in c for x in ("COMIS", "COM.", "ENVIO DE EXTRACTO")):
        return "COMISIONES / MANT Y PORTES"
    if any(x in c for x in ("N/A", "N/D", "NOTA DE CARGO", "ABONO REMESAS", "PROSEGUR")):
        return "MONEDAS"
    if any(x in c for x in ("APERTURA PLAZO", "SPOT MER")):
        return "TRANSFERENCIAS"
    if (c.startswith("CASH") and "CASH DEFINIR" not in c) or any(
        x in c
        for x in ("PAGO PROVEEDORES", "PAGO CUOTA SINDICAL", "CONAFOVICER", "SAT LIMA", "C/PH")
    ):
        return "PAGO PROVEEDORES (PEN)"
    if any(x in c for x in ("CASH DEFINIR", "PAGO P. HABERES", "AFPNET")):
        return "PLANILLA"
    if any(x in c for x in ("ABONO", "TIN0", "FACT")):
        return "ABONO"
    return ""


# Formatos de texto que `pd.to_datetime` sin `dayfirst` resuelve. Los de año
# primero no son ambiguos y no se reportan.
_FORMATOS_ANIO_PRIMERO: Final[tuple[str, ...]] = (
    "%Y-%m-%d",
    "%Y-%m-%d %H:%M:%S",
    "%Y-%m-%dT%H:%M:%S",
    "%Y-%m-%d %H:%M:%S.%f",
    "%Y/%m/%d",
    "%Y/%m/%d %H:%M:%S",
)
_FORMATOS_MES_PRIMERO: Final[tuple[str, ...]] = tuple(
    f"%m{s}%d{s}{anio}{hora}"
    for hora in ("", " %H:%M:%S")
    for anio in ("%Y", "%y")
    for s in ("/", "-", ".")
)
_FORMATOS_DIA_PRIMERO: Final[tuple[str, ...]] = tuple(
    f"%d{s}%m{s}{anio}{hora}"
    for hora in ("", " %H:%M:%S")
    for anio in ("%Y", "%y")
    for s in ("/", "-", ".")
)
"""Respaldo de dateutil: `13/01/2025` no tiene mes 13 y se lee día primero."""


def _texto_a_fecha(texto: str) -> tuple[datetime | None, bool]:
    """`pd.to_datetime(texto)` sin `dayfirst`: la fecha y si el texto era ambiguo.

    Devuelve `(None, False)` por `NaT`. El segundo elemento es `True` cuando
    la fecha salió de un formato con el mes o el día primero, que es lo que
    se reporta como `fecha_texto`.
    """
    limpio = texto.strip()
    for formato in _FORMATOS_ANIO_PRIMERO:
        try:
            return datetime.strptime(limpio, formato), False
        except ValueError:
            continue
    for formato in (*_FORMATOS_MES_PRIMERO, *_FORMATOS_DIA_PRIMERO):
        try:
            return datetime.strptime(limpio, formato), True
        except ValueError:
            continue
    return None, False


def calcular_semanas_mes(mes: int, anio: int) -> tuple[tuple[datetime, datetime], ...]:
    """`calcular_semanas_mes` de los scripts, verbatim (ver la docstring del módulo)."""
    primer_dia_mes = datetime(anio, mes, 1)
    ultimo_dia_mes = (primer_dia_mes.replace(day=28) + timedelta(days=4)).replace(
        day=1
    ) - timedelta(days=1)
    dia_semana = primer_dia_mes.weekday()
    if dia_semana == 6:  # domingo
        fin_sem1 = primer_dia_mes + timedelta(days=7)
    elif dia_semana in (2, 3, 4, 5):  # miércoles a sábado
        fin_sem1 = primer_dia_mes + timedelta(days=13 - dia_semana)
    else:  # lunes y martes
        fin_sem1 = primer_dia_mes + timedelta(days=6 - dia_semana)
    semanas = [(primer_dia_mes, fin_sem1)]
    inicio = fin_sem1 + timedelta(days=1)
    for i in range(2, 5):
        if i < 4:
            fin = inicio + timedelta(days=6)
        else:
            fin = inicio + timedelta(days=(ultimo_dia_mes - inicio).days)
        semanas.append((inicio, fin))
        inicio = fin + timedelta(days=1)
    return tuple(semanas)


def calcular_semana(fecha: datetime | None) -> str:
    """`calcular_semana`: `"Semana N"`, o `""` si la fecha es nula o no cae en ninguna."""
    if fecha is None:
        return ""
    for numero, (inicio, fin) in enumerate(calcular_semanas_mes(fecha.month, fecha.year), 1):
        if inicio <= fecha <= fin:
            return f"Semana {numero}"
    return ""


def _mes(fecha: datetime | None) -> str:
    return "" if fecha is None else MESES_ES[fecha.month - 1]


def _suma_de_grupo(valores: Sequence[Importe]) -> Importe:
    """`groupby(...).sum()`: suma compensada de Kahan, como `group_sum` de pandas.

    Con enteros la suma es exacta y se devuelve `int` (mismo número).
    """
    if all(isinstance(v, int) for v in valores):
        return sum(valores)
    total = 0.0
    compensacion = 0.0
    for valor in valores:
        y = valor - compensacion
        t = total + y
        compensacion = t - total - y
        if compensacion != compensacion:  # NaN: pandas la reinicia igual
            compensacion = 0.0
        total = t
    return total


def _totales(claves: Sequence[tuple[str, ...]], valores: Sequence[Importe]) -> list[Importe]:
    """`groupby(claves)[valores].transform('sum')`: el total del grupo en cada fila."""
    grupos: dict[tuple[str, ...], list[Importe]] = {}
    for clave, valor in zip(claves, valores, strict=True):
        grupos.setdefault(clave, []).append(valor)
    sumas = {clave: _suma_de_grupo(grupo) for clave, grupo in grupos.items()}
    return [sumas[clave] for clave in claves]


def _redondeo_numpy(valor: float) -> float:
    """`Series.round(2)`: `rint(x * 100) / 100`, con el redondeo al par de numpy."""
    if not math.isfinite(valor):
        return valor
    return round(valor * 100.0) / 100.0


_SMALL_QUICKSORT: Final[int] = 15
"""`SMALL_QUICKSORT` de `numpy/_core/src/npysort/npysort_common.h`."""


def _orden_de_pandas(fechas: Sequence[datetime | None]) -> list[int]:
    """Índices de `sort_values(kind='quicksort')` sobre una columna de fechas.

    Replica `nargsort` de pandas —las nulas se apartan y van al final, en su
    orden original— y el `aquicksort_` genérico de numpy sobre lo demás. No
    es estable, y ése es el punto: los empates salen como los sacaba pandas.
    """
    validos = [i for i, fecha in enumerate(fechas) if fecha is not None]
    nulos = [i for i, fecha in enumerate(fechas) if fecha is None]
    valores: list[datetime] = [f for f in fechas if f is not None]
    orden = _aquicksort(valores)
    return [validos[i] for i in orden] + nulos


def _aquicksort(v: Sequence[datetime]) -> list[int]:
    """`aquicksort_` de numpy (introsort: quicksort con mediana de tres,
    inserción por debajo de 16 elementos y heapsort si se agota la
    profundidad), traducido línea por línea."""
    num = len(v)
    tosort = list(range(num))
    if num < 2:
        return tosort
    pila: list[tuple[int, int, int]] = []
    pl, pr = 0, num - 1
    cdepth = (num.bit_length() - 1) * 2  # npy_get_msb(num) * 2
    while True:
        if cdepth < 0:
            # Como en C, la profundidad sólo se mira al tomar un tramo de la pila.
            _aheapsort(v, tosort, pl, pr - pl + 1)
        else:
            while pr - pl > _SMALL_QUICKSORT:
                pm = pl + ((pr - pl) >> 1)
                if v[tosort[pm]] < v[tosort[pl]]:
                    tosort[pm], tosort[pl] = tosort[pl], tosort[pm]
                if v[tosort[pr]] < v[tosort[pm]]:
                    tosort[pr], tosort[pm] = tosort[pm], tosort[pr]
                if v[tosort[pm]] < v[tosort[pl]]:
                    tosort[pm], tosort[pl] = tosort[pl], tosort[pm]
                vp = v[tosort[pm]]
                pi = pl
                pj = pr - 1
                tosort[pm], tosort[pj] = tosort[pj], tosort[pm]
                while True:
                    pi += 1
                    while v[tosort[pi]] < vp:
                        pi += 1
                    pj -= 1
                    while vp < v[tosort[pj]]:
                        pj -= 1
                    if pi >= pj:
                        break
                    tosort[pi], tosort[pj] = tosort[pj], tosort[pi]
                pk = pr - 1
                tosort[pi], tosort[pk] = tosort[pk], tosort[pi]
                cdepth -= 1
                # El tramo más grande va a la pila; se sigue con el más chico.
                if pi - pl < pr - pi:
                    pila.append((pi + 1, pr, cdepth))
                    pr = pi - 1
                else:
                    pila.append((pl, pi - 1, cdepth))
                    pl = pi + 1
            for pi in range(pl + 1, pr + 1):  # inserción
                vi = tosort[pi]
                valor = v[vi]
                pj = pi
                pk = pi - 1
                while pj > pl and valor < v[tosort[pk]]:
                    tosort[pj] = tosort[pk]
                    pj -= 1
                    pk -= 1
                tosort[pj] = vi
        if not pila:
            break
        pl, pr, cdepth = pila.pop()
    return tosort


def _aheapsort(v: Sequence[datetime], tosort: list[int], inicio: int, n: int) -> None:
    """`aheapsort_` de numpy sobre `tosort[inicio:inicio + n]` (índices desde 1)."""

    def a(k: int) -> int:
        return tosort[inicio + k - 1]

    def poner(k: int, valor: int) -> None:
        tosort[inicio + k - 1] = valor

    for l_ in range(n >> 1, 0, -1):
        tmp = a(l_)
        i, j = l_, l_ << 1
        while j <= n:
            if j < n and v[a(j)] < v[a(j + 1)]:
                j += 1
            if v[tmp] < v[a(j)]:
                poner(i, a(j))
                i = j
                j += j
            else:
                break
        poner(i, tmp)
    while n > 1:
        tmp = a(n)
        poner(n, a(1))
        n -= 1
        i, j = 1, 2
        while j <= n:
            if j < n and v[a(j)] < v[a(j + 1)]:
                j += 1
            if v[tmp] < v[a(j)]:
                poner(i, a(j))
                i = j
                j += j
            else:
                break
        poner(i, tmp)


# --- Lectura ------------------------------------------------------------------


@dataclass(frozen=True, slots=True)
class _Fila:
    """Una fila leída: su número en Excel (desde 1) y las columnas configuradas."""

    numero: int
    valores: tuple[object, ...]


@dataclass(frozen=True, slots=True)
class _Libro:
    archivo: str
    tipo: str
    hojas: Mapping[str, tuple[_Fila, ...]]
    ausentes: tuple[str, ...]


def _normalizar(texto: object) -> str:
    """Rótulo comparable: sin espacios de más, sin tildes, sin mayúsculas."""
    descompuesto = unicodedata.normalize("NFKD", str(texto))
    sin_tildes = "".join(c for c in descompuesto if not unicodedata.combining(c))
    return " ".join(sin_tildes.split()).casefold()


def _validar_encabezado(fila: Sequence[object], config: _ConfigHoja, archivo: str) -> None:
    """Falla con `columna_faltante` en la primera columna cuyo rótulo no coincide.

    `ContextoContenido` tiene un solo campo libre, así que la columna viaja
    como `"<hoja>: <rótulo esperado>"`.
    """
    for columna, esperado in zip(config.columnas, config.encabezados, strict=True):
        indice = columna - 1
        real = fila[indice] if indice < len(fila) else None
        if real is None or _normalizar(real) != _normalizar(esperado):
            raise ErrorContenido.columna_faltante(
                archivo=archivo, columna=f"{config.nombre}: {esperado}"
            )


def _leer_hoja(
    hoja: ReadOnlyWorksheet, config: _ConfigHoja, archivo: str, *, con_filas: bool
) -> tuple[_Fila, ...]:
    """Valida el encabezado y devuelve las filas desde `fila_inicio`.

    `reset_dimensions()` es lo que hace pandas antes de leer en modo sólo
    lectura. openpyxl entrega las filas vacías intermedias como tuplas
    vacías, así que la posición de cada fila es la de Excel, como en el
    `DataFrame` de `read_excel(header=None)`. Se omiten la fila de encabezado
    (ver `HOJAS_OPERATIVAS`) y las filas cuyas columnas configuradas son
    todas NaN (`dropna(how='all')`). Sin `con_filas` sólo se lee hasta el
    encabezado.
    """
    hoja.reset_dimensions()
    indices = tuple(columna - 1 for columna in config.columnas)
    filas: list[_Fila] = []
    encabezado_visto = False
    for posicion, fila in enumerate(hoja.values):
        if posicion == config.fila_encabezado:
            _validar_encabezado(fila, config, archivo)
            encabezado_visto = True
            if not con_filas:
                break
            continue
        if posicion < config.fila_inicio:
            continue
        largo = len(fila)
        valores = tuple(_celda(fila[i]) if i < largo else None for i in indices)
        if all(_es_nulo(v) for v in valores):
            continue
        filas.append(_Fila(numero=posicion + 1, valores=valores))
    if not encabezado_visto:
        _validar_encabezado((), config, archivo)
    return tuple(filas)


def _detectar_tipo(hojas: Sequence[str], archivo: str) -> str:
    """El tipo del libro, decidido en este orden:

    1. Por las hojas configuradas presentes (nombres EXACTOS, como los leía
       `pd.read_excel(sheet_name=...)`): sólo de Fideicomiso → Fideicomisas;
       sólo de Operativas o `BCO NACION` → Operativas; de los dos →
       `tipo_no_reconocido`.
    2. Sin ninguna, la regla de los scripts: el nombre del archivo contiene
       "Fideicomisas" u "Operativas" (sin distinguir mayúsculas, como el
       `glob` de Windows). Ninguno o los dos → `tipo_no_reconocido`.
    """
    fideicomiso = any(config.nombre in hojas for config in _HOJAS_POR_TIPO[TIPO_FIDEICOMISAS])
    operativas = any(config.nombre in hojas for config in _HOJAS_POR_TIPO[TIPO_OPERATIVAS])
    if fideicomiso and operativas:
        raise ErrorContenido.tipo_no_reconocido(archivo=archivo)
    if fideicomiso:
        return TIPO_FIDEICOMISAS
    if operativas:
        return TIPO_OPERATIVAS
    nombre = archivo.casefold()
    por_nombre = [tipo for tipo in TIPOS if tipo.casefold() in nombre]
    if len(por_nombre) == 1:
        return por_nombre[0]
    raise ErrorContenido.tipo_no_reconocido(archivo=archivo)


def _leer_libro(entrada: ArchivoEntrada, *, con_filas: bool) -> _Libro:
    """Detecta el tipo, valida los encabezados y, si se pide, lee las filas.

    **Se abre un descriptor y se le pasa el objeto, nunca la ruta**: los
    temporales de `app/recepcion.py` no tienen extensión y `load_workbook`
    rechaza la ruta por eso. Un archivo que ni siquiera abre como Excel es
    `cero_filas`, como en los demás procesadores; una falla DESPUÉS de
    abrirlo no está prevista y se propaga.
    """
    archivo = entrada.nombre_original
    with entrada.ruta_temporal.open("rb") as descriptor:
        try:
            libro = load_workbook(descriptor, read_only=True, data_only=True)
        except Exception as error:
            raise ErrorContenido.cero_filas(archivo=archivo) from error
        try:
            # Antes de cerrar el descriptor: en modo `read_only` el libro lee
            # perezosamente y lo sigue necesitando mientras se itera.
            tipo = _detectar_tipo(libro.sheetnames, archivo)
            configuradas = _HOJAS_POR_TIPO[tipo]
            hojas: dict[str, tuple[_Fila, ...]] = {}
            ausentes: list[str] = []
            for config in configuradas:
                hoja = libro[config.nombre] if config.nombre in libro.sheetnames else None
                if not isinstance(hoja, ReadOnlyWorksheet):
                    # Ausente, o una hoja de gráfico con el nombre esperado.
                    ausentes.append(config.nombre)
                    continue
                hojas[config.nombre] = _leer_hoja(hoja, config, archivo, con_filas=con_filas)
        finally:
            libro.close()
    if not hojas:
        raise ErrorContenido.hoja_faltante(archivo=archivo, hoja=configuradas[0].nombre)
    return _Libro(
        archivo=archivo, tipo=tipo, hojas=MappingProxyType(hojas), ausentes=tuple(ausentes)
    )


def _exigir_tipo_nuevo(libro: _Libro, vistos: set[str]) -> None:
    if libro.tipo in vistos:
        raise ErrorContenido.tipo_duplicado(archivo=libro.archivo)
    vistos.add(libro.tipo)


# --- Procesamiento ------------------------------------------------------------


@dataclass(frozen=True, slots=True)
class Observacion:
    """Algo que los scripts habrían hecho en silencio, con dónde y con qué valor.

    `fila` es la de Excel (desde 1), o `None` si la observación es de la hoja
    entera.
    """

    archivo: str
    hoja: str
    fila: int | None
    motivo: str
    valor: str


@dataclass(slots=True)
class _Contexto:
    archivo: str
    observaciones: list[Observacion] = field(default_factory=list)

    def observar(self, hoja: str, fila: int | None, motivo: str, valor: object) -> None:
        self.observaciones.append(
            Observacion(
                archivo=self.archivo,
                hoja=hoja,
                fila=fila,
                motivo=motivo,
                valor="" if valor is None else str(valor),
            )
        )


@dataclass(frozen=True, slots=True)
class _Movimiento:
    """Una fila que pasó los filtros de su hoja, antes de fechas y totales."""

    hoja: str
    fila: int
    fecha: object
    valores: tuple[object, ...]


def _hojas_de(
    libro: _Libro, configs: Sequence[_ConfigHoja], contexto: _Contexto
) -> list[tuple[_ConfigHoja, tuple[_Fila, ...]]]:
    """Las hojas presentes, en el orden de la configuración; las ausentes se reportan."""
    presentes: list[tuple[_ConfigHoja, tuple[_Fila, ...]]] = []
    for config in configs:
        filas = libro.hojas.get(config.nombre)
        if filas is None:
            contexto.observar(config.nombre, None, MOTIVO_HOJA_AUSENTE, "")
        else:
            presentes.append((config, filas))
    return presentes


def _fecha_del_consolidado(movimiento: _Movimiento, contexto: _Contexto) -> datetime | None:
    """`pd.to_datetime(columna, errors='coerce')` de `generar_columnas_mes_semana`."""
    valor = movimiento.fecha
    if isinstance(valor, datetime):
        return valor
    if isinstance(valor, date):
        return datetime(valor.year, valor.month, valor.day)
    if isinstance(valor, str):
        fecha, ambigua = _texto_a_fecha(valor)
        if fecha is None:
            contexto.observar(movimiento.hoja, movimiento.fila, MOTIVO_FECHA, valor)
        elif ambigua:
            contexto.observar(movimiento.hoja, movimiento.fila, MOTIVO_FECHA_TEXTO, valor)
        return fecha
    # Un número (o una hora suelta) no es una fecha: ver la docstring del módulo.
    contexto.observar(movimiento.hoja, movimiento.fila, MOTIVO_FECHA, valor)
    return None


def _fideicomiso(libro: _Libro, contexto: _Contexto) -> list[list[object]]:
    """`procesar_todas_hojas_fideicomiso` + Mes/Semana + totales + orden.

    Devuelve las filas de `Consolidado Fideicomiso`, ya ordenadas, en el orden
    de `COLUMNAS_FIDEICOMISO`.
    """
    movimientos: list[_Movimiento] = []
    for config, filas in _hojas_de(libro, HOJAS_FIDEICOMISO, contexto):
        antes = len(movimientos)
        for fila in filas:
            fecha, concepto, crudo = fila.valores
            if _es_nulo(fecha):
                continue
            importe = _a_numero(crudo)
            if importe is None:
                if not _es_blanco(crudo):
                    contexto.observar(config.nombre, fila.numero, MOTIVO_IMPORTE, crudo)
                continue
            if not importe > 0:  # sólo importes positivos
                continue
            movimientos.append(
                _Movimiento(
                    hoja=config.nombre,
                    fila=fila.numero,
                    fecha=fecha,
                    valores=(
                        None if _es_nulo(concepto) else concepto,
                        importe,
                        config.moneda,
                        config.categoria,
                        config.excel,
                        config.nombre,
                    ),
                )
            )
        if len(movimientos) == antes:
            contexto.observar(config.nombre, None, MOTIVO_HOJA_SIN_FILAS, "")
    if not movimientos:
        return []

    fechas = [_fecha_del_consolidado(m, contexto) for m in movimientos]
    meses = [_mes(fecha) for fecha in fechas]
    semanas = [calcular_semana(fecha) for fecha in fechas]
    importes = [_importe(m.valores[1]) for m in movimientos]
    excel = [str(m.valores[4]) for m in movimientos]
    total_hoja = _totales(
        [
            (e, m.hoja, mes, s)
            for e, m, mes, s in zip(excel, movimientos, meses, semanas, strict=True)
        ],
        importes,
    )
    total_mes = _totales(
        [(e, mes, s) for e, mes, s in zip(excel, meses, semanas, strict=True)], importes
    )
    filas_salida: list[list[object]] = [
        [fecha, *m.valores, mes, semana, t_hoja, t_mes]
        for fecha, m, mes, semana, t_hoja, t_mes in zip(
            fechas, movimientos, meses, semanas, total_hoja, total_mes, strict=True
        )
    ]
    return [filas_salida[i] for i in _orden_de_pandas(fechas)]


def _importe(valor: object) -> Importe:
    """Estrecha el tipo de un importe ya calculado por este módulo."""
    if isinstance(valor, int | float):
        return valor
    raise TypeError(f"importe no numérico: {valor!r}")


def _operativas_de_hoja(
    config: _ConfigHoja, filas: Sequence[_Fila], contexto: _Contexto
) -> list[_Movimiento]:
    """`procesar_hoja_operativas`: cargos y abonos de una hoja, sin fechas ni totales."""
    movimientos: list[_Movimiento] = []
    for fila in filas:
        if _es_nulo(fila.valores[0]):
            continue
        fecha, concepto = fila.valores[0], fila.valores[1]
        if config.procesamiento == CARGOS_ABONOS_SEPARADOS:
            cargo, cargo_invalido = limpiar_valor_monetario(fila.valores[2])
            abono, abono_invalido = limpiar_valor_monetario(fila.valores[3])
            if cargo_invalido:
                contexto.observar(config.nombre, fila.numero, MOTIVO_IMPORTE, fila.valores[2])
            if abono_invalido:
                contexto.observar(config.nombre, fila.numero, MOTIVO_IMPORTE, fila.valores[3])
            cargos = abs(cargo)
            abonos = abs(abono)
            if not (cargos > 0.01 or abonos > 0.01):
                continue
        else:
            importe, invalido = limpiar_valor_monetario(fila.valores[2])
            if invalido:
                contexto.observar(config.nombre, fila.numero, MOTIVO_IMPORTE, fila.valores[2])
            if not abs(importe) > 0.01:
                continue
            cargos = abs(importe) if importe < 0 else 0.0
            abonos = importe if importe > 0 else 0.0
        movimientos.append(
            _Movimiento(
                hoja=config.nombre,
                fila=fila.numero,
                fecha=fecha,
                valores=(
                    None if _es_nulo(concepto) else concepto,
                    cargos,
                    abonos,
                    config.moneda,
                    config.excel,
                    config.nombre,
                    categorizar(concepto),
                ),
            )
        )
    return movimientos


def _operativas(libro: _Libro, contexto: _Contexto) -> list[list[object]]:
    """`procesar_todas_hojas_operativas` + Mes/Semana + seis totales + orden + redondeo.

    Devuelve las filas de `Consolidado Operativas`, en el orden de
    `COLUMNAS_OPERATIVAS`.
    """
    movimientos: list[_Movimiento] = []
    for config, filas in _hojas_de(libro, HOJAS_OPERATIVAS, contexto):
        de_la_hoja = _operativas_de_hoja(config, filas, contexto)
        if not de_la_hoja:
            contexto.observar(config.nombre, None, MOTIVO_HOJA_SIN_FILAS, "")
        movimientos.extend(de_la_hoja)
    if not movimientos:
        return []

    fechas = [_fecha_del_consolidado(m, contexto) for m in movimientos]
    meses = [_mes(fecha) for fecha in fechas]
    semanas = [calcular_semana(fecha) for fecha in fechas]
    cargos = [_importe(m.valores[1]) for m in movimientos]
    abonos = [_importe(m.valores[2]) for m in movimientos]
    monedas = [str(m.valores[3]) for m in movimientos]
    categorias = [str(m.valores[6]) for m in movimientos]
    por_hoja = [(m.hoja, mes, s) for m, mes, s in zip(movimientos, meses, semanas, strict=True)]
    por_moneda = [(mes, s, mon) for mes, s, mon in zip(meses, semanas, monedas, strict=True)]
    por_categoria = [
        (cat, m.hoja, mes, s)
        for cat, m, mes, s in zip(categorias, movimientos, meses, semanas, strict=True)
    ]
    totales = (
        _totales(por_hoja, cargos),
        _totales(por_hoja, abonos),
        _totales(por_moneda, cargos),
        _totales(por_moneda, abonos),
        _totales(por_categoria, cargos),
        _totales(por_categoria, abonos),
    )
    filas_salida: list[list[object]] = []
    for i, (fecha, m) in enumerate(zip(fechas, movimientos, strict=True)):
        concepto, cargo, abono, *resto = m.valores
        filas_salida.append(
            [
                fecha,
                concepto,
                # El redondeo es del export: los totales usan el valor completo.
                _redondeo_numpy(float(_importe(cargo))),
                _redondeo_numpy(float(_importe(abono))),
                *resto,
                meses[i],
                semanas[i],
                *(total[i] for total in totales),
            ]
        )
    return [filas_salida[i] for i in _orden_de_pandas(fechas)]


def _fecha_banco_nacion(texto: str, hoja: str, fila: int, contexto: _Contexto) -> datetime | None:
    """`convertir_fecha_banco_nacion`: `.` por `-` y `pd.to_datetime` sin `dayfirst`."""
    formateada = texto.replace(".", "-")
    fecha, ambigua = _texto_a_fecha(formateada)
    if fecha is None:
        contexto.observar(hoja, fila, MOTIVO_FECHA, texto)
    elif ambigua:
        contexto.observar(hoja, fila, MOTIVO_FECHA_TEXTO, texto)
    return fecha


def _banco_nacion(filas: Sequence[_Fila], contexto: _Contexto) -> list[list[object]]:
    """`flujoCaja_BancoNacion.py`: filas de la hoja `Banco Nacion`, en su orden original.

    `F. Operación` sale como el TEXTO leído con `dtype=str`. Un importe que no
    es número vale 0, como el `to_numeric(errors='coerce').fillna(0)` del
    script; si no estaba en blanco, se reporta.
    """
    hoja = HOJA_BANCO_NACION.nombre
    registros: list[tuple[str, Importe, Importe, datetime | None]] = []
    for fila in filas:
        crudo_fecha, crudo_cargo, crudo_abono = fila.valores
        if _es_nulo(crudo_fecha):
            continue
        texto = _como_texto(crudo_fecha)
        if texto.strip() == "" or _LETRAS.search(texto):
            continue
        importes: list[Importe] = []
        for crudo in (crudo_cargo, crudo_abono):
            numero = _a_numero_de_texto(crudo)
            if numero is None:
                if not _es_blanco(crudo):
                    contexto.observar(hoja, fila.numero, MOTIVO_IMPORTE, crudo)
                numero = 0
            importes.append(numero)
        fecha = _fecha_banco_nacion(texto, hoja, fila.numero, contexto)
        registros.append((texto, importes[0], importes[1], fecha))

    meses = [_mes(fecha) for *_resto, fecha in registros]
    semanas = [calcular_semana(fecha) for *_resto, fecha in registros]
    claves = list(zip(meses, semanas, strict=True))
    total_cargos = _totales(claves, [cargo for _t, cargo, _a, _f in registros])
    total_abonos = _totales(claves, [abono for _t, _c, abono, _f in registros])
    return [
        [
            texto,
            cargo,
            abono,
            HOJA_BANCO_NACION.moneda,
            HOJA_BANCO_NACION.excel,
            hoja,
            meses[i],
            semanas[i],
            total_cargos[i],
            total_abonos[i],
        ]
        for i, (texto, cargo, abono, _fecha) in enumerate(registros)
    ]


# --- Escritura ----------------------------------------------------------------

_BORDE_FINO: Final[Side] = Side(style="thin")


def _celdas_de_salida(hoja: WriteOnlyWorksheet, fila: Sequence[object]) -> list[object]:
    """Las fechas llevan el formato con que `to_excel` las escribía."""
    celdas: list[object] = []
    for valor in fila:
        if isinstance(valor, datetime):
            celda = WriteOnlyCell(hoja, value=valor)
            celda.number_format = FORMATO_FECHA_SALIDA
            celdas.append(celda)
        else:
            celdas.append(valor)
    return celdas


def _encabezado(hoja: WriteOnlyWorksheet, columnas: Sequence[str]) -> list[Cell]:
    """La fila de títulos con el estilo de `to_excel`: negrita, borde fino, centrada."""
    celdas: list[Cell] = []
    for titulo in columnas:
        celda = WriteOnlyCell(hoja, value=titulo)
        celda.font = Font(bold=True)
        celda.border = Border(
            left=_BORDE_FINO, right=_BORDE_FINO, top=_BORDE_FINO, bottom=_BORDE_FINO
        )
        celda.alignment = Alignment(horizontal="center", vertical="top")
        celdas.append(celda)
    return celdas


_HojaDeSalida = tuple[str, Sequence[str], Sequence[Sequence[object]]]
"""Nombre de la hoja, títulos de columna y filas."""


def _salida_xlsx(ruta: Path, hojas: Sequence[_HojaDeSalida]) -> ArchivoSalida:
    _escribir_xlsx(hojas, ruta)
    return ArchivoSalida(nombre_propuesto=ruta.name, ruta_temporal=ruta, tipo_mime=MIME_XLSX)


def _escribir_xlsx(hojas: Sequence[_HojaDeSalida], ruta: Path) -> None:
    """Una hoja por salida, con fila de títulos y sin índice (`index=False`).

    `""` se escribe como `""`, igual que pandas: openpyxl lo emite como celda
    vacía. `None` (NaN, NaT) también queda vacía.
    """
    libro = Workbook(write_only=True)
    for nombre, columnas, filas in hojas:
        hoja = libro.create_sheet(title=nombre)
        hoja.append(_encabezado(hoja, columnas))
        for fila in filas:
            hoja.append(_celdas_de_salida(hoja, fila))
    libro.save(ruta)


def _limpio(texto: str) -> str:
    return texto.replace("\t", " ").replace("\r", " ").replace("\n", " ")


def _escribir_observaciones(observaciones: Sequence[Observacion], ruta: Path) -> None:
    """Reporte de observaciones, una por línea, con `FIN_DE_LINEA` explícito."""
    conteo: dict[str, int] = {}
    for observacion in observaciones:
        conteo[observacion.motivo] = conteo.get(observacion.motivo, 0) + 1
    resumen = ", ".join(f"{motivo}={total}" for motivo, total in sorted(conteo.items()))

    lineas = [
        f"Observaciones: {len(observaciones)} ({resumen})",
        "",
        "\t".join(("archivo", "hoja", "fila", "motivo", "valor")),
    ]
    lineas.extend(
        "\t".join(
            (
                _limpio(o.archivo),
                _limpio(o.hoja),
                "" if o.fila is None else str(o.fila),
                o.motivo,
                _limpio(o.valor),
            )
        )
        for o in observaciones
    )
    with ruta.open("w", encoding=CODIFICACION, newline="") as salida:
        salida.write(FIN_DE_LINEA.join(lineas) + FIN_DE_LINEA)


@dataclass(slots=True)
class _Resultado:
    """Lo calculado para un libro, antes de escribir nada."""

    libro: _Libro
    contexto: _Contexto
    fideicomiso: list[list[object]] = field(default_factory=list)
    operativas: list[list[object]] = field(default_factory=list)
    banco_nacion: list[list[object]] | None = None

    def filas(self) -> int:
        return (
            len(self.fideicomiso)
            + len(self.operativas)
            + (len(self.banco_nacion) if self.banco_nacion is not None else 0)
        )


def _calcular(libro: _Libro) -> _Resultado:
    """Las dos lógicas de un libro, separadas: el consolidado y, si es Operativas,
    Banco de la Nación. Un libro que no aporta ninguna fila es `cero_filas`."""
    resultado = _Resultado(libro=libro, contexto=_Contexto(archivo=libro.archivo))
    if libro.tipo == TIPO_FIDEICOMISAS:
        resultado.fideicomiso = _fideicomiso(libro, resultado.contexto)
    else:
        resultado.operativas = _operativas(libro, resultado.contexto)
        filas_bn = libro.hojas.get(HOJA_BANCO_NACION.nombre)
        if filas_bn is None:
            resultado.contexto.observar(HOJA_BANCO_NACION.nombre, None, MOTIVO_HOJA_AUSENTE, "")
        else:
            resultado.banco_nacion = _banco_nacion(filas_bn, resultado.contexto)
            if not resultado.banco_nacion:
                resultado.contexto.observar(
                    HOJA_BANCO_NACION.nombre, None, MOTIVO_HOJA_SIN_FILAS, ""
                )
    if resultado.filas() == 0:
        raise ErrorContenido.cero_filas(archivo=libro.archivo)
    return resultado


class FlujoCajaIngresos(Procesador):
    """El procesador real: uno o dos Excel entran; el consolidado y Banco de la
    Nación salen.

    Más un `observaciones.txt` cuando —y sólo cuando— hubo algo que reportar.
    `app.core.empaquetado` decide si eso viaja suelto o comprimido; este módulo
    nunca arma un ZIP (ADR 0006).
    """

    def __init__(self, clave: str = CLAVE) -> None:
        self.clave = clave

    def validar(self, archivos: list[ArchivoEntrada]) -> ErrorTipificado | None:
        """Tipo, encabezados y unicidad de tipo, leyendo SÓLO hasta los encabezados.

        Devuelve el error en vez de levantarlo, como fija `Procesador`.
        """
        vistos: set[str] = set()
        try:
            for entrada in archivos:
                _exigir_tipo_nuevo(_leer_libro(entrada, con_filas=False), vistos)
        except ErrorContenido as error:
            return error
        return None

    def procesar(self, archivos: list[ArchivoEntrada]) -> list[ArchivoSalida]:
        """Calcula todo antes de escribir el primer byte, y devuelve las salidas.

        Ninguna excepción se traga; lo no previsto se propaga y
        `app/core/ejecucion.py` lo convierte en `FalloDelModulo`.
        """
        vistos: set[str] = set()
        resultados: list[_Resultado] = []
        for entrada in archivos:
            libro = _leer_libro(entrada, con_filas=True)
            _exigir_tipo_nuevo(libro, vistos)
            resultados.append(_calcular(libro))

        fideicomiso = [f for r in resultados for f in r.fideicomiso]
        operativas = [f for r in resultados for f in r.operativas]
        banco_nacion = next(
            (r.banco_nacion for r in resultados if r.banco_nacion is not None), None
        )

        directorio = archivos[0].ruta_temporal.parent
        salidas: list[ArchivoSalida] = []
        hojas_consolidado: list[_HojaDeSalida] = []
        if fideicomiso:
            hojas_consolidado.append((HOJA_SALIDA_FIDEICOMISO, COLUMNAS_FIDEICOMISO, fideicomiso))
        if operativas:
            hojas_consolidado.append((HOJA_SALIDA_OPERATIVAS, COLUMNAS_OPERATIVAS, operativas))
        if hojas_consolidado:
            salidas.append(_salida_xlsx(directorio / NOMBRE_CONSOLIDADO, hojas_consolidado))
        if banco_nacion is not None:
            # Con cero filas se escribe igual, sólo con títulos: es lo que hacía
            # el script (el caso ya quedó en las observaciones).
            salidas.append(
                _salida_xlsx(
                    directorio / NOMBRE_BANCO_NACION,
                    [(HOJA_SALIDA_BANCO_NACION, COLUMNAS_BANCO_NACION, banco_nacion)],
                )
            )

        observaciones = [o for r in resultados for o in r.contexto.observaciones]
        if observaciones:
            ruta_observaciones = directorio / NOMBRE_OBSERVACIONES
            _escribir_observaciones(observaciones, ruta_observaciones)
            salidas.append(
                ArchivoSalida(
                    nombre_propuesto=NOMBRE_OBSERVACIONES,
                    ruta_temporal=ruta_observaciones,
                    tipo_mime=MIME_TEXTO,
                )
            )
            # El logging se configura ACÁ y no al importar: este código corre
            # en el hijo de `spawn`, que no hereda la configuración del padre
            # (ver `contado_carga`). Una línea por archivo con observaciones.
            configurar_logging()
            for resultado in resultados:
                if resultado.contexto.observaciones:
                    registrar_descartes(
                        clave=self.clave,
                        archivo=resultado.libro.archivo,
                        filas_procesadas=resultado.filas(),
                        descartes=[
                            (o.motivo, o.fila or 0) for o in resultado.contexto.observaciones
                        ],
                    )
        return salidas
