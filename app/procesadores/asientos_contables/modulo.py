"""Migración de `AsientosContables_Carga` (PRD "Endurecimiento obligatorio").

Traduce `AsientosContables_Carga.py` —el script manual de la raíz del
repositorio— a un `Procesador` de este servicio, siguiendo el precedente de
`app/procesadores/contado_carga/`. Las reglas de negocio son las de ese
archivo y esta migración **no las reinterpreta**: cuentas, claves contables,
códigos de documento, descripciones, signos y el layout de 26 columnas salen
de ahí tal cual.

**Qué hace.** Cada Excel subido es de UN tipo (AMEX, DINNERS, MASTERCARD,
SAFETYPAY, VISA, IZIPAY o EFECTIVO) y produce UN libro
`resultado_<tipo>_procesado.xlsx` con las hojas COMISION, COBRANZA y RECARGA
—sólo las que tengan filas, en ese orden—, sin encabezado. Cada asiento es
una línea tipo 1 (cabecera), una o más tipo 2 (detalle) y una separadora en
blanco. Una ejecución acepta hasta siete archivos, uno por tipo;
`app.core.empaquetado` decide si la salida viaja suelta o en ZIP.

Lo que cambia respecto del script, porque el PRD lo exige:

1. **El tipo se detecta por las HOJAS, no por el nombre del archivo.** El
   script buscaba `*AMEX*.xlsx` en una carpeta; acá el nombre que manda el
   cliente es sólo el último recurso (ver `_detectar_tipo`). Dos archivos del
   mismo tipo en una ejecución son `tipo_duplicado`: sus salidas chocarían.
2. **Columnas por nombre, no por posición.** Los encabezados varían entre
   archivos para la misma posición (`OBS`/`Ajuste`/`Referencia`, `Fecha
   Cierre`/`Fecha cierre`), así que cada columna acepta alias, comparados
   tras `strip()` y `casefold()`. Leer por nombre elimina además un bug real
   del script: agrega las columnas auxiliares `fecha` y `monto` ANTES de
   comprobar `len(columns) >= 7`, de modo que en una hoja angosta
   `iloc[:, 6]` leía el MONTO como fecha de cierre.
3. **Hoja o columna obligatoria ausente es error tipificado**
   (`hoja_faltante`, `columna_faltante`), no un `print` y una salida con
   hojas de menos.
4. **Las pérdidas silenciosas se cuentan y se reportan** en
   `observaciones.txt` (sólo si las hay) y en el log: filas que pasaron el
   filtro y se perdieron por fecha ilegible, fechas de cierre ilegibles
   reemplazadas por hoy, categorías que el filtro acepta pero el bucle
   ignora, y categorías de comisión IZIPAY sin cuenta, que se escriben con
   cuenta 0 (se conserva el 0 por paridad).
5. **Ninguna excepción se traga.** Desaparecen los `try/except: print` por
   hoja. Un archivo sin ninguna fila de salida es `cero_filas`, y cualquier
   falla hace fallar la ejecución entera: la API no tiene "éxito parcial".

**La fecha de cierre cae a HOY, y eso ES la regla de negocio.** En los
archivos reales la columna de fecha de cierre viene casi siempre vacía y el
script usa `datetime.now()`: las siete salidas históricas llevan 20250826,
el día en que se corrió (una sola cabecera AMEX trae un 20250701 real). Se
conserva, con el reloj inyectable (`AsientosContables(hoy=...)`) para que las
pruebas y la paridad fijen el día.

**Rarezas del script migradas VERBATIM** —parecen errores, pero cambiarlas
cambia asientos:

- La clave `"01"` es TEXTO y las demás (40, 50, 11, 31, 21) son enteros; el
  Excel de salida conserva esa diferencia.
- Comisión estándar: el filtro es `Comisión|Extorno` sin distinguir
  mayúsculas, pero el bucle exige `"Comisión"` / `"Extorno"` EXACTOS, así
  que "comisión" en minúscula pasa el filtro y no genera línea (sí puede
  generar una cabecera vacía y aportar la fecha de cierre del grupo).
- COBRANZA estándar filtra con `contains('DZ')` (subcadena, sin distinguir
  mayúsculas); la de IZIPAY exige `== "DZ"`.
- RECARGA estándar filtra `RG` y `AJUSTE` como subcadenas; la de IZIPAY sólo
  exige `== "AJUSTE"` y NO mira la clase de documento.
- RECARGA estándar trata el cero como positivo (`>= 0`); la de IZIPAY, como
  negativo (`> 0`).
- IZIPAY COBRANZA: MASTERCARD e IZIPAY se suman juntos y su signo INVERTIDO
  elige la clave (negativo → `"01"` sin `Z001`, si no → 11 con `Z001`). Las
  categorías suman TODAS las filas del día, no sólo las DZ.
- IZIPAY COMISION agrupa TODAS las filas por fecha, así que emite cabecera y
  separadora para cada día aunque no tenga comisiones; y toma la fecha de
  cierre de la PRIMERA fila del día, no de la primera no vacía.
- Los importes NO se redondean.
- Las fechas se leen con el día primero (`dayfirst=True`) salvo en las tres
  lecturas de IZIPAY que el script hace con `pd.to_datetime` pelado.

**Emulación de pandas.** Este módulo no importa pandas (el hijo de `spawn`
lo reimportaría en cada petición; ver el README de la migración de
`contado_carga`), así que reproduce lo observable: `limpiar_montos` en
`_a_importe`, los textos que `read_excel` convierte en NaN en `_NULOS`, el
`astype(str)` de los filtros en `_como_texto`, la suma compensada (Kahan)
del `groupby().sum()` y la suma por pares de numpy del `Series.sum()`. Los
grupos se ordenan por fecha ascendente y la fecha nula no forma grupo, como
en `groupby`.

**Decisiones de esta migración que el script no tomaba:**

- **La columna de fecha de cierre es OBLIGATORIA** en las tres clases de
  hoja. Las siete muestras reales la traen siempre. El script caía a
  `now()` si faltaba —o, por el bug del punto 2, leía otra cosa—; acá un
  encabezado renombrado que trajera fechas reales se reemplazaría en
  silencio por hoy, que es justo el modo de fallo que el endurecimiento
  viene a cerrar.
- **Donde el script se caía, acá se pierde UNA fila y se reporta.** Una
  RECARGA con fecha ilegible (`NaT.strftime`), o cualquier fecha ilegible
  en las tres lecturas estrictas de IZIPAY, hacía que el script perdiera la
  HOJA de salida entera con un `print`. Acá la fila va a
  `observaciones.txt` con motivo `fecha_invalida` y el resto sigue. Una
  fecha de cierre ilegible cae a hoy (como el `errors='coerce'` del resto
  del script) y se reporta como `fecha_cierre_invalida`. Ninguna de las
  siete muestras reales cae en estos casos.
- **Un número en una columna de fecha es fecha ilegible.** `pd.to_datetime`
  lo interpretaba como nanosegundos desde 1970 y producía un asiento con
  fecha 19700101; acá se reporta. Mismo criterio que `contado_carga`.

**Código muerto del script que NO se migró:** `procesar_por_fecha` nunca se
llama, y `procesar_hoja_ajuste` está definida dos veces, idénticas (vale la
segunda). `buscar_archivos` y `main` son la plomería de carpetas que este
servicio reemplaza.

Este módulo importa la biblioteca estándar, `openpyxl` y —de `core/`— la
interfaz, los tipos, los errores y el registro operativo. **No importa
`app.core.db`** ni nada que lo alcance (ADR 0013), no crea procesos ni hilos,
y no tiene efecto de importación alguno.
"""

from __future__ import annotations

import math
import re
from collections.abc import Callable, Iterable, Mapping, Sequence
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from types import MappingProxyType
from typing import Any, Final

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet._read_only import ReadOnlyWorksheet

from app.core.errores import ErrorContenido, ErrorTipificado
from app.core.interfaz import Procesador
from app.core.registro import configurar_logging, registrar_descartes
from app.core.tipos import ArchivoEntrada, ArchivoSalida

CLAVE: Final[str] = "asientos-contables"

TIPOS_DE_TARJETA: Final[tuple[str, ...]] = (
    "AMEX",
    "DINNERS",
    "MASTERCARD",
    "SAFETYPAY",
    "VISA",
    "IZIPAY",
)
"""Los tipos cuya hoja principal se llama igual que el tipo. "DINNERS" lleva
la doble N del script y de los archivos reales; la cuenta IZIPAY de más abajo
dice "DINERS" con una sola, y las dos se conservan como están."""

TIPO_IZIPAY: Final[str] = "IZIPAY"
TIPO_EFECTIVO: Final[str] = "EFECTIVO"
TIPOS: Final[tuple[str, ...]] = (*TIPOS_DE_TARJETA, TIPO_EFECTIVO)
"""Los siete, en el orden en que el script los buscaba por nombre de archivo."""

HOJA_COMPENSACION: Final[str] = "COMPENSACION"
HOJA_EFECTIVO: Final[str] = "Efectivo"
"""Hoja que identifica un archivo de efectivo, comparada sin distinguir
mayúsculas. NO se lee: el script sólo procesa `COMPENSACION` para EFECTIVO."""

HOJA_COMISION: Final[str] = "COMISION"
HOJA_COBRANZA: Final[str] = "COBRANZA"
HOJA_RECARGA: Final[str] = "RECARGA"

COLUMNAS_POR_FILA: Final[int] = 26

CONFIGURACIONES: Final[Mapping[str, Mapping[str, Mapping[str, int]]]] = MappingProxyType(
    {
        "AMEX": {
            "COMISION": {
                "linea2_campo14": 1106012033,
                "linea3_campo13": 50,
                "linea3_campo14": 1101020065,
            },
            "COMPENSACION": {
                "linea2_campo14": 1101020064,
                "linea3_campo13": 11,
                "linea3_campo14": 1001420,
            },
            "AJUSTE": {"campo14": 1001420},
        },
        "DINNERS": {
            "COMISION": {
                "linea2_campo14": 1106012034,
                "linea3_campo13": 50,
                "linea3_campo14": 1101020065,
            },
            "COMPENSACION": {
                "linea2_campo14": 1101020064,
                "linea3_campo13": 11,
                "linea3_campo14": 1001418,
            },
            "AJUSTE": {"campo14": 1001418},
        },
        "MASTERCARD": {
            "COMISION": {
                "linea2_campo14": 1106012032,
                "linea3_campo13": 50,
                "linea3_campo14": 1101020065,
            },
            "COMPENSACION": {
                "linea2_campo14": 1101020064,
                "linea3_campo13": 11,
                "linea3_campo14": 1001417,
            },
            "AJUSTE": {"campo14": 1001417},
        },
        "SAFETYPAY": {
            "COMISION": {
                "linea2_campo14": 1106012031,
                "linea3_campo13": 50,
                "linea3_campo14": 1101020065,
            },
            "COMPENSACION": {
                "linea2_campo14": 1101020064,
                "linea3_campo13": 11,
                "linea3_campo14": 1001419,
            },
            "AJUSTE": {"campo14": 1001419},
        },
        "VISA": {
            "COMISION": {
                "linea2_campo14": 1106012030,
                "linea3_campo13": 50,
                "linea3_campo14": 1101020065,
            },
            "COMPENSACION": {
                "linea2_campo14": 1101020064,
                "linea3_campo13": 11,
                "linea3_campo14": 1001416,
            },
            "AJUSTE": {"campo14": 1001416},
        },
        "EFECTIVO": {
            "COMPENSACION": {
                "linea2_campo14": 1101020064,
                "linea3_campo13": 11,
                "linea3_campo14": 1001415,
            },
            "AJUSTE": {"campo14": 1001415},
        },
    }
)
"""Copiado verbatim del script. IZIPAY no figura: sus cuentas son las dos
tablas de abajo y las constantes `CUENTA_*`."""

CODIGO_CATEGORIAS_IZIPAY: Final[Mapping[str, int]] = MappingProxyType(
    {"MC/VISA": 1106012032, "Diners": 1106012034, "Amex": 1106012033}
)
"""Subcadena EXACTA de la categoría → cuenta. El orden importa: gana la
primera que aparezca. Sin coincidencia, el script escribe cuenta 0."""

CODIGO_COMPENSACION_IZIPAY: Final[Mapping[str, int]] = MappingProxyType(
    {"VISA": 1001416, "DINERS": 1001418, "AMEX": 1001420, "MASTERCARD+IZIPAY": 1001417}
)
"""Texto EXACTO → cuenta, en este orden de emisión. `MASTERCARD+IZIPAY` no es
un texto: agrupa las filas cuyo texto es `MASTERCARD` o `IZIPAY`."""

_MASTERCARD_MAS_IZIPAY: Final[str] = "MASTERCARD+IZIPAY"
_TEXTOS_MASTERCARD_MAS_IZIPAY: Final[frozenset[str]] = frozenset({"MASTERCARD", "IZIPAY"})

CUENTA_BANCO: Final[int] = 1101020064
CUENTA_COMISION_POR_PAGAR: Final[int] = 1101020065
CUENTA_EXTORNO: Final[int] = 2107808010
CUENTA_RECARGA: Final[int] = 3000013956
CUENTA_RECARGA_IZIPAY: Final[int] = 1001417

NOMBRE_OBSERVACIONES: Final[str] = "observaciones.txt"
FIN_DE_LINEA: Final[str] = "\r\n"
"""El del `descartes.txt` de `contado_carga`, por la misma razón: fijado, no
heredado de la plataforma."""
CODIFICACION: Final[str] = "utf-8"

MIME_XLSX: Final[str] = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MIME_TEXTO: Final[str] = "text/plain; charset=utf-8"

MOTIVO_FECHA: Final[str] = "fecha_invalida"
MOTIVO_FECHA_CIERRE: Final[str] = "fecha_cierre_invalida"
MOTIVO_CATEGORIA: Final[str] = "categoria_ignorada"
MOTIVO_CUENTA: Final[str] = "cuenta_no_mapeada"


def nombre_de_salida(tipo: str) -> str:
    """El nombre que el script le daba a la salida de cada tipo."""
    return f"resultado_{tipo.lower()}_procesado.xlsx"


# --- Columnas -----------------------------------------------------------------


@dataclass(frozen=True, slots=True)
class _Columna:
    """Un campo que el módulo lee, con los encabezados que acepta para él.

    Los alias se comparan tras `strip()` + `casefold()`. Si varios aparecen en
    la misma hoja gana el PRIMERO de esta tupla, y entre columnas con el mismo
    encabezado, la de más a la izquierda. El primer alias es el que nombra el
    error `columna_faltante`.
    """

    campo: str
    alias: tuple[str, ...]


_ALIAS_OBSERVACION: Final[tuple[str, ...]] = (
    "OBS",
    "Ajuste",
    "OBSERVACIÓN",
    "OBSERVACION",
    "Referencia",
)
"""La sexta columna de `COMPENSACION`, que es la que marca los `AJUSTE`. Su
encabezado real varía: `OBS` (AMEX, DINNERS, MASTERCARD, SAFETYPAY,
EFECTIVO), `Ajuste` (VISA), `Referencia` (IZIPAY); `OBSERVACIÓN` aparece en
otras hojas de los mismos libros. `OBS` va primero porque, donde conviven
`OBS` y `REFERENCIA`, la posición que el script leía es la de `OBS`."""

_FECHA_CIERRE: Final[_Columna] = _Columna("fecha_cierre", ("Fecha Cierre",))

# Relevados de las siete muestras reales (agosto de 2025). Las posiciones
# coinciden con los `iloc` del script: 0, 2, 3 y 4 en la hoja de la tarjeta;
# 1, 2, 3, 5 y 6 en COMPENSACION; y 1 a 7 en la COMPENSACION de IZIPAY.
_COLUMNAS_COMISION: Final[tuple[_Columna, ...]] = (
    _Columna("fecha", ("Fecha1",)),
    _Columna("monto", ("Monto",)),
    _Columna("categoria", ("Categoria", "Categoría")),
    _FECHA_CIERRE,
)
_COLUMNAS_COMPENSACION: Final[tuple[_Columna, ...]] = (
    _Columna("clase", ("Clase de documento",)),
    _Columna("fecha", ("Fecha de documento",)),
    _Columna("importe", ("Importe en moneda local",)),
    _Columna("observacion", _ALIAS_OBSERVACION),
    _FECHA_CIERRE,
)
_COLUMNAS_COMPENSACION_IZIPAY: Final[tuple[_Columna, ...]] = (
    _Columna("clase", ("Clase de documento",)),
    _Columna("fecha", ("Fecha de documento",)),
    _Columna("importe", ("Importe en moneda local",)),
    _Columna("texto", ("Texto",)),
    _Columna("observacion", _ALIAS_OBSERVACION),
    _Columna("fecha_dia", ("fecha dia", "fecha día")),
    _FECHA_CIERRE,
)


def _hojas_requeridas(tipo: str) -> tuple[tuple[str, tuple[_Columna, ...]], ...]:
    """Las hojas de entrada que cada tipo necesita, con sus columnas."""
    if tipo == TIPO_EFECTIVO:
        return ((HOJA_COMPENSACION, _COLUMNAS_COMPENSACION),)
    if tipo == TIPO_IZIPAY:
        return (
            (TIPO_IZIPAY, _COLUMNAS_COMISION),
            (HOJA_COMPENSACION, _COLUMNAS_COMPENSACION_IZIPAY),
        )
    return ((tipo, _COLUMNAS_COMISION), (HOJA_COMPENSACION, _COLUMNAS_COMPENSACION))


def _normalizar(encabezado: object) -> str:
    return str(encabezado).strip().casefold()


def _indices_de_columnas(
    encabezado: Sequence[object], columnas: Sequence[_Columna], archivo: str
) -> dict[str, int]:
    """Mapa campo → índice. Falla en la PRIMERA columna que falte.

    El contexto del error nombra la columna y no la hoja: `ContextoContenido`
    tiene un solo campo libre. El único nombre compartido por las dos hojas
    de un libro es `Fecha Cierre`.
    """
    posiciones: dict[str, int] = {}
    for indice, nombre in enumerate(encabezado):
        if nombre is not None:
            posiciones.setdefault(_normalizar(nombre), indice)
    indices: dict[str, int] = {}
    for columna in columnas:
        encontrado = next(
            (posiciones[a] for a in map(_normalizar, columna.alias) if a in posiciones), None
        )
        if encontrado is None:
            raise ErrorContenido.columna_faltante(archivo=archivo, columna=columna.alias[0])
        indices[columna.campo] = encontrado
    return indices


# --- Conversiones que emulan a pandas -----------------------------------------

_NULOS: Final[frozenset[str]] = frozenset(
    {
        "",
        "#N/A",
        "#N/A N/A",
        "#NA",
        "-1.#IND",
        "-1.#QNAN",
        "-NaN",
        "-nan",
        "1.#IND",
        "1.#QNAN",
        "<NA>",
        "N/A",
        "NA",
        "NULL",
        "NaN",
        "None",
        "n/a",
        "nan",
        "null",
    }
)
"""Los textos que `pd.read_excel` convierte en NaN al leer (`na_values` por
defecto). Comparados EXACTOS, sin `strip()`, como hace pandas."""

_FORMATOS_DIA_PRIMERO: Final[tuple[str, ...]] = (
    "%d/%m/%Y",
    "%d-%m-%Y",
    "%d.%m.%Y",
    "%d/%m/%y",
    "%d-%m-%y",
    "%d.%m.%y",
    "%Y-%m-%d",
    "%Y/%m/%d",
    "%d/%m/%Y %H:%M:%S",
    "%Y-%m-%d %H:%M:%S",
)
"""`dayfirst=True`: `01/02/2025` es el 1 de febrero. Un `.xlsx` bien formado
entrega `datetime` y no pasa por acá; esto cubre la fecha guardada como texto
(hay algunas en COMPENSACION de MASTERCARD y SAFETYPAY)."""

_FORMATOS_MES_PRIMERO: Final[tuple[str, ...]] = (
    "%m/%d/%Y",
    "%m-%d-%Y",
    "%m.%d.%Y",
    "%m/%d/%y",
    "%m-%d-%y",
    "%m.%d.%y",
    "%Y-%m-%d",
    "%Y/%m/%d",
    "%m/%d/%Y %H:%M:%S",
    "%Y-%m-%d %H:%M:%S",
    *_FORMATOS_DIA_PRIMERO,
)
"""`pd.to_datetime` pelado (las lecturas de IZIPAY): mes primero, y día
primero sólo como respaldo cuando el "mes" pasa de 12, como hace dateutil."""

# Los `str.contains(..., case=False)` del script: regex con IGNORECASE.
_FILTRO_COMISION: Final[re.Pattern[str]] = re.compile("Comisión|Extorno", re.IGNORECASE)
_FILTRO_DZ: Final[re.Pattern[str]] = re.compile("DZ", re.IGNORECASE)
_FILTRO_RG: Final[re.Pattern[str]] = re.compile("RG", re.IGNORECASE)
_FILTRO_AJUSTE: Final[re.Pattern[str]] = re.compile("AJUSTE", re.IGNORECASE)

Importe = int | float


def _es_nulo(valor: object) -> bool:
    """Lo que pandas habría leído como NaN."""
    if valor is None:
        return True
    if isinstance(valor, float):
        return math.isnan(valor)
    return isinstance(valor, str) and valor in _NULOS


def _como_texto(valor: object) -> str:
    """Equivalente de `astype(str)`: lo nulo es `"nan"`, lo demás `str()`."""
    return "nan" if _es_nulo(valor) else str(valor)


def _a_importe(valor: object) -> Importe:
    """Equivalente de `limpiar_montos`: quita comas, y lo que no parsea es 0.

    Mismo criterio que `contado_carga._a_importe`, con dos ajustes de
    fidelidad: `"nan"` y los demás nulos de pandas dan 0 (no NaN), y `1_000`
    no parsea (Python lo acepta, pandas no). Un valor entero se devuelve como
    `int`, que es lo que pandas escribía para las columnas sin decimales; el
    valor numérico es el mismo en cualquier caso.
    """
    if _es_nulo(valor) or isinstance(valor, bool):
        return 0
    if isinstance(valor, int):
        return valor
    if isinstance(valor, float):
        return int(valor) if valor.is_integer() else valor
    texto = str(valor).replace(",", "").strip()
    if texto in {"", "-"} or "_" in texto:
        return 0
    try:
        numero = float(texto)
    except ValueError:
        return 0
    if math.isnan(numero):
        return 0
    return int(numero) if numero.is_integer() else numero


def _a_fecha(valor: object, *, dia_primero: bool) -> datetime | None:
    """Equivalente de `pd.to_datetime(valor, errors='coerce')`: `None` por `NaT`.

    Conserva la hora: el script agrupa por `Timestamp`, y dos horas distintas
    del mismo día son dos grupos (dos cabeceras). Un número no es fecha (ver
    la docstring del módulo).
    """
    if isinstance(valor, datetime):
        return valor
    if isinstance(valor, date):
        return datetime(valor.year, valor.month, valor.day)
    if not isinstance(valor, str) or _es_nulo(valor):
        return None
    texto = valor.strip()
    for formato in _FORMATOS_DIA_PRIMERO if dia_primero else _FORMATOS_MES_PRIMERO:
        try:
            return datetime.strptime(texto, formato)
        except ValueError:
            continue
    return None


def _suma_de_grupo(valores: Sequence[Importe]) -> Importe:
    """`groupby(...).sum()`: suma compensada de Kahan, como `group_sum` de pandas.

    Con enteros la suma es exacta y se devuelve `int`.
    """
    if all(isinstance(v, int) for v in valores):
        return sum(valores)
    total = 0.0
    compensacion = 0.0
    for valor in valores:
        y = valor - compensacion
        t = total + y
        compensacion = t - total - y
        if compensacion != compensacion:  # NaN: pandas la reinicia igual
            compensacion = 0.0
        total = t
    return total


def _suma_de_serie(valores: Sequence[Importe]) -> Importe:
    """`Series.sum()`: la suma por pares de numpy (bloques de 8 y de 128)."""
    if all(isinstance(v, int) for v in valores):
        return sum(valores)
    return _suma_por_pares([float(v) for v in valores])


def _suma_por_pares(valores: Sequence[float]) -> float:
    cantidad = len(valores)
    if cantidad < 8:
        resultado = 0.0
        for valor in valores:
            resultado += valor
        return resultado
    if cantidad <= 128:
        parciales = list(valores[:8])
        indice = 8
        while indice < cantidad - cantidad % 8:
            for desplazamiento in range(8):
                parciales[desplazamiento] += valores[indice + desplazamiento]
            indice += 8
        resultado = ((parciales[0] + parciales[1]) + (parciales[2] + parciales[3])) + (
            (parciales[4] + parciales[5]) + (parciales[6] + parciales[7])
        )
        for valor in valores[indice:]:
            resultado += valor
        return resultado
    mitad = cantidad // 2
    mitad -= mitad % 8
    return _suma_por_pares(valores[:mitad]) + _suma_por_pares(valores[mitad:])


# --- Líneas de salida ---------------------------------------------------------


def _linea_base(
    *, codigo: str, fecha: date, fecha_cierre: date, descripcion: str, tipo: str
) -> list[Any]:
    """`crear_linea_base` del script, tipo 1. El año y el mes son los del CIERRE."""
    linea: list[Any] = [""] * COLUMNAS_POR_FILA
    linea[0] = 1
    linea[2] = "PEXP"
    linea[4] = fecha_cierre.year
    linea[5] = codigo
    linea[6] = fecha.strftime("%Y%m%d")
    linea[7] = fecha_cierre.strftime("%Y%m%d")
    linea[8] = fecha_cierre.strftime("%m")
    linea[9] = descripcion
    linea[10] = tipo
    linea[11] = "PEN"
    return linea


def _linea_detalle(campo13: int | str, campo14: int, monto: Importe, tipo: str) -> list[Any]:
    """`crear_linea_detalle` del script, tipo 2."""
    linea: list[Any] = [""] * COLUMNAS_POR_FILA
    linea[0] = 2
    linea[12] = campo13
    linea[13] = campo14
    linea[17] = monto
    linea[25] = tipo
    return linea


def _separadora() -> list[Any]:
    return [""] * COLUMNAS_POR_FILA


def _lineas_de_extorno(monto: Importe, tipo: str) -> list[list[Any]]:
    """El par de líneas de un extorno, idéntico en la comisión estándar y la de IZIPAY."""
    if monto > 0:
        return [
            _linea_detalle(40, CUENTA_BANCO, abs(monto), tipo),
            _linea_detalle(50, CUENTA_EXTORNO, abs(monto), tipo),
        ]
    return [
        _linea_detalle(40, CUENTA_EXTORNO, abs(monto), tipo),
        _linea_detalle(50, CUENTA_COMISION_POR_PAGAR, abs(monto), tipo),
    ]


def _lineas_de_recarga(
    *,
    fecha: date,
    fecha_cierre: date,
    monto: Importe,
    es_positivo: bool,
    cuenta: int,
    tipo: str,
) -> list[list[Any]]:
    """Las cuatro filas de una recarga. El signo lo decide quien llama (`>=` o `>`)."""
    cabecera = _linea_base(
        codigo="RG",
        fecha=fecha,
        fecha_cierre=fecha_cierre,
        descripcion="RECARGA POSITIVO" if es_positivo else "RECARGA NEGATIVO",
        tipo=tipo,
    )
    if es_positivo:
        linea2 = _linea_detalle("01", cuenta, abs(monto), tipo)
        linea3 = _linea_detalle(31, CUENTA_RECARGA, abs(monto), tipo)
    else:
        linea2 = _linea_detalle(11, cuenta, abs(monto), tipo)
        linea3 = _linea_detalle(21, CUENTA_RECARGA, abs(monto), tipo)
        linea2[20] = "Z001"
    return [cabecera, linea2, linea3, _separadora()]


# --- Lectura ------------------------------------------------------------------


@dataclass(frozen=True, slots=True)
class _Registro:
    """Una fila no vacía de una hoja de entrada: su número en Excel y sus campos."""

    fila: int
    valores: Mapping[str, object]


@dataclass(frozen=True, slots=True)
class _Hoja:
    nombre: str
    registros: tuple[_Registro, ...]


@dataclass(frozen=True, slots=True)
class _Libro:
    archivo: str
    tipo: str
    hojas: Mapping[str, _Hoja]


def _detectar_tipo(hojas: Sequence[str], archivo: str) -> str:
    """El tipo del archivo, decidido en este orden:

    1. Exactamente una hoja llamada como un tipo de tarjeta → ese tipo. Más de
       una es ambiguo y falla, sin mirar el nombre del archivo.
    2. Ninguna, pero hay `COMPENSACION` y además una hoja `Efectivo` o
       "EFECTIVO" en el nombre del archivo (sin distinguir mayúsculas) →
       EFECTIVO.
    3. La regla del script: el nombre del archivo contiene el tipo (sin
       distinguir mayúsculas, como el `glob` de Windows). Sirve para que un
       "AMEX.xlsx" sin su hoja falle con `hoja_faltante: AMEX` en vez de con
       un `tipo_no_reconocido` que no dice qué falta.

    Cualquier otro caso —incluido un nombre que contenga dos tipos— es
    `tipo_no_reconocido`. Los nombres de hoja se comparan EXACTOS, como los
    leía `pd.read_excel(sheet_name=...)`.
    """
    tarjetas = [tipo for tipo in TIPOS_DE_TARJETA if tipo in hojas]
    if len(tarjetas) == 1:
        return tarjetas[0]
    if len(tarjetas) > 1:
        raise ErrorContenido.tipo_no_reconocido(archivo=archivo)

    nombre = archivo.casefold()
    hoja_efectivo = any(hoja.casefold() == HOJA_EFECTIVO.casefold() for hoja in hojas)
    if HOJA_COMPENSACION in hojas and (hoja_efectivo or TIPO_EFECTIVO.casefold() in nombre):
        return TIPO_EFECTIVO

    por_nombre = [tipo for tipo in TIPOS if tipo.casefold() in nombre]
    if len(por_nombre) == 1:
        return por_nombre[0]
    raise ErrorContenido.tipo_no_reconocido(archivo=archivo)


def _es_vacio(valor: object) -> bool:
    return valor is None or valor == ""


def _leer_hoja(
    hoja: ReadOnlyWorksheet, columnas: Sequence[_Columna], archivo: str, *, con_registros: bool
) -> _Hoja:
    """Encabezado (fila 1) y filas no vacías de una hoja, por nombre de columna.

    `reset_dimensions()` es lo que hace pandas antes de leer en modo sólo
    lectura: la dimensión declarada en el XML puede mentir y recortar
    columnas. Las filas vacías se saltean una por una y sin cortar antes: la
    COMPENSACION real de EFECTIVO tiene ~711 mil filas con formato y su
    ÚLTIMA fila trae datos, así que ninguna heurística de "tantas vacías
    seguidas y corto" es equivalente. Saltearlas no cambia nada: en el script
    son filas NaN que ningún filtro acepta y ningún `groupby` agrupa.
    """
    hoja.reset_dimensions()
    filas = hoja.values
    encabezado = next(filas, ())
    indices = _indices_de_columnas(encabezado, columnas, archivo)
    if not con_registros:
        return _Hoja(nombre=hoja.title, registros=())

    campos = tuple(indices.items())
    registros: list[_Registro] = []
    for numero, fila in enumerate(filas, start=2):
        if fila.count(None) == len(fila):
            continue
        largo = len(fila)
        valores = {campo: (fila[i] if i < largo else None) for campo, i in campos}
        if all(_es_vacio(v) for v in valores.values()):
            continue
        registros.append(_Registro(fila=numero, valores=valores))
    return _Hoja(nombre=hoja.title, registros=tuple(registros))


def _leer_libro(entrada: ArchivoEntrada, *, con_registros: bool) -> _Libro:
    """Detecta el tipo, exige sus hojas y columnas y, si se pide, lee las filas.

    **Se abre un descriptor y se le pasa el objeto, nunca la ruta**: los
    temporales de `app/recepcion.py` no tienen extensión y `load_workbook`
    rechaza la ruta por eso antes de leer un byte (ver `_leer_filas` en
    `contado_carga`). Un archivo que ni siquiera abre como Excel es
    `cero_filas`, igual que allá. Una falla DESPUÉS de abrirlo no se traduce:
    no está prevista y se propaga.

    Con `con_registros=False` sólo se leen los encabezados: es lo que usa
    `validar`, y evita recorrer dos veces el millón de filas vacías de la
    hoja DINNERS real.
    """
    archivo = entrada.nombre_original
    with entrada.ruta_temporal.open("rb") as descriptor:
        try:
            libro = load_workbook(descriptor, read_only=True, data_only=True)
        except Exception as error:
            raise ErrorContenido.cero_filas(archivo=archivo) from error
        try:
            # Antes de cerrar el descriptor: en modo `read_only` el libro lee
            # perezosamente y lo sigue necesitando mientras se itera.
            tipo = _detectar_tipo(libro.sheetnames, archivo)
            requeridas = _hojas_requeridas(tipo)
            for nombre, _columnas in requeridas:
                if nombre not in libro.sheetnames:
                    raise ErrorContenido.hoja_faltante(archivo=archivo, hoja=nombre)
            hojas: dict[str, _Hoja] = {}
            for nombre, columnas in requeridas:
                hoja = libro[nombre]
                if not isinstance(hoja, ReadOnlyWorksheet):
                    # Una hoja de gráfico con el nombre esperado no es la hoja.
                    raise ErrorContenido.hoja_faltante(archivo=archivo, hoja=nombre)
                hojas[nombre] = _leer_hoja(hoja, columnas, archivo, con_registros=con_registros)
        finally:
            libro.close()
    return _Libro(archivo=archivo, tipo=tipo, hojas=MappingProxyType(hojas))


def _exigir_tipo_nuevo(libro: _Libro, vistos: set[str]) -> None:
    if libro.tipo in vistos:
        raise ErrorContenido.tipo_duplicado(archivo=libro.archivo)
    vistos.add(libro.tipo)


# --- Procesamiento ------------------------------------------------------------


@dataclass(frozen=True, slots=True)
class Observacion:
    """Algo que el script habría hecho en silencio, con dónde y con qué valor."""

    archivo: str
    hoja: str
    fila: int
    motivo: str
    valor: str


@dataclass(slots=True)
class _Contexto:
    """Lo que comparten las transformaciones de un mismo archivo."""

    archivo: str
    tipo: str
    hoy: date
    observaciones: list[Observacion] = field(default_factory=list)

    def observar(self, hoja: _Hoja, registro: _Registro, motivo: str, valor: object) -> None:
        self.observaciones.append(
            Observacion(
                archivo=self.archivo,
                hoja=hoja.nombre,
                fila=registro.fila,
                motivo=motivo,
                valor="" if valor is None else str(valor),
            )
        )

    def fecha_cierre(
        self, hoja: _Hoja, registro: _Registro, *, dia_primero: bool
    ) -> datetime | None:
        """La fecha de cierre de una fila, o `None`; ilegible se reporta y es `None`."""
        crudo = registro.valores["fecha_cierre"]
        if _es_nulo(crudo):
            return None
        fecha = _a_fecha(crudo, dia_primero=dia_primero)
        if fecha is None:
            self.observar(hoja, registro, MOTIVO_FECHA_CIERRE, crudo)
        return fecha


@dataclass(frozen=True, slots=True)
class _Movimiento:
    registro: _Registro
    categoria: str
    monto: Importe
    fecha_cierre: datetime | None


def _primera_fecha(fechas: Iterable[datetime | None], hoy: date) -> date:
    """La primera fecha de cierre no nula del grupo, o HOY (`datetime.now()`)."""
    return next((fecha for fecha in fechas if fecha is not None), hoy)


def _comision(hoja: _Hoja, contexto: _Contexto) -> list[list[Any]]:
    """`procesar_hoja_comision`: comisiones y extornos de la hoja de la tarjeta."""
    config = CONFIGURACIONES[contexto.tipo]["COMISION"]
    grupos: dict[datetime, list[_Movimiento]] = {}
    for registro in hoja.registros:
        categoria = _como_texto(registro.valores["categoria"])
        if not _FILTRO_COMISION.search(categoria):
            continue
        fecha = _a_fecha(registro.valores["fecha"], dia_primero=True)
        if fecha is None:
            contexto.observar(hoja, registro, MOTIVO_FECHA, registro.valores["fecha"])
            continue
        if "Comisión" not in categoria and "Extorno" not in categoria:
            # Sigue en el grupo: puede crear la cabecera y aportar el cierre.
            contexto.observar(hoja, registro, MOTIVO_CATEGORIA, categoria)
        grupos.setdefault(fecha, []).append(
            _Movimiento(
                registro=registro,
                categoria=categoria,
                monto=_a_importe(registro.valores["monto"]),
                fecha_cierre=contexto.fecha_cierre(hoja, registro, dia_primero=True),
            )
        )

    filas: list[list[Any]] = []
    for fecha in sorted(grupos):
        grupo = grupos[fecha]
        cierre = _primera_fecha((m.fecha_cierre for m in grupo), contexto.hoy)
        filas.append(
            _linea_base(
                codigo="AB",
                fecha=fecha,
                fecha_cierre=cierre,
                descripcion="COMISION",
                tipo=contexto.tipo,
            )
        )
        suma: Importe = 0
        for movimiento in grupo:
            if "Comisión" in movimiento.categoria and movimiento.monto != 0:
                monto = abs(movimiento.monto)
                filas.append(_linea_detalle(40, config["linea2_campo14"], monto, contexto.tipo))
                suma += monto
        if suma > 0:
            filas.append(
                _linea_detalle(
                    config["linea3_campo13"], config["linea3_campo14"], suma, contexto.tipo
                )
            )
        for movimiento in grupo:
            if "Extorno" in movimiento.categoria and movimiento.monto != 0:
                filas.extend(_lineas_de_extorno(movimiento.monto, contexto.tipo))
        filas.append(_separadora())
    return filas


def _cobranza(hoja: _Hoja, contexto: _Contexto) -> list[list[Any]]:
    """`procesar_hoja_compensacion`: una cobranza por fecha de documento de las filas DZ."""
    config = CONFIGURACIONES[contexto.tipo]["COMPENSACION"]
    grupos: dict[datetime, list[_Movimiento]] = {}
    for registro in hoja.registros:
        if not _FILTRO_DZ.search(_como_texto(registro.valores["clase"])):
            continue
        fecha = _a_fecha(registro.valores["fecha"], dia_primero=True)
        if fecha is None:
            contexto.observar(hoja, registro, MOTIVO_FECHA, registro.valores["fecha"])
            continue
        grupos.setdefault(fecha, []).append(
            _Movimiento(
                registro=registro,
                categoria="",
                monto=abs(_a_importe(registro.valores["importe"])),
                fecha_cierre=contexto.fecha_cierre(hoja, registro, dia_primero=True),
            )
        )

    filas: list[list[Any]] = []
    for fecha in sorted(grupos):
        grupo = grupos[fecha]
        monto = _suma_de_grupo([m.monto for m in grupo])
        cierre = _primera_fecha((m.fecha_cierre for m in grupo), contexto.hoy)
        linea3 = _linea_detalle(
            config["linea3_campo13"], config["linea3_campo14"], monto, contexto.tipo
        )
        linea3[20] = "Z001"
        filas.extend(
            [
                _linea_base(
                    codigo="DZ",
                    fecha=fecha,
                    fecha_cierre=cierre,
                    descripcion="COBRANZA",
                    tipo=contexto.tipo,
                ),
                _linea_detalle(40, config["linea2_campo14"], monto, contexto.tipo),
                linea3,
                _separadora(),
            ]
        )
    return filas


def _recarga(hoja: _Hoja, contexto: _Contexto) -> list[list[Any]]:
    """`procesar_hoja_ajuste`: una recarga por fila RG marcada como AJUSTE."""
    cuenta = CONFIGURACIONES[contexto.tipo]["AJUSTE"]["campo14"]
    filas: list[list[Any]] = []
    for registro in hoja.registros:
        if not (
            _FILTRO_RG.search(_como_texto(registro.valores["clase"]))
            and _FILTRO_AJUSTE.search(_como_texto(registro.valores["observacion"]))
        ):
            continue
        fecha = _a_fecha(registro.valores["fecha"], dia_primero=True)
        if fecha is None:
            # El script se caía acá (`NaT.strftime`) y perdía la hoja entera.
            contexto.observar(hoja, registro, MOTIVO_FECHA, registro.valores["fecha"])
            continue
        monto = _a_importe(registro.valores["importe"])
        cierre = contexto.fecha_cierre(hoja, registro, dia_primero=True) or contexto.hoy
        filas.extend(
            _lineas_de_recarga(
                fecha=fecha,
                fecha_cierre=cierre,
                monto=monto,
                es_positivo=monto >= 0,
                cuenta=cuenta,
                tipo=contexto.tipo,
            )
        )
    return filas


def _fecha_estricta(crudo: object) -> datetime | None:
    """Las lecturas de IZIPAY: el valor crudo, leído sin `dayfirst`."""
    return None if _es_nulo(crudo) else _a_fecha(crudo, dia_primero=False)


def _cuenta_izipay(categoria: str) -> int:
    """La primera subcadena de `CODIGO_CATEGORIAS_IZIPAY` que aparezca, o 0."""
    return next((c for clave, c in CODIGO_CATEGORIAS_IZIPAY.items() if clave in categoria), 0)


def _izipay_comision(hoja: _Hoja, contexto: _Contexto) -> list[list[Any]]:
    """`procesar_izipay_comision`: agrupa TODAS las filas por fecha, sin filtro previo."""
    grupos: dict[datetime, list[_Movimiento]] = {}
    for registro in hoja.registros:
        categoria = _como_texto(registro.valores["categoria"])
        monto = _a_importe(registro.valores["monto"])
        fecha = _fecha_estricta(registro.valores["fecha"])
        if fecha is None:
            if monto != 0 and ("Comisión" in categoria or "Extorno" in categoria):
                contexto.observar(hoja, registro, MOTIVO_FECHA, registro.valores["fecha"])
            continue
        if _FILTRO_COMISION.search(categoria) and not (
            "Comisión" in categoria or "Extorno" in categoria
        ):
            contexto.observar(hoja, registro, MOTIVO_CATEGORIA, categoria)
        grupos.setdefault(fecha, []).append(
            _Movimiento(registro=registro, categoria=categoria, monto=monto, fecha_cierre=None)
        )

    filas: list[list[Any]] = []
    for fecha in sorted(grupos):
        grupo = grupos[fecha]
        # La PRIMERA fila del día decide, esté vacía o no (`grupo.iloc[0, 4]`).
        cierre = contexto.fecha_cierre(hoja, grupo[0].registro, dia_primero=False)
        filas.append(
            _linea_base(
                codigo="AB",
                fecha=fecha,
                fecha_cierre=cierre or contexto.hoy,
                descripcion="COMISION",
                tipo=contexto.tipo,
            )
        )
        suma: Importe = 0
        for movimiento in grupo:
            monto = abs(movimiento.monto)
            if "Comisión" in movimiento.categoria and monto > 0:
                cuenta = _cuenta_izipay(movimiento.categoria)
                if cuenta == 0:
                    contexto.observar(
                        hoja, movimiento.registro, MOTIVO_CUENTA, movimiento.categoria
                    )
                filas.append(_linea_detalle(40, cuenta, monto, contexto.tipo))
                suma += monto
        if suma > 0:
            filas.append(_linea_detalle(50, CUENTA_COMISION_POR_PAGAR, suma, contexto.tipo))
        for movimiento in grupo:
            if "Extorno" in movimiento.categoria and movimiento.monto != 0:
                filas.extend(_lineas_de_extorno(movimiento.monto, contexto.tipo))
        filas.append(_separadora())
    return filas


def _importes(registros: Iterable[_Registro]) -> list[Importe]:
    return [_a_importe(registro.valores["importe"]) for registro in registros]


def _izipay_cobranza(hoja: _Hoja, contexto: _Contexto) -> list[list[Any]]:
    """`procesar_izipay_compensacion`: agrupa por `fecha dia`, no por fecha de documento.

    Una `fecha dia` VACÍA no se reporta: es como el operador deja una fila
    fuera del asiento del día (585 de las ~900 filas de la muestra real de
    agosto de 2025, incluidas 63 DZ de días anteriores). Sólo una `fecha dia`
    escrita e ilegible es una pérdida, y ésa sí va a las observaciones.
    """
    grupos: dict[datetime, list[_Registro]] = {}
    cierres: dict[int, datetime | None] = {}
    for registro in hoja.registros:
        crudo = registro.valores["fecha_dia"]
        fecha = _fecha_estricta(crudo)
        if fecha is None:
            if not _es_nulo(crudo):
                contexto.observar(hoja, registro, MOTIVO_FECHA, crudo)
            continue
        cierres[registro.fila] = contexto.fecha_cierre(hoja, registro, dia_primero=True)
        grupos.setdefault(fecha, []).append(registro)

    filas: list[list[Any]] = []
    for fecha in sorted(grupos):
        grupo = grupos[fecha]
        cierre = _primera_fecha((cierres[r.fila] for r in grupo), contexto.hoy)
        filas.append(
            _linea_base(
                codigo="DZ",
                fecha=fecha,
                fecha_cierre=cierre,
                descripcion="COBRANZA",
                tipo=contexto.tipo,
            )
        )
        importes_dz = _importes(r for r in grupo if r.valores["clase"] == "DZ")
        if importes_dz:
            filas.append(
                _linea_detalle(40, CUENTA_BANCO, abs(_suma_de_serie(importes_dz)), contexto.tipo)
            )
        for categoria, cuenta in CODIGO_COMPENSACION_IZIPAY.items():
            textos = (
                _TEXTOS_MASTERCARD_MAS_IZIPAY
                if categoria == _MASTERCARD_MAS_IZIPAY
                else frozenset({categoria})
            )
            importes = _importes(r for r in grupo if r.valores["texto"] in textos)
            if not importes:
                continue
            suma = _suma_de_serie(importes)
            clave: int | str = 11
            if categoria == _MASTERCARD_MAS_IZIPAY and suma < 0:
                clave = "01"
            linea = _linea_detalle(clave, cuenta, abs(suma), contexto.tipo)
            if clave != "01":
                linea[20] = "Z001"
            filas.append(linea)
        filas.append(_separadora())
    return filas


def _izipay_recarga(hoja: _Hoja, contexto: _Contexto) -> list[list[Any]]:
    """`procesar_izipay_ajustes`: una recarga por fila `AJUSTE`, sin mirar la clase."""
    filas: list[list[Any]] = []
    for registro in hoja.registros:
        if registro.valores["observacion"] != "AJUSTE":
            continue
        fecha = _fecha_estricta(registro.valores["fecha"])
        if fecha is None:
            # El script se caía acá (`pd.to_datetime` sin `coerce`).
            contexto.observar(hoja, registro, MOTIVO_FECHA, registro.valores["fecha"])
            continue
        monto = _a_importe(registro.valores["importe"])
        cierre = contexto.fecha_cierre(hoja, registro, dia_primero=False) or contexto.hoy
        filas.extend(
            _lineas_de_recarga(
                fecha=fecha,
                fecha_cierre=cierre,
                monto=monto,
                es_positivo=monto > 0,
                cuenta=CUENTA_RECARGA_IZIPAY,
                tipo=contexto.tipo,
            )
        )
    return filas


def _transformar(libro: _Libro, contexto: _Contexto) -> list[tuple[str, list[list[Any]]]]:
    """Las hojas de salida NO vacías de un libro, en el orden del script."""
    compensacion = libro.hojas[HOJA_COMPENSACION]
    if libro.tipo == TIPO_IZIPAY:
        hojas = [
            (HOJA_COMISION, _izipay_comision(libro.hojas[TIPO_IZIPAY], contexto)),
            (HOJA_COBRANZA, _izipay_cobranza(compensacion, contexto)),
            (HOJA_RECARGA, _izipay_recarga(compensacion, contexto)),
        ]
    elif libro.tipo == TIPO_EFECTIVO:
        hojas = [
            (HOJA_COBRANZA, _cobranza(compensacion, contexto)),
            (HOJA_RECARGA, _recarga(compensacion, contexto)),
        ]
    else:
        hojas = [
            (HOJA_COMISION, _comision(libro.hojas[libro.tipo], contexto)),
            (HOJA_COBRANZA, _cobranza(compensacion, contexto)),
            (HOJA_RECARGA, _recarga(compensacion, contexto)),
        ]
    return [(nombre, filas) for nombre, filas in hojas if filas]


# --- Escritura ----------------------------------------------------------------


def _escribir_xlsx(hojas: Sequence[tuple[str, list[list[Any]]]], ruta: Path) -> None:
    """Una hoja por salida no vacía, sin encabezado ni índice.

    Las celdas `""` se escriben como `""`, igual que las escribía pandas:
    openpyxl las emite como celdas vacías, que se leen de vuelta como `None`,
    y la separadora conserva sus 26 columnas.
    """
    libro = Workbook(write_only=True)
    for nombre, filas in hojas:
        hoja = libro.create_sheet(title=nombre)
        for fila in filas:
            hoja.append(fila)
    libro.save(ruta)


def _limpio(texto: str) -> str:
    return texto.replace("\t", " ").replace("\r", " ").replace("\n", " ")


def _escribir_observaciones(observaciones: Sequence[Observacion], ruta: Path) -> None:
    """Reporte de observaciones, una por línea, con `FIN_DE_LINEA` explícito."""
    conteo: dict[str, int] = {}
    for observacion in observaciones:
        conteo[observacion.motivo] = conteo.get(observacion.motivo, 0) + 1
    resumen = ", ".join(f"{motivo}={total}" for motivo, total in sorted(conteo.items()))

    lineas = [
        f"Observaciones: {len(observaciones)} ({resumen})",
        "",
        "\t".join(("archivo", "hoja", "fila", "motivo", "valor")),
    ]
    lineas.extend(
        "\t".join((_limpio(o.archivo), _limpio(o.hoja), str(o.fila), o.motivo, _limpio(o.valor)))
        for o in observaciones
    )
    with ruta.open("w", encoding=CODIFICACION, newline="") as salida:
        salida.write(FIN_DE_LINEA.join(lineas) + FIN_DE_LINEA)


def _filas_de_asiento(hojas: Sequence[tuple[str, list[list[Any]]]]) -> int:
    """Líneas tipo 1 y tipo 2 escritas; las separadoras no cuentan."""
    return sum(1 for _nombre, filas in hojas for fila in filas if fila[0] != "")


class AsientosContables(Procesador):
    """El procesador real: hasta siete Excel entran, un Excel por archivo sale.

    Más un `observaciones.txt` cuando —y sólo cuando— hubo algo que reportar.
    `app.core.empaquetado` decide si eso viaja suelto o comprimido; este módulo
    nunca arma un ZIP (ADR 0006).

    `hoy` es el reloj de la fecha de cierre por defecto. La instancia del
    registro usa `date.today`; las pruebas y la paridad inyectan un día fijo.
    Se consulta UNA vez por ejecución, así que todos los archivos de una
    misma ejecución comparten el día aunque crucen la medianoche.
    """

    def __init__(self, clave: str = CLAVE, hoy: Callable[[], date] = date.today) -> None:
        self.clave = clave
        self._hoy = hoy

    def validar(self, archivos: list[ArchivoEntrada]) -> ErrorTipificado | None:
        """Tipo, hojas, columnas y unicidad de tipo, leyendo SÓLO los encabezados.

        Devuelve el error en vez de levantarlo, como fija `Procesador`.
        """
        vistos: set[str] = set()
        try:
            for entrada in archivos:
                _exigir_tipo_nuevo(_leer_libro(entrada, con_registros=False), vistos)
        except ErrorContenido as error:
            return error
        return None

    def procesar(self, archivos: list[ArchivoEntrada]) -> list[ArchivoSalida]:
        """Transforma cada archivo y devuelve un Excel por archivo.

        Todo se calcula antes de escribir el primer byte: un archivo con cero
        filas de salida levanta `ErrorContenido.cero_filas` y no queda ningún
        Excel a medias. Ninguna excepción se traga; una fila que rompa de una
        forma no prevista se propaga y `app/core/ejecucion.py` la convierte en
        `FalloDelModulo`.
        """
        hoy = self._hoy()
        vistos: set[str] = set()
        resultados: list[tuple[_Libro, list[tuple[str, list[list[Any]]]], _Contexto]] = []
        for entrada in archivos:
            libro = _leer_libro(entrada, con_registros=True)
            _exigir_tipo_nuevo(libro, vistos)
            contexto = _Contexto(archivo=libro.archivo, tipo=libro.tipo, hoy=hoy)
            hojas = _transformar(libro, contexto)
            if not hojas:
                raise ErrorContenido.cero_filas(archivo=libro.archivo)
            resultados.append((libro, hojas, contexto))

        directorio = archivos[0].ruta_temporal.parent
        salidas: list[ArchivoSalida] = []
        for libro, hojas, _contexto in resultados:
            ruta = directorio / nombre_de_salida(libro.tipo)
            _escribir_xlsx(hojas, ruta)
            salidas.append(
                ArchivoSalida(nombre_propuesto=ruta.name, ruta_temporal=ruta, tipo_mime=MIME_XLSX)
            )

        observaciones = [o for _libro, _hojas, c in resultados for o in c.observaciones]
        if observaciones:
            ruta_observaciones = directorio / NOMBRE_OBSERVACIONES
            _escribir_observaciones(observaciones, ruta_observaciones)
            salidas.append(
                ArchivoSalida(
                    nombre_propuesto=NOMBRE_OBSERVACIONES,
                    ruta_temporal=ruta_observaciones,
                    tipo_mime=MIME_TEXTO,
                )
            )
            # El logging se configura ACÁ y no al importar: este código corre
            # en el hijo de `spawn`, que no hereda la configuración del padre
            # (ver `contado_carga`). Una línea por archivo con observaciones.
            configurar_logging()
            for libro, hojas, contexto in resultados:
                if contexto.observaciones:
                    registrar_descartes(
                        clave=self.clave,
                        archivo=libro.archivo,
                        filas_procesadas=_filas_de_asiento(hojas),
                        descartes=[(o.motivo, o.fila) for o in contexto.observaciones],
                    )
        return salidas
